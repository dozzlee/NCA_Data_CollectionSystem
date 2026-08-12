from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


SOURCE = Path(r"C:\Users\dozzl\Downloads\Product Requirements Document - Development Ready.docx")
OUTPUT = Path(r"C:\Users\dozzl\OneDrive\Documents\testing and fixing\reports\Product Requirements Document - Development Ready - Updated.docx")


def cell(table, row, column, text):
    target = table.rows[row].cells[column]
    target.text = text
    for paragraph in target.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)


def add_row(table, values):
    row = table.add_row()
    for index, value in enumerate(values):
        row.cells[index].text = value
        for paragraph in row.cells[index].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
    return row


def replace_paragraph(document, index, text):
    document.paragraphs[index].text = text


def bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


def set_table_header_repeat(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


if not SOURCE.exists():
    raise FileNotFoundError(SOURCE)

doc = Document(SOURCE)
tables = doc.tables

# Document identity and planning decisions.
cell(tables[0], 1, 1, "1.2 - Delivered-system publication remediation update")
cell(tables[0], 2, 1, "August 12, 2026")
cell(tables[0], 3, 1, "National Communications Authority publication readiness and operational acceptance")
cell(tables[0], 4, 1, "Align the PRD with the delivered five-role system, seven corrected Section 11 forms, governed workflows, security controls, private files, data requests, aggregate dashboard and explicit release gates.")
cell(tables[0], 5, 1, "Original development-ready PRD and Section 11 requirements, preserved historical system records, corrected PRD-derived form manifests, implementation tests and publication-readiness evidence. Exact source-workbook fidelity remains subject to approved original regulatory forms.")

cell(tables[1], 0, 0, "The delivered release is a governed regulatory data collection, provider approval, NCA review, compliance and controlled data-delivery system. It prioritizes manual structured entry, immutable form/submission versions, targeted corrections, validation, private scanned files, role-sanitized aggregate dashboard access, approved-data exports, notifications, auditability and recovery evidence. General spreadsheet import, GIS exploration, real email sending, MFA and automated retention deletion are not represented as delivered capabilities. Private Excel backup upload is a secured operational extension and is not a data-import path.")

cell(tables[2], 2, 1, "Use a configurable, immutable form-family/version engine with sections, scalar fields, fixed/repeatable grids, options, allow-listed validations, formulas, conditions, declarations, Section 11 gaps and form-specific KMZ rules. Historical versions remain renderable.")
cell(tables[2], 4, 1, "Operational and request exports use canonical long/narrow rows. The Industry Dashboard exposes server-authorized aggregates and a generated aggregate XLSX; source workbooks and provider-level dashboard rows remain NCA-only.")
cell(tables[2], 6, 1, "V1 remains manual web-form entry. A private Excel backup may be uploaded or replaced while editable as secured evidence; it is quarantined, scanned and never imported into form fields. General document/spreadsheet intake remains deferred.")
cell(tables[2], 7, 1, "The system creates approved/versioned message content and durable drafts/queue entries. Without configured delivery credentials and provider evidence, a message is never represented as SENT. Microsoft Graph/shared-mailbox delivery is advisory/deferred.")
cell(tables[2], 8, 1, "Only DC-DBS05 requires topology KMZ in the corrected Section 11 set. KMZ and Excel backups are private, hashed, quarantined and download-gated by malware scan and role authorization. GIS rendering remains deferred.")
cell(tables[2], 10, 1, "Five stored roles are retained: NCA Admin; NCA Officer; Provider Data Entry; Provider Approver; and NCA Data Requester (stored as NCA_VIEWER for compatibility). NCA Officer has operational Admin parity except Users/divisions and Data Requests.")

# Five-role permissions table.
role_rows = [
    ("NCA Admin", "Complete system access: users/divisions, providers, form versions and publication, assignments, periods, submissions, compliance, exports, governance, support, audit and Data Requests/artifacts.", "Cannot act as a provider user; high-risk actions remain audited and subject to readiness gates."),
    ("NCA Officer", "Operational Admin parity for forms, providers, assignments, periods, review, compliance, exports, dashboard configuration, governance, audit and support.", "No Users/division/account management and no Data Request queue, decisions or artifacts."),
    ("Provider Data Entry", "Own-organization draft entry, scalar/grid snapshots, validation, allowed private uploads, provider correction/resubmission and technical support.", "Cannot officially submit to NCA, see another provider, or use compliance/general contact actions."),
    ("Provider Approver", "Own-organization approver queue, complete review, targeted provider corrections, official submission, compliance contact and technical support.", "Cannot perform NCA review, internal exports, user administration or cross-provider access."),
    ("NCA Data Requester", "Metadata-only catalog, own requests, portal notifications and authenticated released-artifact download; named account with managed division and required grade.", "No operational submissions, providers, compliance, periods, raw exports, dashboard source/detail or another requester's records."),
]
while len(tables[4].rows) > 1:
    tables[4]._tbl.remove(tables[4].rows[-1]._tr)
for values in role_rows:
    add_row(tables[4], values)

# Form inventory alignment.
versions = {
    "MNO-MONTHLY": "Corrected PRD-derived immutable version 3.0",
    "DC-TB02": "Corrected PRD-derived immutable version 3.0",
    "DC-ISP06": "Corrected PRD-derived immutable version 3.0",
    "DC-ITC04": "Corrected PRD-derived immutable version 3.0",
    "TOWER-MAIN-ANNUAL": "Corrected PRD-derived immutable version 2.0; code remains visibly provisional",
    "DC-DBS05": "Corrected PRD-derived immutable version 2.0; topology KMZ required",
    "DC-SUB03": "Corrected PRD-derived immutable version 2.0; no KMZ requirement",
}
for row in tables[5].rows[1:]:
    code = row.cells[0].text.strip()
    if code in versions:
        row.cells[3].text = versions[code]
        if code == "DC-SUB03":
            row.cells[4].text = "Separate global and Ghana capacity metrics, client categories, fibre-cut events, checklist, comments and declaration. Future source attachments are non-KMZ and remain deferred."
        elif code == "DC-DBS05":
            row.cells[4].text += " This is the only corrected form that requires topology KMZ."

cell(tables[6], 1, 1, "Not Started, Draft, Pending Provider Approval, Provider Changes Requested, Provider Resubmitted, Submitted, Under NCA Review, Correction Requested, Resubmitted, Approved, Rejected, Archived")
cell(tables[6], 1, 2, "Provider and regulatory stages are explicit. Official provider and NCA versions are immutable; regulatory correction creates a cloned version and locks unaffected content.")

# Data entities extended without removing original entries.
cell(tables[7], 1, 2, "user_id, organization_id, role, name, unique email, status, division_id and grade for requesters, failed_login_attempts, locked_until, must_change_password, last_login_at")
add_row(tables[7], ("SubmissionEvent / SubmissionNotification", "Immutable provider/regulatory workflow history and own-user portal notification stream.", "submission_id, event_type, status transition, actor, audience, reason/metadata, recipient, read_at, created_at"))
add_row(tables[7], ("DataRequest / Artifact", "Governed division request, immutable approval manifest and private released file.", "requester snapshots, division, grade, purpose, scope, status, due date, manifest, format, hash, size, rows, expiry"))
add_row(tables[7], ("PrivateUpload / Scan", "Quarantined KMZ and Excel backup metadata outside public media.", "submission_id, requirement, storage reference, SHA-256, scan status/result, uploader, timestamps"))

cell(tables[9], 2, 2, "Declarations and applicable structured values are required before transition. General audited-financial-statement attachments remain deferred; private Excel backup is optional evidence, not a form value.")
cell(tables[9], 9, 2, "Topology KMZ is required only for DC-DBS05 in the corrected form set and must be clean or explicitly accepted before approval.")

# Section 11 corrections.
cell(tables[11], 8, 2, "Source formulas/references are converted into explicit regional grids and calculated totals. TOWER-MAIN-ANNUAL is published as a visibly provisional code; an official replacement creates a new family/version rather than rewriting history.")
cell(tables[11], 10, 2, "This is the only corrected form with required topology KMZ. It also uses fixed Ghana-region grids, exact fibre-cut classifications, prescribed km/Gbps/GHS/hours/count units, and percentage validation.")
cell(tables[11], 12, 2, "Global cable and Ghana segment metrics remain distinct. Section 11 future attachments are non-KMZ; no KMZ requirement is configured for DC-SUB03.")

# Portal-first notifications and disabled delivery claims.
cell(tables[12], 1, 2, "Creates portal notification and, when applicable, an approved outbound message queue item. No message is marked sent without a configured provider and delivery evidence.")
cell(tables[12], 5, 2, "Portal notification is immediate and contains targeted correction instructions. Email remains a draft/queue item unless delivery is configured.")
cell(tables[12], 6, 2, "Immutable receipt/event and portal notifications are created transactionally. Provider and assigned NCA users see audience-appropriate status updates.")

# Release plan and backlog statements.
cell(tables[13], 1, 2, "Seven corrected immutable versions pass objective Section 11 structural checks. Historical versions remain readable; Tower code is visibly provisional and exact source-workbook fidelity is not claimed.")
cell(tables[13], 3, 1, "Explicit provider state machine, targeted correction/resubmission, tenant-isolated approver queue, full validation and immutable provider workflow events/notifications.")
cell(tables[13], 4, 1, "Exact-version NCA review, cloned regulatory corrections, unaffected-content locks, non-filled dispositions, receipts and approval blockers.")
cell(tables[13], 5, 1, "Portal compliance workflow plus provider-neutral message drafts/queue. Manual Sent controls are removed; live delivery remains disabled without credentials/evidence.")
cell(tables[13], 6, 1, "Approved-only operational CSV/PDF, aggregate dashboard XLSX, private file scan/download controls, cookie/CSRF authentication, audits, recovery tooling and publication evidence.")

cell(tables[14], 1, 1, "Five-role login/logout, HttpOnly cookie JWT, CSRF, rotation/revocation, lockout, Admin temporary-password reset, mandatory password change and provider/requester isolation. MFA is deferred.")
cell(tables[14], 3, 1, "Immutable family/version management, source map, preview, grid columns/fixed rows/options/conditions/formulas/validated rules, DBS05 KMZ, strict gaps and publication gates.")
cell(tables[14], 6, 1, "Private, quarantined, hashed and scan-gated DBS05 topology KMZ; authenticated role-scoped download. Other corrected forms have no topology KMZ requirement.")
cell(tables[14], 7, 1, "Submit to Approver, targeted provider correction, resubmit and official submission with immutable events, reasons, targets, timestamps and portal notifications.")
cell(tables[14], 10, 1, "Approved-only canonical operational CSV/PDF, request CSV/XLSX/PDF artifacts and role-sanitized aggregate dashboard XLSX; all authenticated and audited.")

# Open questions are now decisions or explicit external gates.
decisions = {
    1: ("What is the official form code for Infrastructure Main Companies annual workbook?", "TOWER-MAIN-ANNUAL remains visibly provisional. An approved code creates a new family/version; historical references are not rewritten.", "NCA data owner - external release decision"),
    2: ("Should operational exports include submitted-but-unapproved records?", "Resolved: operational exports include only NCA-approved submission versions.", "Resolved in delivered system"),
    3: ("Should provider users be able to download official receipts?", "Resolved: authenticated provider receipt download is supported and audited when the receipt is available.", "Resolved in delivered system"),
    4: ("What is the approved penalty/escalation wording?", "Still required before penalty wording is enabled in approved templates.", "NCA Legal/compliance - external input"),
    5: ("What NCA operational scope applies?", "Resolved for this release: NCA Officer has operational parity, with explicit Users/divisions and Data Request exclusions.", "Resolved in delivered system"),
    6: ("What are the approved retention schedules and recovery targets?", "Still required. Automated disposition remains globally disabled; production requires approved RPO/RTO and verified restore evidence.", "NCA IT/security/legal - release gate"),
    7: ("Is live Microsoft Graph/shared-mailbox delivery required for launch?", "Advisory/deferred. Without credentials and delivery evidence, messages remain DRAFT or QUEUED and are never represented as sent.", "Product owner and IT - deferred integration"),
    8: ("Which corrected forms require topology KMZ?", "Resolved: DC-DBS05 only. File-size and malware policies remain configuration-controlled.", "Resolved in delivered system"),
}
for row_number, values in decisions.items():
    for column, value in enumerate(values):
        cell(tables[15], row_number, column, value)

# Correct conflicting narrative requirements while retaining the original topology.
replace_paragraph(doc, 19, "Allow providers to manually enter data, save transactional section snapshots, pass validation, route a draft to their own Provider Approver, upload private evidence where allowed, and receive portal confirmation.")
replace_paragraph(doc, 20, "Allow NCA reviewers to inspect the submission's exact immutable form version, scalar and grid values, clean uploads, statuses, explanations, validation evidence, Section 11 coverage, timeline and correction comparisons.")
replace_paragraph(doc, 21, "Allow privileged NCA staff to identify overdue or incomplete submissions, generate approved message drafts/queue entries and preserve the compliance trail. Live sending is not required for this release.")
replace_paragraph(doc, 22, "Support approved-only operational CSV/PDF, data-request CSV/XLSX/PDF artifacts, and aggregate-only dashboard XLSX with authenticated audited delivery.")
replace_paragraph(doc, 23, "Maintain hash-chained audit evidence for authentication, configuration, workflow, validation, upload/scan/download, compliance, support, export, governance and recovery actions.")
replace_paragraph(doc, 26, "Advanced intelligence, AI anomaly detection and executive analytics beyond the delivered aggregate Industry Dashboard and operational metrics.")
replace_paragraph(doc, 31, "Provider/raw Excel import remains out of scope. The delivered aggregate dashboard XLSX, approved request XLSX and secured Excel backup extension are explicitly governed exceptions.")
replace_paragraph(doc, 33, "Public or unauthenticated dashboards and disclosure of source workbooks/provider-level dashboard rows.")
replace_paragraph(doc, 45, "NCA Officer is a distinct stored role with operational Admin parity except Users/divisions and Data Requests. Requesters use named individual accounts with a managed division and required grade.")
replace_paragraph(doc, 46, "Passwords use approved hashing, strong validation, timeout, failed-login counting, throttling and lockout. Authentication uses rotating/revoked JWTs in Secure, HttpOnly, SameSite cookies with CSRF protection. MFA is explicitly deferred and is not represented as enforced.")
replace_paragraph(doc, 65, "A draft is editable only by authorized Data Entry users while its provider-stage status permits entry. Provider and regulatory correction stages reopen only targeted content; official historical versions remain immutable.")
replace_paragraph(doc, 66, "Provider and NCA transitions rerun full validation. Approval is blocked by required content, clean required uploads, blocking rules, unresolved dispositions, open correction items and failed high/blocker Section 11 requirements unless an authorized audited exception applies.")
replace_paragraph(doc, 68, "Provider Approvers review tenant-isolated drafts, complete values, validation, explanations and clean allowed files before official submission.")
replace_paragraph(doc, 69, "Approvers can request a targeted provider correction with reason and section/field/grid-cell targets, or approve after a fresh readiness check. Returned drafts preserve decision history and resubmission count.")
replace_paragraph(doc, 70, "Official provider submission creates an immutable version. Regulatory correction clones it into a new version; the official source version is never edited.")
replace_paragraph(doc, 71, "Provider transitions create immutable workflow events, responsible-user/timestamp evidence, audience-scoped portal notifications and audit records.")
replace_paragraph(doc, 73, "NCA review always renders the exact stored form template/version, including scalar/calculated values, fixed and repeatable grids, statuses, explanations, uploads/scans, validation results, Section 11 results and blockers.")
replace_paragraph(doc, 75, "Regulatory correction targets sections, fields or individual grid cells. A cloned version carries forward values and instructions; unaffected content is locked in API and UI and old/new comparison is retained.")
replace_paragraph(doc, 76, "Provider resubmission after regulatory correction uses the cloned version and returns through Provider Approver before reaching NCA as Resubmitted.")
replace_paragraph(doc, 78, "Corrected Section 11 topology KMZ applies only to DC-DBS05. Private Excel backup is also permitted as secured evidence while editable; neither path imports values into the form.")
replace_paragraph(doc, 79, "Each private upload records provider, submission, exact form/period, requirement/category, uploader/time, original name, size, private storage reference, SHA-256, quarantine/scan state and review state.")
replace_paragraph(doc, 80, "General document, CSV and spreadsheet import remains deferred. DC-SUB03 future attachments are non-KMZ and do not create a topology KMZ rule.")
replace_paragraph(doc, 81, "Private files are type/size validated, quarantined and malware-scanned. Downloads are blocked until clean/accepted, streamed only through authenticated authorization and audited; `/media/` never serves them in production.")
replace_paragraph(doc, 84, "Providers enter canonical values directly in scalar fields and fixed/repeatable grids. The optional Excel backup is evidence only and is not parsed or mapped into form values.")
replace_paragraph(doc, 85, "Bulk spreadsheet/CSV import, general document intake and automated mapping remain future capabilities. The private Excel backup extension and DBS05 topology KMZ are not import workflows.")
replace_paragraph(doc, 89, "The system renders approved template content and durable outbound message records for notices, reminders, missing data/KMZ, correction and escalation. Portal notifications handle workflow updates independently.")
replace_paragraph(doc, 90, "Without configured delivery credentials and provider delivery evidence, messages remain DRAFT or QUEUED and never SENT. The misleading manual Sent action is removed.")
replace_paragraph(doc, 91, "The provider-neutral delivery interface and event/webhook contracts allow future Microsoft Graph shared-mailbox integration without changing workflow records.")
replace_paragraph(doc, 93, "Operational dashboards retain submission/compliance metrics. The workbook-driven Industry Dashboard is an intentional shared aggregate feature: every authenticated role may view/filter aggregates and generate aggregate XLSX.")
replace_paragraph(doc, 94, "Dashboard data and capabilities are authorized server-side. Only NCA Admin/Officer can update source data, use chart-building/configuration tools or view source/provenance/provider-level rows.")
replace_paragraph(doc, 97, "Operational approved-data exports are CSV and PDF. Data-request artifacts may be CSV, XLSX or PDF. Industry Dashboard export is a system-generated aggregate XLSX.")
replace_paragraph(doc, 100, "Every export and private-file download is authenticated, role/tenant scoped and audited with actor, filters, record count, hash/reference and timestamp.")
replace_paragraph(doc, 109, "Provider Data Entry submits a validated draft to Provider Approver (Pending Provider Approval).")
replace_paragraph(doc, 110, "Provider Approver requests targeted changes, receives Provider Resubmitted, or approves to create the official Submitted version.")
replace_paragraph(doc, 111, "NCA explicitly starts review, then approves, rejects or requests targeted regulatory correction.")
replace_paragraph(doc, 112, "A regulatory correction creates a cloned version with unaffected content locked; it returns through provider approval and reaches NCA as Resubmitted.")
replace_paragraph(doc, 132, "Production runs behind NCA-managed TLS ingress with strict hosts/origins, HTTPS redirect, HSTS, secure cookies, CSRF, frame denial, safe secrets and trusted proxy CIDRs. Unsafe production configuration fails startup.")
replace_paragraph(doc, 133, "KMZ, Excel backups, receipts, operational exports and data-request artifacts are stored outside public media and downloaded only through authenticated streaming endpoints.")
replace_paragraph(doc, 134, "Central audit emission covers authentication, users/configuration, forms, periods, workflow/corrections, validations, overrides, uploads/scans/downloads, messages, compliance, exports, support, retention and recovery; events are hash chained and daily anchored.")
replace_paragraph(doc, 135, "Back up PostgreSQL/WAL and private files to approved encrypted immutable storage and verify restoration in isolation against configured RPO/RTO. Credentials, targets and completed recovery evidence are external release inputs.")
replace_paragraph(doc, 137, "Automated disposition remains globally disabled until approved record-class retention policies exist. Legal holds override disposition, and every approved run begins with an auditable dry run.")
replace_paragraph(doc, 145, "The seven PRD-derived versions are structurally verified and published. TOWER-MAIN-ANNUAL remains visibly provisional pending an approved official code; a future decision creates a new family/version.")
replace_paragraph(doc, 146, "Section 11 structural maps, units, conditions, formulas, validations, declarations and DBS05-only topology KMZ are enforced. Exact source-workbook fidelity remains pending approved original forms.")
replace_paragraph(doc, 147, "MFA and live Microsoft Graph delivery are documented deferred work. Manual web entry, private Excel backup evidence and DBS05 topology KMZ are accepted release boundaries.")
replace_paragraph(doc, 148, "Security controls are implemented. Approved retention schedules, production domains/secrets, trusted ingress, immutable backup destination, RPO/RTO, restore evidence and formal UAT are external readiness gates.")
replace_paragraph(doc, 151, "Providers can enter scalar/grid data, save optimistic snapshots, complete provider correction/resubmission and have a Provider Approver make the official submission.")
replace_paragraph(doc, 152, "NCA staff review the exact immutable schema, request targeted cloned corrections, decide non-filled explanations, inspect clean private files and approve/reject with a complete audit trail.")
replace_paragraph(doc, 153, "NCA Admin/Officer can identify overdue providers, create reminder drafts/queue records and view follow-up history; delivery is not falsely represented when no provider is configured.")
replace_paragraph(doc, 154, "Approved-only operational CSV/PDF, governed request artifacts and role-sanitized aggregate dashboard XLSX work with server-side permissions and audit records.")
replace_paragraph(doc, 161, "Required DBS05 topology KMZ and optional private Excel backups can be uploaded, quarantined, scanned, reviewed, downloaded by authorized roles and represented in metadata/audit evidence.")
replace_paragraph(doc, 167, "Publication acceptance is based on reliable seven-form structured collection, explicit Provider Data Entry-to-Approver and NCA correction workflows, immutable history, exact-version review, approved-only exports, governed data requests, aggregate dashboard controls, private scanned files, cookie/CSRF authentication, portal notifications, audit evidence and verified recovery. MFA, live email delivery, automated retention deletion, production credentials/targets, official Tower code and formal NCA/provider UAT remain explicit deferred or external gates; none is represented as complete without evidence.")

# Add the delivered-system alignment section in the same document styles.
doc.add_page_break()
doc.add_heading("20. Delivered System Alignment (August 2026)", level=1)
doc.add_paragraph("This section records the delivered behaviour that supersedes earlier planning assumptions. It does not alter historical submission/form records, and it does not claim exact source-workbook fidelity without the approved original regulatory forms.")

doc.add_heading("20.1 Roles and isolation", level=2)
for text in [
    "The five stored roles and their exclusions are enforced at both API and navigation layers.",
    "Every NCA Data Requester is a named individual with unique email, managed active division and required grade; request snapshots preserve name, email, division and grade.",
    "Provider tenancy applies to obligations, drafts, values, grids, files, approver queues, correction actions, history and notifications.",
]:
    bullet(doc, text)

doc.add_heading("20.2 Corrected immutable forms", level=2)
form_table = doc.add_table(rows=1, cols=4)
form_table.style = tables[5].style
set_table_header_repeat(form_table.rows[0])
for index, value in enumerate(("Family", "Version", "Frequency", "Publication note")):
    form_table.rows[0].cells[index].text = value
for values in [
    ("MNO-MONTHLY", "3.0", "Monthly", "Thirteen PRD sections, fixed dictionaries, traffic distinctions, mobile money and calculated totals."),
    ("DC-ISP06", "3.0", "Annual", "Supplier/package/PoP/backhaul/coordinate grids and IRU/GIX/VSAT conditions."),
    ("DC-ITC04", "3.0", "Annual", "Regional distribution, tenant/permit/addition/decommission/acquisition and coordinate structures."),
    ("DC-TB02", "3.0", "Annual", "Service selections, regional dealers, bouquet/tariff/channel and satellite grids, declaration."),
    ("TOWER-MAIN-ANNUAL", "2.0", "Annual", "Five regional grids and totals; form code remains visibly provisional."),
    ("DC-DBS05", "2.0", "Annual", "Fibre/capacity/coverage, Ghana regions and the only required topology KMZ."),
    ("DC-SUB03", "2.0", "Annual", "Separate global/Ghana metrics and fibre cuts; no topology KMZ."),
]:
    add_row(form_table, values)

doc.add_heading("20.3 Provider and regulatory state machines", level=2)
doc.add_paragraph("Provider stage: DRAFT → PENDING_APPROVAL → PROVIDER_CHANGES_REQUESTED → PROVIDER_RESUBMITTED → SUBMITTED.")
doc.add_paragraph("Regulatory stage: SUBMITTED → UNDER_REVIEW → CORRECTION_REQUESTED → RESUBMITTED → APPROVED or REJECTED.")
for text in [
    "Section saves are transactional authoritative snapshots with optimistic revision checks, exact scalar/grid targets, fixed-row preservation and omitted editable-repeatable-row deletion.",
    "Non-filled status clears a prior value and requires an explanation; conditional visibility reads the complete submission; allow-listed formulas and rules are recalculated on save/transitions.",
    "Provider and NCA correction instructions require reasons and section/field/grid-cell targets. Regulatory correction clones the official version and locks unaffected content.",
]:
    bullet(doc, text)

doc.add_heading("20.4 Data delivery, dashboard and private files", level=2)
for text in [
    "NCA Admin can download exactly the READY, unexpired artifact released to a requester; owning requester access is isolated and NCA Officer remains excluded from Data Requests.",
    "All authenticated roles may view/filter Industry Dashboard aggregates and generate aggregate XLSX. Source workbooks, provider rows and configuration controls remain NCA Admin/Officer only.",
    "Operational CSV/PDF uses only NCA-approved submission versions and one canonical scalar/grid row shape with status, explanation, review and compliance context.",
    "DBS05 KMZ and optional Excel backups remain outside public media, are SHA-256 hashed, quarantined and scan-gated, and use authenticated audited streaming downloads.",
]:
    bullet(doc, text)

doc.add_heading("20.5 Authentication, notifications and audit", level=2)
for text in [
    "Rotating/revoked JWT sessions are stored in Secure, HttpOnly, SameSite cookies; unsafe browser requests require CSRF. Failed-login counting, throttling, lockout and Admin-mediated temporary-password reset are implemented.",
    "MFA routes/claims are removed for this release and MFA remains deferred.",
    "Submission notifications are separate from Data Request notifications and are visible only to intended users; timeline audiences prevent internal NCA notes from leaking to providers.",
    "Audit events are immutable and hash linked; daily signed anchors and governance readiness evidence support tamper detection and release review.",
]:
    bullet(doc, text)

doc.add_heading("20.6 External release gates and genuinely deferred work", level=2)
gate_table = doc.add_table(rows=1, cols=3)
gate_table.style = tables[15].style
set_table_header_repeat(gate_table.rows[0])
for index, value in enumerate(("Item", "Current classification", "Required evidence")):
    gate_table.rows[0].cells[index].text = value
for values in [
    ("Managed TLS ingress and production configuration", "External release gate", "Approved domains, secrets, HTTPS origins and trusted proxy CIDRs."),
    ("Backup and recovery", "External release gate", "Immutable destination/credentials, approved RPO/RTO and successful isolated restore."),
    ("Retention disposition", "Disabled", "Approved record-class schedules, maker/checker approvals and legal-hold validation."),
    ("Tower official form code", "Provisional", "Approved NCA source-owner decision; replacement is a new family/version."),
    ("Microsoft Graph/shared mailbox", "Advisory/deferred", "Tenant credentials, approved mailbox/template governance and delivery-event UAT."),
    ("MFA", "Deferred", "Approved policy, method, recovery flow and security/UAT evidence."),
    ("Formal publication acceptance", "External release gate", "Named NCA/provider, browser, accessibility and spreadsheet/PDF sign-offs."),
]:
    add_row(gate_table, values)

doc.add_heading("20.7 Acceptance evidence", level=2)
doc.add_paragraph("Publication readiness may be claimed only when forward migrations, Django checks/tests, frontend lint/type/build, dependency audits, production configuration/build, private-file and security checks, full smoke test, audit verification, recovery evidence, rendered PRD inspection and named UAT are all attached to the release record. Missing external inputs remain visible blockers and are never replaced with invented credentials, schedules, targets or approvals.")

# Preserve header/footer structure while updating stale draft wording.
for section in doc.sections:
    for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
        for run in paragraph.runs:
            run.text = run.text.replace("Development-ready planning draft | May 28, 2026", "Publication remediation update | August 12, 2026")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
