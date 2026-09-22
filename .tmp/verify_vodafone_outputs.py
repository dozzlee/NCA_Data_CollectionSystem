import json
import math
from pathlib import Path

from openpyxl import load_workbook

root = Path(r"C:\Users\dozzl\OneDrive\Documents\testing and fixing")
output = root / "outputs" / "vodafone-monthly-2026"
manifest = json.loads((root / ".tmp" / "vodafone_values_manifest.json").read_text(encoding="utf-8"))
paths = {
    "september": output / "NCA_Monthly_Report_Vodafone_September_2026.xlsx",
    "half": output / "NCA_Monthly_Report_Vodafone_October_2026_Half_Filled.xlsx",
    "absurd": output / "NCA_Monthly_Report_Vodafone_October_2026_Absurd_Growth.xlsx",
}
books = {key: load_workbook(path, read_only=True, data_only=True) for key, path in paths.items()}


def same_number(actual, expected):
    return actual is not None and math.isclose(float(actual), float(expected), rel_tol=1e-10, abs_tol=1e-6)


for entry in manifest["entries"]:
    sheet = entry["sheet"]
    row = entry["row"]
    september_values = [books[key][sheet].cell(row, 46).value for key in books]
    if not all(same_number(value, entry["september"]) for value in september_values):
        raise AssertionError((entry["target_key"], september_values, entry["september"]))

half_count = sum(
    books["half"][entry["sheet"]].cell(entry["row"], 47).value is not None
    for entry in manifest["entries"]
)
full_count = sum(
    books["absurd"][entry["sheet"]].cell(entry["row"], 47).value is not None
    for entry in manifest["entries"]
)
september_october_count = sum(
    books["september"][entry["sheet"]].cell(entry["row"], 47).value is not None
    for entry in manifest["entries"]
)
if (half_count, full_count, september_october_count) != (141, 281, 0):
    raise AssertionError((half_count, full_count, september_october_count))

for book in books.values():
    book.close()

print(json.dumps({
    "september_identical_in_all_files": True,
    "september_indicator_count": 281,
    "half_october_indicator_count": half_count,
    "full_october_indicator_count": full_count,
    "october_in_september_only_file": september_october_count,
    "files": {key: {"path": str(path), "bytes": path.stat().st_size} for key, path in paths.items()},
}, indent=2))
