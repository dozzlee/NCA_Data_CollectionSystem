import json
import os
import re
import sys
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings.development")

import django

django.setup()

from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.forms_engine.models import FormField, FormGrid, FormTemplate
from apps.providers.models import ProviderProfile
from apps.submissions.models import (
    ExpectedSubmission,
    ProviderApprovalDecision,
    ReviewAction,
    Submission,
    SubmissionEvent,
    SubmissionValue,
)
from apps.submissions.readiness import refresh_submission_completion
from apps.users.models import User


SOURCE_WORKBOOK = r"C:\Users\dozzl\Downloads\NCA_Monthly_Report.xlsx"
SOURCE_VALUE_COLUMN = 45  # AS: latest populated historical column in the supplied workbook.
YEAR = 2026
MONTH = 9


def decimal_from_cell(raw):
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw))
    cleaned = re.sub(r"[^0-9.\-]", "", str(raw))
    if cleaned in {"", "-", ".", "-."}:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def quantized(value, places=0):
    quantum = Decimal("1") if places == 0 else Decimal("1." + ("0" * places))
    return value.quantize(quantum, rounding=ROUND_HALF_UP)


def generated_value(field, raw):
    base = decimal_from_cell(raw)
    if field.field_type == "percentage":
        if base is None:
            base = Decimal(55 + (field.id % 36))
        elif abs(base) <= 1:
            base *= 100
        return float(max(Decimal("1"), min(Decimal("99.5"), quantized(abs(base), 2))))
    if base is None:
        base = Decimal(((field.source_row or field.id) * 7919 + field.id * 37) % 900000 + 100)
    factor = Decimal("0.30") + Decimal(field.id % 7) / Decimal("100")
    scaled = abs(base) * factor
    if scaled == 0:
        scaled = Decimal((field.id % 997) + 1)
    places = 2 if scaled != scaled.to_integral_value() else 0
    return float(quantized(scaled, places)) if places else int(quantized(scaled, 0))


SPECIAL_SEPTEMBER = {
    "TOTAL_MOBILE_CELLULAR_VOICE_SUBSCRIPTIONS_PREPAID_POSTPAID": 8_750_000,
    "MOBILE_CELLULAR_VOICE_SUBSCRIPTIONS_PREPAID_ONLY": 8_640_000,
    "MOBILE_CELLULAR_VOICE_SUBSCRIPTIONS_POSTPAID_ONLY": 110_000,
    "TOTAL_MOBILE_CELLULAR_DATA_SUBSCRIPTIONS": 6_950_000,
    "TOTAL_MOBILE_CELLULAR_SUBSCRIPTION_PREPAID_DATA": 6_870_000,
    "TOTAL_MOBILE_CELLULAR_SUBSCRIPTION_POSTPAID_DATA": 80_000,
    "NEWLY_ACTIVATED_SIMS_FOR_THE_MONTH": 245_000,
    "NUMBER_OF_BLOCKED_MOBILE_SUBSCRIPTIONS": 12_500,
}


def october_value(value, field_type):
    number = Decimal(str(value))
    if field_type == "percentage":
        return float(min(Decimal("99.9"), quantized(number + Decimal("0.35"), 2)))
    updated = number * Decimal("1.025")
    places = 2 if number != number.to_integral_value() else 0
    return float(quantized(updated, places)) if places else int(quantized(updated, 0))


def db_value(value):
    number = Decimal(str(value))
    if number == number.to_integral_value():
        return str(int(number))
    return format(number.normalize(), "f")


workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
manifest = []

with connection.cursor() as cursor:
    cursor.execute("PRAGMA busy_timeout = 120000")

with transaction.atomic():
    provider = ProviderProfile.objects.select_related("organization").get(provider_code="VOD")
    form = FormTemplate.objects.get(form_code="MNO-MONTHLY", status="ACTIVE")
    expected = ExpectedSubmission.objects.select_for_update().get(
        provider=provider,
        form_template=form,
        period__year=YEAR,
        period__month=MONTH,
    )
    submission = Submission.objects.select_for_update().get(expected=expected, version=1)
    approver = User.objects.get(organization=provider.organization, role="PROVIDER_APPROVER", is_active=True)
    reviewer = User.objects.filter(role="NCA_ADMIN", is_active=True).order_by("id").first()
    if reviewer is None:
        reviewer = User.objects.filter(role="NCA_OFFICER", is_active=True).order_by("id").first()
    if reviewer is None:
        raise RuntimeError("No active NCA Admin or Officer is available to approve the submission.")

    SubmissionValue.objects.filter(submission=submission).delete()
    target_index = 0
    values_to_create = []

    fields = FormField.objects.filter(section__form_template=form).select_related("section").order_by(
        "section__sort_order", "sort_order", "id"
    )
    for field in fields:
        raw = workbook[field.source_sheet].cell(field.source_row, SOURCE_VALUE_COLUMN).value
        september = SPECIAL_SEPTEMBER.get(field.field_code, generated_value(field, raw))
        october = october_value(september, field.field_type)
        is_mobile_subscription = "MOBILE" in field.field_code and "SUBSCRIPTION" in field.field_code
        october_absurd = october * 10000 if is_mobile_subscription else october
        values_to_create.append(
            SubmissionValue(
                submission=submission,
                field=field,
                value=db_value(september),
                value_status="PROVIDED",
                value_source="MANUAL",
                source_reference="vodafone-september-2026-demo",
                updated_by=approver,
            )
        )
        manifest.append({
            "target_key": f"FIELD:{field.id}",
            "target_type": "FIELD",
            "section": field.section.section_code,
            "field_code": field.field_code,
            "label": field.label,
            "field_type": field.field_type,
            "sheet": field.source_sheet,
            "row": field.source_row,
            "september": september,
            "october": october,
            "october_absurd": october_absurd,
            "selected_half": target_index % 2 == 0,
            "is_mobile_subscription": is_mobile_subscription,
        })
        target_index += 1

    grids = FormGrid.objects.filter(section__form_template=form).select_related("section").prefetch_related(
        "columns", "fixed_rows"
    ).order_by("section__sort_order", "sort_order", "id")
    for grid in grids:
        columns = list(grid.columns.all().order_by("sort_order", "id"))
        rows = list(grid.fixed_rows.all().order_by("sort_order", "id"))
        for grid_row in rows:
            for column_index, column in enumerate(columns):
                source_row = grid_row.source_rows[column_index]
                raw = workbook[column.source_sheet or grid.source_sheet].cell(source_row, SOURCE_VALUE_COLUMN).value
                pseudo_field = type("GridField", (), {
                    "field_type": column.field_type,
                    "id": column.id + grid_row.id,
                    "source_row": source_row,
                })()
                september = generated_value(pseudo_field, raw)
                october = october_value(september, column.field_type)
                values_to_create.append(
                    SubmissionValue(
                        submission=submission,
                        grid=grid,
                        grid_row_id=str(grid_row.id),
                        grid_column=column,
                        value=db_value(september),
                        value_status="PROVIDED",
                        value_source="MANUAL",
                        source_reference="vodafone-september-2026-demo",
                        updated_by=approver,
                    )
                )
                manifest.append({
                    "target_key": f"GRID:{grid.id}:{grid_row.id}:{column.id}",
                    "target_type": "GRID_CELL",
                    "section": grid.section.section_code,
                    "field_code": f"{grid.grid_code}.{grid_row.row_label}.{column.column_code}",
                    "label": f"{grid.title} - {grid_row.row_label} - {column.label}",
                    "field_type": column.field_type,
                    "sheet": column.source_sheet or grid.source_sheet,
                    "row": source_row,
                    "september": september,
                    "october": october,
                    "october_absurd": october,
                    "selected_half": target_index % 2 == 0,
                    "is_mobile_subscription": False,
                })
                target_index += 1

    SubmissionValue.objects.bulk_create(values_to_create)
    readiness = refresh_submission_completion(submission)
    if readiness["completion_pct"] != 100.0 or readiness["missing_indicator_count"] != 0:
        raise RuntimeError(f"Submission did not reach full completion: {readiness}")

    now = timezone.now()
    previous_status = expected.workflow_status
    expected.workflow_status = "APPROVED"
    expected.provider_status = "CLOSED"
    expected.due_state = "CLOSED"
    expected.save(update_fields=["workflow_status", "provider_status", "due_state"])
    submission.regulatory_status = "APPROVED"
    submission.completion_pct = 100
    submission.submitted_by = approver
    submission.submitted_at = submission.submitted_at or now
    submission.reviewed_by = reviewer
    submission.reviewed_at = now
    submission.last_edited_by = approver
    submission.last_edited_at = now
    submission.revision += 1
    submission.save(update_fields=[
        "regulatory_status", "completion_pct", "submitted_by", "submitted_at", "reviewed_by",
        "reviewed_at", "last_edited_by", "last_edited_at", "revision",
    ])
    ProviderApprovalDecision.objects.update_or_create(
        submission=submission,
        defaults={
            "approver": approver,
            "attestation": True,
            "approval_note": "All September indicators reviewed and approved for official submission.",
            "change_summary": "",
            "approver_edited": True,
            "edit_batch_count": 0,
        },
    )
    if not ReviewAction.objects.filter(
        submission=submission, action="APPROVE", comment="Vodafone September 2026 fully populated demonstration submission."
    ).exists():
        ReviewAction.objects.create(
            submission=submission,
            action="APPROVE",
            target_type="SUBMISSION",
            target_id=str(submission.id),
            comment="Vodafone September 2026 fully populated demonstration submission.",
            is_provider_visible=True,
            created_by=reviewer,
        )
    if not SubmissionEvent.objects.filter(
        submission=submission,
        event_type="SUBMISSION_APPROVED",
        metadata__seed_key="vodafone-september-2026-full",
    ).exists():
        SubmissionEvent.objects.create(
            submission=submission,
            event_type="SUBMISSION_APPROVED",
            from_status=previous_status,
            to_status="APPROVED",
            message="NCA approved the fully completed Vodafone September 2026 submission.",
            audience="BOTH",
            metadata={"seed_key": "vodafone-september-2026-full", "indicator_count": len(manifest)},
            actor=reviewer,
        )

workbook.close()
result = {
    "submission_id": submission.id,
    "submission_reference": submission.submission_reference,
    "expected_submission_id": expected.id,
    "completion_pct": float(submission.completion_pct),
    "workflow_status": expected.workflow_status,
    "regulatory_status": submission.regulatory_status,
    "indicator_count": len(manifest),
    "half_filled_count": sum(1 for item in manifest if item["selected_half"]),
    "entries": manifest,
}
print("MANIFEST_START")
print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
print("MANIFEST_END")
