from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase
from openpyxl import Workbook
from io import BytesIO
from django.utils import timezone

from apps.users.models import User
from .models import (
    FormCodeCatalog, FormField, FormGapAssessment, FormHeading, FormRequirement, FormSection,
    FormTemplate, GridRow, ValidationRule, FormWorkbookImport,
)
from .workbook_import import PARSER_VERSION, _parse_workbook_streaming, parse_workbook, validate_schema


class PRDSection11FormTests(TestCase):
    def setUp(self):
        User.objects.create_user("prd-admin@nca.test", "password", name="Admin", role="NCA_ADMIN")
        User.objects.create_user("prd-officer@nca.test", "password", name="Officer", role="NCA_OFFICER")

    def test_loader_builds_exact_section_11_versions_and_publishes_provisional_tower_code(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "Product Requirements Document - Development Ready.docx"
            source.write_bytes(b"section-11-test-source")
            call_command("load_prd_section11", source=str(source), verbosity=0)
        active = FormTemplate.objects.filter(status="ACTIVE", approval_status="APPROVED")
        self.assertEqual(active.count(), 7)
        self.assertEqual(set(active.values_list("form_code", flat=True)), {"MNO-MONTHLY", "DC-TB02", "DC-ISP06", "DC-ITC04", "TOWER-MAIN-ANNUAL", "DC-DBS05", "DC-SUB03"})
        tower = FormTemplate.objects.get(form_code="TOWER-MAIN-ANNUAL", version="2.0")
        self.assertEqual(tower.status, "ACTIVE")
        self.assertEqual(tower.family.code_status, "PROVISIONAL")
        self.assertTrue(tower.gap_assessments.filter(requirement__requirement_type="SOURCE_DECISION", status="MATCHED").exists())
        mno = active.get(form_code="MNO-MONTHLY")
        brands = list(GridRow.objects.filter(grid__section__form_template=mno, grid__grid_code="phone_brand_counts").values_list("row_label", flat=True))
        self.assertIn("Samsung", brands)
        self.assertIn("Itel", brands)
        self.assertFalse(mno.sections.filter(grids__grid_code="industry_subscriptions_metrics").exists())
        self.assertEqual(mno.version, "3.0")
        self.assertEqual(active.get(form_code="DC-DBS05").frequency, "ANNUAL")
        self.assertTrue(active.get(form_code="DC-DBS05").kmz_requirements.exists())
        self.assertFalse(active.get(form_code="DC-SUB03").kmz_requirements.exists())
        for rule in ValidationRule.objects.filter(form_template__in=active):
            rule.full_clean()


class WorkbookFormImportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user("workbook-admin@nca.test", "password", name="Admin", role="NCA_ADMIN")
        self.client.force_authenticate(self.admin)
        definitions = {
            "NEW-QUARTERLY": ("New Quarterly Return", "QUARTERLY"),
            "EXISTING-MONTHLY": ("Existing Monthly Return", "MONTHLY"),
            "CUSTOM-LAYOUT": ("Custom Layout", "ANNUAL"),
            "BROKEN-WORKBOOK": ("Broken Workbook", "ANNUAL"),
            "LARGE-WORKBOOK": ("Large Workbook", "ANNUAL"),
        }
        for order, (code, (name, frequency)) in enumerate(definitions.items(), start=100):
            FormCodeCatalog.objects.update_or_create(code=code, defaults={
                "name": name, "sector": "TELECOM", "provider_category": "ISP",
                "frequency": frequency, "source_filename": f"{code}.xlsx",
                "is_active": True, "sort_order": order,
            })

    def workbook_file(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Subscriber Data"
        sheet.append(["NCA quarterly subscriber return"])
        sheet.append(["Industry Data", "Definitions", "Data Type", "Unit", "Mandatory", "Options", "June 2025"])
        sheet.append(["Subscriber indicators", "", "", "", "", "", "Historical heading"])
        sheet.append(["Total subscribers", "Number of active subscriptions", "Number", "Subscribers", "No", "", 999999])
        sheet.append(["Customer segment", "Select the customer category", "Select", "", "Yes", "Prepaid; Postpaid", "Prepaid"])
        out = BytesIO(); workbook.save(out)
        return SimpleUploadedFile("source.xlsx", out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @override_settings(MALWARE_SCANNER_REQUIRED=False)
    def test_workbook_import_generates_schema_only_and_confirms_draft(self):
        with TemporaryDirectory() as private_root:
            with override_settings(PRIVATE_UPLOAD_ROOT=private_root):
                response = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "NEW-QUARTERLY", "name": "New Quarterly Return", "version": "1.0",
                    "sector": "TELECOM", "provider_category": "ISP", "frequency": "QUARTERLY",
                    "file": self.workbook_file(),
                }, format="multipart")
                self.assertEqual(response.status_code, 201, response.data)
                item = FormWorkbookImport.objects.get(pk=response.data["id"])
                self.assertEqual(item.parse_status, "READY")
                self.assertNotIn("999999", str(item.detected_schema))
                self.assertEqual(item.parser_version, PARSER_VERSION)
                section = item.detected_schema["sections"][0]
                self.assertEqual(section["headings"][0]["title"], "Subscriber indicators")
                self.assertEqual(section["fields"][0]["help_text"], "Number of active subscriptions")
                self.assertFalse(section["fields"][0]["is_required"])
                self.assertTrue(section["fields"][1]["is_required"])
                confirmed = self.client.post(f"/api/v1/form-workbook-imports/{item.id}/confirm/", {}, format="json")
                self.assertEqual(confirmed.status_code, 201, confirmed.data)
                form = FormTemplate.objects.get(pk=confirmed.data["id"])
                self.assertEqual(form.form_code, "NEW-QUARTERLY")
                self.assertEqual(form.frequency, "QUARTERLY")
                self.assertEqual(form.mapping_basis, "SOURCE_FORM")
                heading = FormHeading.objects.get(section__form_template=form)
                self.assertEqual(heading.title, "Subscriber indicators")
                fields = list(FormField.objects.filter(section__form_template=form).order_by("sort_order"))
                self.assertEqual(fields[0].heading, heading)
                self.assertFalse(fields[0].is_required)
                self.assertTrue(fields[1].is_required)
                updated_heading = self.client.patch(
                    f"/api/v1/form-templates/{form.id}/sections/{heading.section_id}/headings/{heading.id}/",
                    {"title": "Subscriber measures"}, format="json",
                )
                self.assertEqual(updated_heading.status_code, 200, updated_heading.data)
                added_heading = self.client.post(
                    f"/api/v1/form-templates/{form.id}/sections/{heading.section_id}/headings/",
                    {"heading_code": "ADDITIONAL", "title": "Additional measures", "level": 2},
                    format="json",
                )
                self.assertEqual(added_heading.status_code, 201, added_heading.data)
                removed_heading = self.client.delete(
                    f"/api/v1/form-templates/{form.id}/sections/{heading.section_id}/headings/{added_heading.data['id']}/",
                )
                self.assertEqual(removed_heading.status_code, 204)

    @override_settings(MALWARE_SCANNER_REQUIRED=False)
    def test_workbook_import_creates_new_immutable_version_in_existing_family(self):
        with TemporaryDirectory() as private_root:
            with override_settings(PRIVATE_UPLOAD_ROOT=private_root):
                first = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "EXISTING-MONTHLY", "name": "Existing Monthly Return", "version": "1.0",
                    "sector": "TELECOM", "provider_category": "MNO", "frequency": "MONTHLY",
                    "file": self.workbook_file(),
                }, format="multipart")
                self.assertEqual(first.status_code, 201, first.data)
                first_draft = self.client.post(
                    f"/api/v1/form-workbook-imports/{first.data['id']}/confirm/", {}, format="json",
                )
                self.assertEqual(first_draft.status_code, 201, first_draft.data)

                replacement = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "EXISTING-MONTHLY", "name": "Existing Monthly Return", "version": "2.0",
                    "sector": "TELECOM", "provider_category": "MNO", "frequency": "MONTHLY",
                    "file": self.workbook_file(),
                }, format="multipart")
                self.assertEqual(replacement.status_code, 201, replacement.data)
                replacement_draft = self.client.post(
                    f"/api/v1/form-workbook-imports/{replacement.data['id']}/confirm/", {}, format="json",
                )
                self.assertEqual(replacement_draft.status_code, 201, replacement_draft.data)

                original = FormTemplate.objects.get(pk=first_draft.data["id"])
                updated = FormTemplate.objects.get(pk=replacement_draft.data["id"])
                self.assertEqual(original.family_id, updated.family_id)
                self.assertEqual(original.version, "1.0")
                self.assertEqual(updated.version, "2.0")
                self.assertEqual(original.sections.count(), 1)
                self.assertEqual(updated.sections.count(), 1)

    @override_settings(MALWARE_SCANNER_REQUIRED=False)
    def test_missing_mapping_is_blocking_until_reparsed_with_selected_columns(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Custom Layout"
        sheet.append(["Metric", "Meaning", "Kind"])
        sheet.append(["Customers", "Active customer count", "Number"])
        out = BytesIO(); workbook.save(out)
        with TemporaryDirectory() as private_root:
            with override_settings(PRIVATE_UPLOAD_ROOT=private_root):
                response = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "CUSTOM-LAYOUT", "name": "Custom Layout", "version": "1.0",
                    "sector": "TELECOM", "provider_category": "ISP", "frequency": "ANNUAL",
                    "file": SimpleUploadedFile("custom.xlsx", out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                }, format="multipart")
                self.assertEqual(response.status_code, 201, response.data)
                item_id = response.data["id"]
                code = "MISSING_DEFINITION_COLUMNS_1"
                self.assertIn(code, {warning["code"] for warning in response.data["warnings"]})
                patched = self.client.patch(
                    f"/api/v1/form-workbook-imports/{item_id}/",
                    {"mapping_decisions": {"resolved_warning_codes": [code]}}, format="json",
                )
                self.assertEqual(patched.status_code, 200, patched.data)
                denied = self.client.post(f"/api/v1/form-workbook-imports/{item_id}/confirm/", {}, format="json")
                self.assertEqual(denied.status_code, 409, denied.data)
                reparsed = self.client.post(
                    f"/api/v1/form-workbook-imports/{item_id}/reparse/",
                    {"column_mappings": {"Custom Layout": {
                        "header_row": 1, "indicator_column": 1,
                        "definition_column": 2, "data_type_column": 3,
                    }}}, format="json",
                )
                self.assertEqual(reparsed.status_code, 200, reparsed.data)
                self.assertEqual(reparsed.data["detected_schema"]["sections"][0]["fields"][0]["label"], "Customers")
                confirmed = self.client.post(f"/api/v1/form-workbook-imports/{item_id}/confirm/", {}, format="json")
                self.assertEqual(confirmed.status_code, 201, confirmed.data)

    @override_settings(MALWARE_SCANNER_REQUIRED=False)
    def test_failed_parse_returns_diagnostic_import_resource(self):
        with TemporaryDirectory() as private_root:
            with override_settings(PRIVATE_UPLOAD_ROOT=private_root):
                response = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "BROKEN-WORKBOOK", "name": "Broken Workbook", "version": "1.0",
                    "sector": "TELECOM", "provider_category": "ISP", "frequency": "ANNUAL",
                    "file": SimpleUploadedFile(
                        "broken.xlsx", b"not-an-excel-archive",
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                }, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["scan_status"], "CLEAN")
        self.assertEqual(response.data["parse_status"], "FAILED")
        self.assertIn("could not be read", response.data["scan_details"])
        self.assertEqual(response.data["detected_schema"], {})

    @override_settings(MALWARE_SCANNER_REQUIRED=False, FORM_WORKBOOK_ASYNC_THRESHOLD_BYTES=1)
    @patch("apps.forms_engine.views._schedule_workbook_import")
    def test_large_upload_returns_pending_preview_resource_and_schedules_processing(self, schedule):
        with TemporaryDirectory() as private_root:
            with override_settings(PRIVATE_UPLOAD_ROOT=private_root):
                response = self.client.post("/api/v1/form-workbook-imports/", {
                    "form_code": "LARGE-WORKBOOK", "name": "Large Workbook", "version": "1.0",
                    "sector": "TELECOM", "provider_category": "ISP", "frequency": "ANNUAL",
                    "file": self.workbook_file(),
                }, format="multipart")
        self.assertEqual(response.status_code, 202, response.data)
        self.assertEqual(response.data["scan_status"], "PENDING")
        self.assertEqual(response.data["parse_status"], "PENDING")
        schedule.assert_called_once_with(response.data["id"])


class WorksheetWorkbookGroupingTests(TestCase):
    def test_device_brand_pairs_generate_one_two_column_fixed_grid(self):
        brands = [
            "Samsung", "Apple", "Huawei", "Honor", "Nokia", "Xiaomi", "OPPO", "LG",
            "Vivo", "Lenovo", "Tecno", "Infinix", "Google", "Motorola", "Sony", "Realme",
            "Hisense", "TCL", "HTC", "OnePlus", "BLU", "Itel", "Others",
        ]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Device Brands"
        sheet.append(["Indicator", "Definition", "Data Type"])
        sheet.append(["Smart/Feature phone brands", "", ""])
        for brand in brands:
            sheet.append([brand, "Smart Phones =", "Number"])
            sheet.append(["", "Feature and Basic Phones =", "Number"])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "device-brands.xlsx"
            workbook.save(path)
            normal, warnings = parse_workbook(path)
            streamed, stream_warnings = _parse_workbook_streaming(path)

        section = normal["sections"][0]
        self.assertEqual(section["fields"], [])
        self.assertEqual(section["counts"], {"scalar_field_count": 0, "table_count": 1, "grid_input_count": 46})
        self.assertEqual(len(section["grids"]), 1)
        grid = section["grids"][0]
        self.assertEqual(grid["fixed_rows"], brands)
        self.assertEqual([column["label"] for column in grid["columns"]], [
            "Smart Phones", "Feature and Basic Phones",
        ])
        self.assertEqual(warnings, stream_warnings)
        self.assertEqual(normal["sections"], streamed["sections"])

    def test_incomplete_device_brand_pair_blocks_preview(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Device Brands"
        sheet.append(["Indicator", "Definition", "Data Type"])
        sheet.append(["Device brands", "", ""])
        sheet.append(["Samsung", "Smart Phones =", "Number"])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "broken-device-brands.xlsx"
            workbook.save(path)
            schema, warnings = parse_workbook(path)
        self.assertEqual(schema["sections"][0]["grids"], [])
        self.assertIn("INCOMPLETE_DEVICE_BRAND_PAIR_1_3", {warning["code"] for warning in warnings})
        self.assertIn("BLOCKING", {warning["severity"] for warning in warnings})

    def test_hidden_master_rows_are_excluded_from_each_worksheet_section(self):
        workbook = Workbook()
        voice = workbook.active
        voice.title = "Voice"
        voice.append(["Indicator", "Definition", "Data Type"])
        voice.append(["Voice minutes", "Outgoing voice traffic", "Number"])
        voice.append(["Data subscribers", "Active data subscriptions", "Number"])
        voice.append(["Customer complaints", "Complaints received", "Number"])
        voice.row_dimensions[3].hidden = True
        voice.row_dimensions[4].hidden = True

        data = workbook.create_sheet("Data and Support")
        data.append(["Indicator", "Definition", "Data Type"])
        data.append(["Voice minutes", "Outgoing voice traffic", "Number"])
        data.append(["Data subscribers", "Active data subscriptions", "Number"])
        data.append(["Customer complaints", "Complaints received", "Number"])
        data.row_dimensions[2].hidden = True

        with TemporaryDirectory() as directory:
            path = Path(directory) / "hidden-master-rows.xlsx"
            workbook.save(path)
            normal, _ = parse_workbook(path)
            streamed, _ = _parse_workbook_streaming(path)

        self.assertEqual(normal["grouping"]["isolation"], "worksheet-isolated")
        self.assertEqual(
            [[field["label"] for field in section["fields"]] for section in normal["sections"]],
            [["Voice minutes"], ["Data subscribers", "Customer complaints"]],
        )
        self.assertEqual(
            [section["row_visibility"]["excluded_hidden_row_count"] for section in normal["sections"]],
            [2, 1],
        )
        self.assertEqual(streamed["sections"], normal["sections"])

    def test_current_schema_rejects_cross_worksheet_children(self):
        schema = {
            "parser_version": PARSER_VERSION,
            "sections": [{
                "section_code": "VOICE",
                "title": "Voice",
                "source": {"sheet": "Voice", "sheet_index": 1},
                "column_mapping": {"indicator_column": 1, "definition_column": 2},
                "headings": [],
                "fields": [{
                    "field_code": "DATA_SUBSCRIBERS",
                    "label": "Data subscribers",
                    "field_type": "number",
                    "source": {"sheet": "Data", "row": 2},
                }],
                "grids": [],
            }],
        }
        with self.assertRaisesRegex(ValueError, "not section worksheet"):
            validate_schema(schema)

    def test_visible_worksheets_become_sections_in_tab_order(self):
        workbook = Workbook()
        voice = workbook.active
        voice.title = "Mobile Voice"
        voice.append(["Contact", "compliance@example.test"])
        voice.append(["Indicator", "Definition", "Data Type", "Required", "June 2025"])
        voice.append(["Mobile services", "", "", "", "Historical header"])
        voice.append(["Voice traffic", "", "", "", "Historical header"])
        voice.append(["Mobile voice minutes", "Total outgoing mobile minutes", "Number", "No", 12345])
        data = workbook.create_sheet("Mobile Data")
        data.append(["Industry Data", "Definitions", "Data Type", "Mandatory", "March 2025"])
        data.append(["Mobile voice minutes", "A deliberately repeated indicator", "Numeric", "Yes", 999])
        empty = workbook.create_sheet("Supporting Notes")
        hidden = workbook.create_sheet("Internal Lookup")
        hidden.sheet_state = "hidden"
        hidden.append(["Must not become a form field", 42])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "worksheet-tabs.xlsx"
            workbook.save(path)
            schema, warnings = parse_workbook(path)
        self.assertEqual(schema["parser_version"], PARSER_VERSION)
        self.assertEqual(schema["grouping"]["strategy"], "worksheet-tabs")
        self.assertEqual([section["title"] for section in schema["sections"]], [
            "Mobile Voice", "Mobile Data", "Supporting Notes",
        ])
        self.assertEqual([item["level"] for item in schema["sections"][0]["headings"]], [1, 2])
        self.assertEqual(schema["sections"][0]["fields"][0]["source"]["heading"], "Voice traffic")
        self.assertEqual(schema["sections"][0]["fields"][0]["help_text"], "Total outgoing mobile minutes")
        self.assertFalse(schema["sections"][0]["fields"][0]["is_required"])
        self.assertEqual(schema["sections"][0]["fields"][0]["parser_version"], PARSER_VERSION)
        self.assertEqual(schema["sections"][1]["fields"][0]["label"], "Mobile voice minutes")
        self.assertTrue(schema["sections"][1]["fields"][0]["is_required"])
        self.assertEqual(schema["sections"][2]["fields"], [])
        self.assertIn("EMPTY_WORKSHEET_SECTION_3", {warning["code"] for warning in warnings})
        self.assertIn("MISSING_DEFINITION_COLUMNS_3", {warning["code"] for warning in warnings})
        self.assertNotIn("Internal Lookup", str(schema))
        self.assertNotIn("12345", str(schema))
        self.assertNotIn("999", str(schema))

    def test_large_worksheet_remains_one_section_and_codes_are_deterministic(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Customer-Data"
        sheet.append(["Indicator", "Definition"])
        for index in range(1, 27):
            sheet.append([f"Customer account {index}", f"Definition {index}"])
        second_sheet = workbook.create_sheet("Customer Data")
        second_sheet.append(["Indicator", "Definition"])
        second_sheet.append(["Customer reference", "Provider reference"])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "large-worksheet.xlsx"; workbook.save(path)
            first, _ = parse_workbook(path); second, _ = parse_workbook(path)
        self.assertEqual([section["section_code"] for section in first["sections"]], ["CUSTOMER_DATA", "CUSTOMER_DATA_2"])
        self.assertEqual(len(first["sections"][0]["fields"]), 26)
        self.assertEqual(first, second)

    def test_data_entry_heading_creates_grid_with_only_blank_columns_editable(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "ISP Locations"
        sheet.append(["Industry Data", "Definitions", "Data Type", "User Input"])
        sheet.append(["8.3 Regional PoPs (Data Entry)", None, None, None])
        sheet.append(["No.", "Region", "Number of PoPs", "Notes"])
        sheet.append([1, "Ahafo", None, None])
        sheet.append([2, "Ashanti", None, None])
        sheet.append([None, None, None, None])
        sheet.append(["Total providers", "Count of providers", "Number", None])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "data-entry-table.xlsx"; workbook.save(path)
            normal, _ = parse_workbook(path)
            streamed, _ = _parse_workbook_streaming(path)
        grid = normal["sections"][0]["grids"][0]
        self.assertEqual(grid["title"], "8.3 Regional PoPs")
        self.assertEqual(grid["fixed_rows"], ["1 | Ahafo", "2 | Ashanti"])
        self.assertEqual([column["label"] for column in grid["columns"]], ["Number of PoPs", "Notes"])
        self.assertNotIn("No.", [column["label"] for column in grid["columns"]])
        self.assertNotIn("Region", [column["label"] for column in grid["columns"]])
        self.assertEqual(streamed["sections"], normal["sections"])

    def test_metadata_types_options_units_and_unknown_type_warning(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Traffic"
        sheet.append(["Indicator", "Definitions", "Data Type", "Unit", "Required", "Options"])
        sheet.append(["Traffic measures", "", "", "", "", ""])
        sheet.append(["Domestic traffic", "Domestic traffic carried", "Number", "Minutes", "true", ""])
        sheet.append(["Traffic category", "Select category", "Text", "", "", "Voice|Data"])
        sheet.append(["Special measure", "Unrecognized type example", "Bespoke", "", "", ""])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "metadata.xlsx"; workbook.save(path)
            schema, warnings = parse_workbook(path)
        fields = schema["sections"][0]["fields"]
        self.assertEqual(fields[0]["field_type"], "number")
        self.assertEqual(fields[0]["unit"], "Minutes")
        self.assertTrue(fields[0]["is_required"])
        self.assertEqual(fields[1]["field_type"], "select")
        self.assertEqual(fields[1]["options"], ["Voice", "Data"])
        self.assertEqual(fields[2]["field_type"], "text")
        self.assertIn("UNKNOWN_DATA_TYPE_1_5", {warning["code"] for warning in warnings})

    def test_streaming_parser_matches_normal_parser_sections_and_warnings(self):
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Subscriber Data"
        sheet.append(["Industry Data", "Definition", "Data Type", "Required"])
        sheet.append(["Subscriber heading", "", "", ""])
        sheet.append(["Prepaid customers", "Active prepaid subscribers", "Number", "Yes"])
        network = workbook.create_sheet("Network")
        network.append(["Indicator", "Definitions", "Data Type"])
        network.append(["Sites on air", "Operational radio sites", "Count"])
        hidden = workbook.create_sheet("Internal"); hidden.sheet_state = "hidden"; hidden.append(["Secret", 1])
        with TemporaryDirectory() as directory:
            path = Path(directory) / "streaming.xlsx"; workbook.save(path)
            normal, normal_warnings = parse_workbook(path)
            streamed, streamed_warnings = _parse_workbook_streaming(path)
        self.assertEqual(streamed["grouping"]["engine"], "streaming")
        self.assertEqual(streamed["sections"], normal["sections"])
        self.assertEqual(streamed_warnings, normal_warnings)
        self.assertNotIn("Internal", str(streamed))


class CustomFormPublicationTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user("custom-admin@nca.test", "password", name="Admin", role="NCA_ADMIN")
        self.client.force_authenticate(self.admin)
        FormCodeCatalog.objects.update_or_create(code="NCA-CUSTOM-01", defaults={
            "name": "Custom Return", "sector": "TELECOM", "provider_category": "MNO",
            "frequency": "MONTHLY", "source_filename": "NCA-CUSTOM-01.xlsx", "is_active": True,
        })

    def _make_publishable(self, form):
        form.mapping_complete = True
        form.source_reference = "NCA custom form approval record"
        form.source_sha256 = "a" * 64
        form.save(update_fields=["mapping_complete", "source_reference", "source_sha256"])
        section = FormSection.objects.create(form_template=form, section_code="CUSTOM_SECTION", title="Custom Section")
        FormField.objects.create(section=section, field_code="CUSTOM_FIELD", label="Custom Field", field_type="text", is_required=True)

    def test_new_manual_family_is_custom_and_does_not_require_section_11_match(self):
        created = self.client.post("/api/v1/form-templates/", {
            "form_code": "NCA-CUSTOM-01", "name": "Tampered Name", "version": "1.0",
            "sector": "BROADCASTING", "provider_category": "PAY_TV", "frequency": "ANNUAL",
            "effective_from": timezone.localdate(),
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        form = FormTemplate.objects.get(pk=created.data["id"])
        self.assertEqual((form.name, form.sector, form.provider_category, form.frequency), (
            "Custom Return", "TELECOM", "MNO", "ANNUAL",
        ))
        self.assertEqual(form.mapping_basis, "CUSTOM")
        self._make_publishable(form)
        FormRequirement.objects.create(
            family=form.family, requirement_key="PRD_ONLY_TEST", requirement_type="SECTION",
            label="PRD-only missing section", description="Not applicable to a custom family.",
            severity="BLOCKER", criteria={"section_codes": ["NOT_PRESENT"]},
        )
        checks = self.client.get(f"/api/v1/form-templates/{form.id}/publication-checks/")
        self.assertEqual(checks.status_code, 200, checks.data)
        self.assertFalse(checks.data["section_11_applicable"])
        self.assertNotIn("section_11_gaps_clear", checks.data["checks"])
        approved = self.client.post(f"/api/v1/form-templates/{form.id}/approve/", {}, format="json")
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertEqual(approved.data["approval_status"], "APPROVED")

    def test_catalog_returns_code_name_and_server_calculated_next_version(self):
        response = self.client.get("/api/v1/form-code-catalog/")
        self.assertEqual(response.status_code, 200, response.data)
        item = next(row for row in response.data["results"] if row["code"] == "NCA-CUSTOM-01")
        self.assertEqual(item["name"], "Custom Return")
        self.assertEqual(item["next_version"], "1.0")

    def test_prd_basis_still_blocks_missing_section_11_requirement(self):
        form = FormTemplate.objects.create(
            form_code="PRD-GATED", name="PRD Gated", version="1.0", sector="TELECOM",
            provider_category="MNO", frequency="MONTHLY", effective_from=timezone.localdate(),
            mapping_basis="PRD_SECTION_11", prepared_by=self.admin,
        )
        self._make_publishable(form)
        FormRequirement.objects.create(
            family=form.family, requirement_key="REQUIRED_PRD_SECTION", requirement_type="SECTION",
            label="Required PRD section", description="A required Section 11 structure.",
            severity="BLOCKER", criteria={"section_codes": ["PRD_REQUIRED"]},
        )
        denied = self.client.post(f"/api/v1/form-templates/{form.id}/approve/", {}, format="json")
        self.assertEqual(denied.status_code, 409, denied.data)
        self.assertIn("Section 11", denied.data["detail"])
