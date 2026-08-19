import re
import posixpath
import os
import re
import zipfile
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from django.conf import settings
from django.db import close_old_connections, transaction
from openpyxl import load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .models import (
    FormFamily, FormTemplate, FormSection, FormHeading, FormField, FormGrid,
    GridColumn, GridRow, SelectOption, FormWorkbookImport,
)
from apps.uploads.scanner import scan_path
PARSER_VERSION = "xlsx-worksheet-v6-visible-rows-matrix-grids"

def parse_document_source(path):
    """Create a schema-only draft from a PDF or Word source document.

    These formats do not carry a reliable spreadsheet schema, so text lines are
    presented as optional indicators for NCA review rather than being treated as
    provider answers.
    """
    extension = os.path.splitext(path)[1].lower()
    text = ""
    if extension == ".docx":
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        text = "\n".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    else:
        raw = open(path, "rb").read().decode("latin1", errors="ignore")
        text = "\n".join(re.findall(r"\(([^()]*)\)", raw))
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if len(line) >= 3]
    fields = []
    for index, line in enumerate(dict.fromkeys(lines), start=1):
        code = re.sub(r"[^A-Za-z0-9]+", "_", line).strip("_").lower()[:80] or f"indicator_{index}"
        fields.append({"field_code": f"DOC_{index}_{code}", "label": line[:255], "field_type": "text", "is_required": False, "help_text": "Imported from source document; verify before publishing.", "source": {"sheet": "Document", "row": index}, "source_order": index})
    return {"parser_version": "document-text-v1", "grouping": {"strategy": "document-text"}, "sections": [{"section_code": "DOCUMENT_CONTENT", "title": "Document content", "instructions": "Review and refine imported indicators before publishing.", "fields": fields, "grids": [], "headings": []}]}, ["PDF/Word sources were converted to optional text indicators; verify the generated structure before confirmation."]
FIXED_ROW_HEADERS = {"region", "category", "country", "operator", "brand", "service", "item", "location"}
STREAMING_THRESHOLD_BYTES = 5 * 1024 * 1024
MAX_SOURCE_ROWS = 2000
MAX_SOURCE_COLUMNS = 100
SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
OFFICE_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"

INDICATOR_HEADERS = {"indicator", "indicatorname", "industrydata", "industryindicator"}
DEFINITION_HEADERS = {"definition", "definitions", "description", "helptext"}
DATA_TYPE_HEADERS = {"datatype", "type", "valuetype"}
UNIT_HEADERS = {"unit", "units", "measurementunit"}
REQUIRED_HEADERS = {"required", "isrequired", "mandatory"}
OPTIONS_HEADERS = {"options", "allowedvalues", "choices"}
TRUE_VALUES = {"1", "true", "yes", "y", "required", "mandatory"}


def normalize_code(value, *, fallback="ITEM", max_length=100):
    code = re.sub(r"[^A-Z0-9]+", "_", str(value or "").strip().upper()).strip("_") or fallback
    return code[:max_length]


def normalize_form_code(value):
    code = re.sub(r"[^A-Z0-9-]+", "-", str(value or "").strip().upper()).strip("-")
    if not code or not re.fullmatch(r"[A-Z0-9]+(?:-[A-Z0-9]+)*", code):
        raise ValueError("Form code must contain uppercase letters, numbers and hyphens only.")
    return code


def _field_type(cell):
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (date, datetime)):
        return "date"
    if isinstance(value, (int, float)):
        number_format = str(cell.number_format or "")
        return "percentage" if "%" in number_format else "number"
    return "text"


def _unit(label):
    match = re.search(r"\(([^()]{1,30})\)\s*$", str(label or ""))
    return match.group(1).strip() if match else ""


def _unique_code(label, used, prefix, *, max_length=100):
    base = normalize_code(label, fallback=prefix, max_length=max_length)
    candidate, index = base, 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = f"{base[:max_length - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _is_heading_row(sheet, row, populated, covered):
    if len(populated) != 1:
        return False
    col, cell = populated[0]
    if not isinstance(cell.value, str) or str(cell.value).startswith("="):
        return False
    merged = any(cell.coordinate in merged_range for merged_range in sheet.merged_cells.ranges)
    styled = bool(cell.font.bold and (cell.font.sz or 0) >= 11)
    filled = bool(cell.fill and cell.fill.fill_type)
    return merged or styled or filled


def _source_metadata(sheet, row, heading):
    return {"sheet": sheet.title, "row": row, "heading": heading or ""}


def _nearest_heading_above(sheet, row):
    for candidate_row in range(row - 1, max(0, row - 12), -1):
        populated = [
            (col, sheet.cell(candidate_row, col))
            for col in range(1, min(sheet.max_column, 100) + 1)
            if sheet.cell(candidate_row, col).value not in (None, "")
        ]
        if _is_heading_row(sheet, candidate_row, populated, set()):
            return str(populated[0][1].value).strip()[:255]
    return ""


def _zip_target(base_path, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_path), target))


def _relationship_map(archive, source_path):
    rels_path = posixpath.join(
        posixpath.dirname(source_path),
        "_rels",
        f"{posixpath.basename(source_path)}.rels",
    )
    if rels_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rels_path))
    return {
        item.attrib["Id"]: _zip_target(source_path, item.attrib["Target"])
        for item in root.findall(f"{PACKAGE_REL_NS}Relationship")
        if item.attrib.get("Id") and item.attrib.get("Target")
    }


def _shared_strings(archive):
    path = "xl/sharedStrings.xml"
    if path not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(path))
    return [
        "".join(node.text or "" for node in item.iter(f"{SHEET_NS}t"))
        for item in root.findall(f"{SHEET_NS}si")
    ]


def _style_metadata(archive):
    path = "xl/styles.xml"
    if path not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(path))
    fonts = []
    fonts_root = root.find(f"{SHEET_NS}fonts")
    for font in list(fonts_root or []):
        bold_node = font.find(f"{SHEET_NS}b")
        size_node = font.find(f"{SHEET_NS}sz")
        bold = bold_node is not None and bold_node.attrib.get("val", "1") not in {"0", "false", "False"}
        try:
            size = float(size_node.attrib.get("val", 0)) if size_node is not None else 0
        except (TypeError, ValueError):
            size = 0
        fonts.append({"bold": bold, "size": size})

    number_formats = dict(BUILTIN_FORMATS)
    formats_root = root.find(f"{SHEET_NS}numFmts")
    for item in list(formats_root or []):
        try:
            number_formats[int(item.attrib["numFmtId"])] = item.attrib.get("formatCode", "")
        except (KeyError, TypeError, ValueError):
            continue

    styles = []
    cell_formats = root.find(f"{SHEET_NS}cellXfs")
    for item in list(cell_formats or []):
        try:
            font_id = int(item.attrib.get("fontId", 0))
            fill_id = int(item.attrib.get("fillId", 0))
            format_id = int(item.attrib.get("numFmtId", 0))
        except (TypeError, ValueError):
            font_id, fill_id, format_id = 0, 0, 0
        font = fonts[font_id] if 0 <= font_id < len(fonts) else {"bold": False, "size": 0}
        number_format = number_formats.get(format_id, "")
        styles.append({
            "heading": bool((font["bold"] and font["size"] >= 11) or fill_id > 1),
            "number_format": number_format,
            "is_date": bool(number_format and is_date_format(number_format)),
            "is_percentage": "%" in number_format,
        })
    return styles


def _stream_cell_value(cell, shared_strings, styles):
    try:
        style_index = int(cell.attrib.get("s", 0) or 0)
    except (TypeError, ValueError):
        style_index = 0
    style = styles[style_index] if 0 <= style_index < len(styles) else {
        "heading": False, "number_format": "", "is_date": False, "is_percentage": False,
    }
    formula_node = cell.find(f"{SHEET_NS}f")
    formula = (formula_node.text or "").strip() if formula_node is not None else ""
    value_node = cell.find(f"{SHEET_NS}v")
    raw_value = value_node.text if value_node is not None else None
    cell_type = cell.attrib.get("t", "n")
    if cell_type == "inlineStr":
        value = "".join(node.text or "" for node in cell.iter(f"{SHEET_NS}t"))
    elif cell_type == "s" and raw_value not in (None, ""):
        try:
            value = shared_strings[int(raw_value)]
        except (IndexError, TypeError, ValueError):
            value = raw_value
    elif cell_type == "b" and raw_value not in (None, ""):
        value = raw_value == "1"
    elif cell_type in {"str", "e"}:
        value = raw_value
    elif raw_value not in (None, ""):
        try:
            numeric = float(raw_value)
            value = int(numeric) if numeric.is_integer() else numeric
        except (TypeError, ValueError):
            value = raw_value
    else:
        value = None
    return {
        "value": value,
        "formula": formula,
        "style": style,
    }


def _stream_field_type(cell_data):
    if cell_data.get("formula"):
        return "formula"
    value = cell_data.get("value")
    style = cell_data.get("style", {})
    if isinstance(value, bool):
        return "boolean"
    if style.get("is_date"):
        return "date"
    if isinstance(value, (int, float)):
        return "percentage" if style.get("is_percentage") else "number"
    return "text"


def _metadata_key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _metadata_field_type(value):
    key = _metadata_key(value)
    if not key:
        return "text", False
    exact = {
        "text": "text", "string": "text", "alphanumeric": "text",
        "number": "number", "numbers": "number", "numeric": "number", "integer": "number",
        "decimal": "number", "float": "number", "count": "number", "volume": "number",
        "currency": "currency", "money": "currency", "ghs": "currency", "usd": "currency", "cedi": "currency",
        "percentage": "percentage", "percent": "percentage", "rate": "percentage",
        "date": "date", "datetime": "date",
        "boolean": "boolean", "yesno": "boolean",
        "select": "select", "dropdown": "select",
        "multiselect": "multiselect", "multiplechoice": "multiselect",
        "textarea": "textarea", "longtext": "textarea",
        "coordinate": "coordinate", "coordinates": "coordinate", "latlong": "coordinate",
        "declaration": "declaration", "attestation": "declaration",
        "formula": "formula", "calculated": "formula",
    }
    return exact.get(key, "text"), key not in exact


def _split_options(value):
    return [item.strip() for item in re.split(r"[\n,;|]+", str(value or "")) if item.strip()]


def _column_for(headers, aliases):
    return next((column for column, value in headers.items() if _metadata_key(value) in aliases), None)


def _detect_column_mapping(rows, override=None):
    override = override if isinstance(override, dict) else {}
    if override.get("indicator_column") and override.get("definition_column"):
        return {
            "detected": False,
            "header_row": int(override.get("header_row") or 1),
            "indicator_column": int(override["indicator_column"]),
            "definition_column": int(override["definition_column"]),
            "data_type_column": int(override["data_type_column"]) if override.get("data_type_column") else None,
            "unit_column": int(override["unit_column"]) if override.get("unit_column") else None,
            "required_column": int(override["required_column"]) if override.get("required_column") else None,
            "options_column": int(override["options_column"]) if override.get("options_column") else None,
        }
    for row_number in sorted(rows):
        headers = {column: cell.get("value") for column, cell in rows[row_number].items()}
        indicator_column = _column_for(headers, INDICATOR_HEADERS)
        definition_column = _column_for(headers, DEFINITION_HEADERS)
        if indicator_column and definition_column:
            return {
                "detected": True,
                "header_row": row_number,
                "indicator_column": indicator_column,
                "definition_column": definition_column,
                "data_type_column": _column_for(headers, DATA_TYPE_HEADERS),
                "unit_column": _column_for(headers, UNIT_HEADERS),
                "required_column": _column_for(headers, REQUIRED_HEADERS),
                "options_column": _column_for(headers, OPTIONS_HEADERS),
            }
    return None


def _column_candidates(rows):
    candidates = []
    for row_number in sorted(rows):
        labels = [
            {"column": column, "label": str(cell.get("value", "")).strip()[:100]}
            for column, cell in sorted(rows[row_number].items())
            if isinstance(cell.get("value"), str) and str(cell.get("value")).strip()
            and not str(cell.get("value")).strip().startswith("=")
        ]
        if len(labels) >= 2:
            candidates.append({"row": row_number, "columns": labels})
        if len(candidates) == 5:
            break
    return candidates


def _definition_content(rows, sheet_title, sheet_index, warnings, override=None):
    mapping = _detect_column_mapping(rows, override)
    if not mapping:
        warnings.append({
            "code": f"MISSING_DEFINITION_COLUMNS_{sheet_index}",
            "severity": "BLOCKING",
            "message": f'{sheet_title} does not contain recognizable Indicator and Definition columns. Select the columns before confirming.',
        })
        return [], [], [], {
            "detected": False, "header_row": None, "indicator_column": None,
            "definition_column": None, "data_type_column": None, "unit_column": None,
            "required_column": None, "options_column": None,
            "candidates": _column_candidates(rows),
        }

    headings = []
    fields = []
    grids = []
    used_heading_codes = set()
    used_field_codes = set()
    current_heading_code = ""
    current_heading_title = ""
    last_kind = ""
    current_level = 0

    # Some NCA workbooks represent a two-dimensional matrix as consecutive
    # metadata rows. Device brands are encoded as a brand row labelled
    # "Smart Phones", followed by a blank-indicator continuation row labelled
    # "Feature and Basic Phones". Preserve that structure as one fixed grid
    # instead of losing the continuation row or creating duplicate scalars.
    ordered_rows = sorted(row for row in rows if row > mapping["header_row"])
    indicator_column = mapping["indicator_column"]
    definition_column = mapping["definition_column"]
    heading_text = " ".join(
        str(rows[row].get(indicator_column, {}).get("value") or "")
        for row in ordered_rows
        if not str(rows[row].get(definition_column, {}).get("value") or "").strip()
    )
    device_brand_context = "brand" in _metadata_key(f"{sheet_title} {heading_text}")
    consumed_matrix_rows = set()
    brand_pairs = []
    if device_brand_context:
        for position, row_number in enumerate(ordered_rows):
            row = rows[row_number]
            brand = str(row.get(indicator_column, {}).get("value") or "").strip()
            category = _metadata_key(row.get(definition_column, {}).get("value") or "")
            if not brand or category not in {"smartphone", "smartphones"}:
                continue
            next_row_number = ordered_rows[position + 1] if position + 1 < len(ordered_rows) else None
            next_row = rows.get(next_row_number, {}) if next_row_number else {}
            continuation_brand = str(next_row.get(indicator_column, {}).get("value") or "").strip()
            continuation_category = _metadata_key(next_row.get(definition_column, {}).get("value") or "")
            if continuation_brand or continuation_category not in {
                "featureandbasicphone", "featureandbasicphones", "featurebasicphone", "featurebasicphones",
            }:
                warnings.append({
                    "code": f"INCOMPLETE_DEVICE_BRAND_PAIR_{sheet_index}_{row_number}",
                    "severity": "BLOCKING",
                    "message": f'{sheet_title} row {row_number} defines Smart Phones for "{brand}" without the required Feature and Basic Phones continuation row.',
                })
                continue
            brand_pairs.append({
                "label": brand[:255], "smart_row": row_number,
                "feature_row": next_row_number,
            })
            consumed_matrix_rows.update({row_number, next_row_number})

    if brand_pairs:
        grid_source_row = min(pair["smart_row"] for pair in brand_pairs)
        grids.append({
            "source_order": grid_source_row,
            "grid_code": "PHONE_COUNTS_BY_BRAND",
            "title": "Smart and Feature/Basic Phone Counts by Brand",
            "row_mode": "FIXED",
            "min_rows": len(brand_pairs),
            "instructions": "Enter Smart Phone and Feature/Basic Phone counts independently for every brand.",
            "columns": [
                {
                    "column_code": "SMART_PHONES", "label": "Smart Phones",
                    "field_type": "number", "unit": "count", "is_required": False,
                    "source": {"sheet": sheet_title, "row": grid_source_row},
                },
                {
                    "column_code": "FEATURE_BASIC_PHONES", "label": "Feature and Basic Phones",
                    "field_type": "number", "unit": "count", "is_required": False,
                    "source": {"sheet": sheet_title, "row": brand_pairs[0]["feature_row"]},
                },
            ],
            "fixed_rows": [pair["label"] for pair in brand_pairs],
            "fixed_row_sources": [
                {
                    "label": pair["label"],
                    "source": {"sheet": sheet_title, "rows": [pair["smart_row"], pair["feature_row"]]},
                }
                for pair in brand_pairs
            ],
            "source": {"sheet": sheet_title, "row": grid_source_row, "heading": "Device Brands"},
            "parser_version": PARSER_VERSION,
        })

    for row_number in ordered_rows:
        if row_number in consumed_matrix_rows:
            continue
        row = rows[row_number]
        indicator = str(row.get(mapping["indicator_column"], {}).get("value") or "").strip()
        if not indicator:
            continue
        definition = str(row.get(mapping["definition_column"], {}).get("value") or "").strip()
        if not definition:
            current_level = min(current_level + 1, 3) if last_kind == "heading" else 1
            current_heading_code = _unique_code(indicator, used_heading_codes, f"HEADING_{sheet_index}_{row_number}")
            current_heading_title = indicator[:255]
            headings.append({
                "heading_code": current_heading_code,
                "title": current_heading_title,
                "level": current_level,
                "source_order": row_number,
                "source_row": row_number,
                "source": {"sheet": sheet_title, "row": row_number},
                "parser_version": PARSER_VERSION,
            })
            last_kind = "heading"
            continue

        data_type_value = row.get(mapping.get("data_type_column"), {}).get("value") if mapping.get("data_type_column") else ""
        field_type, unknown_type = _metadata_field_type(data_type_value)
        if unknown_type:
            warnings.append({
                "code": f"UNKNOWN_DATA_TYPE_{sheet_index}_{row_number}",
                "severity": "WARNING",
                "message": f'{sheet_title} row {row_number} uses unknown data type "{data_type_value}" and was mapped to text.',
            })
        unit_value = row.get(mapping.get("unit_column"), {}).get("value") if mapping.get("unit_column") else ""
        required_value = row.get(mapping.get("required_column"), {}).get("value") if mapping.get("required_column") else ""
        options_value = row.get(mapping.get("options_column"), {}).get("value") if mapping.get("options_column") else ""
        options = _split_options(options_value)
        if options and field_type == "text":
            field_type = "select"
        label = indicator[:255]
        fields.append({
            "source_order": row_number,
            "field_code": _unique_code(label, used_field_codes, f"FIELD_{sheet_index}_{row_number}"),
            "heading_code": current_heading_code,
            "label": label,
            "field_type": field_type,
            "unit": str(unit_value or _unit(label)).strip()[:50],
            "is_required": _metadata_key(required_value) in TRUE_VALUES,
            "help_text": definition,
            "formula": "",
            "options": options,
            "source": {"sheet": sheet_title, "row": row_number, "heading": current_heading_title},
            "parser_version": PARSER_VERSION,
        })
        last_kind = "field"

    mapping["candidates"] = _column_candidates(rows)
    return headings, fields, grids, mapping


def _parse_stream_sheet(archive, sheet_path, shared_strings, styles):
    cells = {}
    merged_ranges = []
    table_relationship_ids = []
    hidden_rows = set()
    with archive.open(sheet_path) as source:
        for _event, element in ET.iterparse(source, events=("end",)):
            if element.tag == f"{SHEET_NS}c":
                coordinate = element.attrib.get("r")
                if coordinate:
                    try:
                        row, column = coordinate_to_tuple(coordinate)
                    except ValueError:
                        row, column = 0, 0
                    if 1 <= row <= MAX_SOURCE_ROWS and 1 <= column <= MAX_SOURCE_COLUMNS:
                        data = _stream_cell_value(element, shared_strings, styles)
                        if data["formula"] or data["value"] not in (None, ""):
                            cells[(row, column)] = data
                element.clear()
            elif element.tag == f"{SHEET_NS}mergeCell":
                reference = element.attrib.get("ref")
                if reference:
                    try:
                        merged_ranges.append(range_boundaries(reference))
                    except ValueError:
                        pass
                element.clear()
            elif element.tag == f"{SHEET_NS}tablePart":
                relationship_id = element.attrib.get(f"{OFFICE_REL_NS}id")
                if relationship_id:
                    table_relationship_ids.append(relationship_id)
                element.clear()
            elif element.tag == f"{SHEET_NS}row":
                if element.attrib.get("hidden", "0").lower() in {"1", "true"}:
                    try:
                        hidden_rows.add(int(element.attrib.get("r", "0")))
                    except ValueError:
                        pass
                element.clear()
    return cells, merged_ranges, table_relationship_ids, hidden_rows


def _streaming_table(archive, table_path, cells, sheet_title, used_grid_codes):
    root = ET.fromstring(archive.read(table_path))
    reference = root.attrib.get("ref")
    if not reference:
        return None
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    if min_row > MAX_SOURCE_ROWS or min_col > MAX_SOURCE_COLUMNS:
        return None
    max_row, max_col = min(max_row, MAX_SOURCE_ROWS), min(max_col, MAX_SOURCE_COLUMNS)
    title = root.attrib.get("displayName") or root.attrib.get("name") or "Table"
    column_root = root.find(f"{SHEET_NS}tableColumns")
    column_names = [item.attrib.get("name") or f"Column {index}" for index, item in enumerate(list(column_root or []), start=1)]
    if not column_names:
        column_names = [
            str(cells.get((min_row, column), {}).get("value") or f"Column {column - min_col + 1}")
            for column in range(min_col, max_col + 1)
        ]
    grid = {
        "grid_code": _unique_code(title, used_grid_codes, "TABLE"),
        "title": title,
        "row_mode": "REPEATABLE", "min_rows": 0, "instructions": "",
        "columns": [], "fixed_rows": [],
        "source": {"sheet": sheet_title, "row": min_row, "heading": ""},
        "parser_version": PARSER_VERSION,
        "bounds": (min_col, min_row, max_col, max_row),
    }
    column_codes = set()
    for offset, label_value in enumerate(column_names[: max_col - min_col + 1]):
        label = str(label_value).strip() or f"Column {offset + 1}"
        sample = next(
            (cells[(row, min_col + offset)] for row in range(min_row + 1, max_row + 1) if (row, min_col + offset) in cells),
            cells.get((min_row, min_col + offset), {"value": label, "formula": "", "style": {}}),
        )
        grid["columns"].append({
            "column_code": _unique_code(label, column_codes, f"COLUMN_{offset + 1}"),
            "label": label,
            "field_type": _stream_field_type(sample),
            "unit": _unit(label),
            "is_required": True,
        })
    first_header = str(column_names[0] if column_names else "").strip().lower()
    if first_header in FIXED_ROW_HEADERS:
        labels = [
            str(cells[(row, min_col)]["value"]).strip()
            for row in range(min_row + 1, max_row + 1)
            if (row, min_col) in cells and cells[(row, min_col)]["value"] not in (None, "")
        ]
        if labels and len(set(labels)) == len(labels):
            grid["row_mode"] = "FIXED"
            grid["fixed_rows"] = labels
            grid["min_rows"] = len(labels)
    return grid


def _parse_workbook_streaming(path, column_mappings=None):
    warnings = []
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError(f"The workbook could not be read: {exc}") from exc
    with archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names:
            raise ValueError("The workbook is missing its Excel workbook definition.")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            warnings.append({
                "code": "MACRO_CONTENT",
                "severity": "BLOCKING",
                "message": "The workbook contains macro content. Save a macro-free .xlsx copy before confirming it.",
            })
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_relationships = _relationship_map(archive, "xl/workbook.xml")
        sheets_root = workbook_root.find(f"{SHEET_NS}sheets")
        visible_sheets = []
        for item in list(sheets_root or []):
            if item.attrib.get("state", "visible") != "visible":
                continue
            relationship_id = item.attrib.get(f"{OFFICE_REL_NS}id")
            sheet_path = workbook_relationships.get(relationship_id)
            if sheet_path and sheet_path in names:
                visible_sheets.append((item.attrib.get("name", "Worksheet"), sheet_path))
        if not visible_sheets:
            raise ValueError("The workbook has no visible worksheets.")
        if any(name.startswith("xl/externalLinks/") for name in names):
            warnings.append({
                "code": "EXTERNAL_LINKS_REMOVED",
                "severity": "BLOCKING",
                "message": "The workbook contains external links. Remove them before confirming the generated template.",
            })

        shared_strings = _shared_strings(archive)
        styles = _style_metadata(archive)
        sections = []
        used_section_codes = set()
        for sheet_index, (sheet_title, sheet_path) in enumerate(visible_sheets, start=1):
            cells, _merged_ranges, _table_ids, hidden_rows = _parse_stream_sheet(
                archive, sheet_path, shared_strings, styles,
            )
            all_rows = {}
            for (row, column), cell_data in cells.items():
                all_rows.setdefault(row, {})[column] = cell_data
            excluded_hidden_rows = sorted(set(all_rows) & hidden_rows)
            rows = {row: values for row, values in all_rows.items() if row not in hidden_rows}
            override = (column_mappings or {}).get(sheet_title)
            headings, fields, grids, column_mapping = _definition_content(
                rows, sheet_title, sheet_index, warnings, override,
            )
            if not fields and not headings and not grids:
                warnings.append({
                    "code": f"EMPTY_WORKSHEET_SECTION_{sheet_index}",
                    "severity": "WARNING",
                    "message": f'{sheet_title} is a visible worksheet, so it was retained as an empty section.',
                })
            sections.append({
                "section_code": _unique_code(sheet_title, used_section_codes, f"WORKSHEET_{sheet_index}", max_length=50),
                "title": str(sheet_title).strip()[:255],
                "instructions": f'Generated from worksheet "{sheet_title}".',
                "worksheet_order": sheet_index,
                "source": {"sheet": sheet_title, "sheet_index": sheet_index},
                "row_visibility": {
                    "policy": "visible-only",
                    "visible_source_row_count": len(rows),
                    "excluded_hidden_row_count": len(excluded_hidden_rows),
                },
                "column_mapping": column_mapping,
                "headings": sorted(headings, key=lambda item: item["source_order"]),
                "fields": sorted(fields, key=lambda item: item["source_order"]),
                "grids": sorted(grids, key=lambda item: item["source_order"]),
                "counts": {
                    "scalar_field_count": len(fields),
                    "table_count": len(grids),
                    "grid_input_count": sum(len(grid.get("fixed_rows", [])) * len(grid.get("columns", [])) for grid in grids),
                },
            })
    return {
        "parser_version": PARSER_VERSION,
        "grouping": {
            "strategy": "worksheet-tabs",
            "isolation": "worksheet-isolated",
            "visible_worksheet_count": len(sections),
            "engine": "streaming",
        },
        "sections": sections,
    }, warnings


def parse_workbook(path, column_mappings=None):
    warnings = []
    if Path(path).stat().st_size >= STREAMING_THRESHOLD_BYTES:
        return _parse_workbook_streaming(path, column_mappings=column_mappings)
    try:
        with zipfile.ZipFile(path) as archive:
            has_external_links = any(name.startswith("xl/externalLinks/") for name in archive.namelist())
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError(f"The workbook could not be read: {exc}") from exc

    if has_external_links or getattr(workbook, "_external_links", None):
        warnings.append({
            "code": "EXTERNAL_LINKS_REMOVED",
            "severity": "BLOCKING",
            "message": "The workbook contains external links. Remove them before confirming the generated template.",
        })

    visible_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if not visible_sheets:
        raise ValueError("The workbook has no visible worksheets.")

    sections = []
    used_section_codes = set()
    for sheet_index, sheet in enumerate(visible_sheets, start=1):
        max_row, max_col = min(sheet.max_row, MAX_SOURCE_ROWS), min(sheet.max_column, MAX_SOURCE_COLUMNS)
        all_rows = {}
        for row in range(1, max_row + 1):
            for column in range(1, max_col + 1):
                cell = sheet.cell(row, column)
                if cell.value not in (None, ""):
                    all_rows.setdefault(row, {})[column] = {
                        "value": cell.value,
                        "formula": str(cell.value)[1:] if isinstance(cell.value, str) and cell.value.startswith("=") else "",
                        "style": {},
                    }
        hidden_rows = {
            row for row in all_rows
            if bool(sheet.row_dimensions[row].hidden)
        }
        rows = {row: values for row, values in all_rows.items() if row not in hidden_rows}
        override = (column_mappings or {}).get(sheet.title)
        headings, fields, grids, column_mapping = _definition_content(
            rows, sheet.title, sheet_index, warnings, override,
        )
        if not fields and not headings and not grids:
            warnings.append({
                "code": f"EMPTY_WORKSHEET_SECTION_{sheet_index}",
                "severity": "WARNING",
                "message": f'{sheet.title} is a visible worksheet, so it was retained as an empty section.',
            })
        sections.append({
            "section_code": _unique_code(
                sheet.title, used_section_codes, f"WORKSHEET_{sheet_index}", max_length=50,
            ),
            "title": str(sheet.title).strip()[:255],
            "instructions": f'Generated from worksheet "{sheet.title}".',
            "worksheet_order": sheet_index,
            "source": {"sheet": sheet.title, "sheet_index": sheet_index},
            "row_visibility": {
                "policy": "visible-only",
                "visible_source_row_count": len(rows),
                "excluded_hidden_row_count": len(hidden_rows),
            },
            "column_mapping": column_mapping,
            "headings": sorted(headings, key=lambda item: item["source_order"]),
            "fields": sorted(fields, key=lambda item: item["source_order"]),
            "grids": sorted(grids, key=lambda item: item["source_order"]),
            "counts": {
                "scalar_field_count": len(fields),
                "table_count": len(grids),
                "grid_input_count": sum(len(grid.get("fixed_rows", [])) * len(grid.get("columns", [])) for grid in grids),
            },
        })

    return {
        "parser_version": PARSER_VERSION,
        "grouping": {
            "strategy": "worksheet-tabs",
            "isolation": "worksheet-isolated",
            "visible_worksheet_count": len(visible_sheets),
        },
        "sections": sections,
    }, warnings


def validate_schema(schema):
    sections = schema.get("sections") if isinstance(schema, dict) else None
    if not isinstance(sections, list) or not sections:
        raise ValueError("At least one section is required.")
    section_codes = set()
    for section in sections:
        code = normalize_code(section.get("section_code"), fallback="SECTION", max_length=50)
        if code in section_codes:
            raise ValueError(f"Duplicate section code: {code}.")
        section_codes.add(code)
        section["section_code"] = code
        if not str(section.get("title", "")).strip():
            raise ValueError(f"Section {code} requires a title.")
        if schema.get("parser_version") == PARSER_VERSION:
            mapping = section.get("column_mapping") or {}
            if not mapping.get("indicator_column") or not mapping.get("definition_column"):
                raise ValueError(
                    f"Section {code} requires explicit Indicator and Definition column mappings."
                )
            source_sheet = str((section.get("source") or {}).get("sheet") or "").strip()
            if not source_sheet:
                raise ValueError(f"Section {code} requires a source worksheet.")
            for item_kind in ("headings", "fields", "grids"):
                for item in section.get(item_kind, []):
                    item_sheet = str((item.get("source") or {}).get("sheet") or "").strip()
                    if item_sheet != source_sheet:
                        item_code = (
                            item.get("heading_code") or item.get("field_code")
                            or item.get("grid_code") or "unknown"
                        )
                        raise ValueError(
                            f"{item_kind[:-1].title()} {item_code} belongs to worksheet "
                            f'"{item_sheet or "unknown"}", not section worksheet "{source_sheet}".'
                        )
            for grid in section.get("grids", []):
                for column in grid.get("columns", []):
                    column_sheet = str((column.get("source") or {}).get("sheet") or "").strip()
                    if column_sheet != source_sheet:
                        raise ValueError(
                            f"Grid column {column.get('column_code') or 'unknown'} does not belong to section worksheet \"{source_sheet}\"."
                        )
                for fixed_row in grid.get("fixed_row_sources", []):
                    row_sheet = str((fixed_row.get("source") or {}).get("sheet") or "").strip()
                    if row_sheet != source_sheet:
                        raise ValueError(
                            f"Grid row {fixed_row.get('label') or 'unknown'} does not belong to section worksheet \"{source_sheet}\"."
                        )
        heading_codes, field_codes, grid_codes = set(), set(), set()
        for heading in section.get("headings", []):
            heading["heading_code"] = _unique_code(
                heading.get("heading_code") or heading.get("title"), heading_codes, "HEADING",
            )
            if not str(heading.get("title", "")).strip():
                raise ValueError(f"Heading {heading['heading_code']} requires a title.")
            level = int(heading.get("level", 1))
            if level not in {1, 2, 3}:
                raise ValueError(f"Heading {heading['heading_code']} level must be between 1 and 3.")
            heading["level"] = level
        for field in section.get("fields", []):
            field["field_code"] = _unique_code(field.get("field_code") or field.get("label"), field_codes, "FIELD")
            if field.get("field_type") not in dict(FormField._meta.get_field("field_type").choices):
                raise ValueError(f"Unsupported field type in {field['field_code']}.")
            if field.get("heading_code") and field["heading_code"] not in heading_codes:
                raise ValueError(f"Field {field['field_code']} references an unknown heading.")
        for grid in section.get("grids", []):
            grid["grid_code"] = _unique_code(grid.get("grid_code") or grid.get("title"), grid_codes, "GRID")
            if not grid.get("columns"):
                raise ValueError(f"Grid {grid['grid_code']} requires at least one column.")
    return schema


def process_workbook_import_record(import_id):
    """Scan and parse one stored workbook outside the upload request when needed."""
    close_old_connections()
    try:
        workbook_import = FormWorkbookImport.objects.get(pk=import_id)
        if workbook_import.parse_status in {"READY", "CONFIRMED"}:
            return workbook_import.parse_status
        full_path = os.path.join(settings.PRIVATE_UPLOAD_ROOT, workbook_import.storage_path)
        try:
            scan = scan_path(full_path)
            workbook_import.scan_status = scan["status"]
            workbook_import.scan_engine = scan["engine"]
            workbook_import.scan_details = scan["details"]
            if scan["status"] == "CLEAN":
                if os.path.splitext(workbook_import.file_name)[1].lower() == ".xlsx":
                    schema, warnings = parse_workbook(full_path, column_mappings=(workbook_import.mapping_decisions or {}).get("column_mappings", {}))
                else:
                    schema, warnings = parse_document_source(full_path)
                workbook_import.detected_schema = schema
                workbook_import.warnings = warnings
                workbook_import.parser_version = PARSER_VERSION
                workbook_import.parse_status = "READY"
            else:
                workbook_import.parse_status = "FAILED"
        except Exception as exc:
            if workbook_import.scan_status == "PENDING":
                workbook_import.scan_status = "ERROR"
            workbook_import.parse_status = "FAILED"
            workbook_import.scan_details = str(exc)
        workbook_import.save()
        return workbook_import.parse_status
    finally:
        close_old_connections()


@transaction.atomic
def create_template_from_import(workbook_import, user):
    if workbook_import.resulting_template_id:
        return workbook_import.resulting_template
    code = normalize_form_code(workbook_import.form_code)
    schema = validate_schema(workbook_import.detected_schema)
    family, created = FormFamily.objects.get_or_create(
        code=code,
        defaults={
            "name": workbook_import.name, "canonical_frequency": workbook_import.frequency,
            "frequency_decision_status": "APPROVED", "frequency_decision_reference": f"Workbook import {workbook_import.id}",
            "source_owner": user.name or user.email,
        },
    )
    if not created and family.canonical_frequency and family.canonical_frequency != workbook_import.frequency:
        raise ValueError("The form family already uses a different canonical frequency.")
    if FormTemplate.objects.filter(family=family, version=workbook_import.version).exists():
        raise ValueError("This form family and version already exist.")
    form = FormTemplate.objects.create(
        family=family, form_code=code, name=workbook_import.name, version=workbook_import.version,
        sector=workbook_import.sector, provider_category=workbook_import.provider_category,
        frequency=workbook_import.frequency, effective_from=date.today(), status="DRAFT",
        source_reference=f"Workbook: {workbook_import.file_name}", source_sha256=workbook_import.sha256,
        mapping_complete=True, mapping_basis="SOURCE_FORM", prepared_by=user,
    )
    for section_order, section_data in enumerate(schema["sections"], start=1):
        section = FormSection.objects.create(
            form_template=form, section_code=section_data["section_code"], title=str(section_data["title"])[:255],
            instructions=section_data.get("instructions", ""), sort_order=section_order,
        )
        heading_map = {}
        for heading_order, heading_data in enumerate(section_data.get("headings", []), start=1):
            heading = FormHeading.objects.create(
                section=section,
                heading_code=heading_data["heading_code"],
                title=str(heading_data.get("title") or heading_data["heading_code"])[:255],
                level=int(heading_data.get("level", 1)),
                sort_order=int(heading_data.get("source_order", heading_order)),
                source_row=max(0, int(heading_data.get("source_row", 0))),
            )
            heading_map[heading.heading_code] = heading
        for field_order, field_data in enumerate(section_data.get("fields", []), start=1):
            field = FormField.objects.create(
                section=section, heading=heading_map.get(field_data.get("heading_code")),
                field_code=field_data["field_code"], label=str(field_data.get("label") or field_data["field_code"])[:255],
                field_type=field_data.get("field_type", "text"), unit=str(field_data.get("unit", ""))[:50],
                is_required=bool(field_data.get("is_required", False)), help_text=field_data.get("help_text", ""),
                formula=field_data.get("formula", ""), sort_order=int(field_data.get("source_order", field_order)),
                source_sheet=str((field_data.get("source") or {}).get("sheet") or "")[:255],
                source_row=(field_data.get("source") or {}).get("row"),
            )
            SelectOption.objects.bulk_create([
                SelectOption(field=field, value=str(option)[:100], label=str(option)[:255], sort_order=index)
                for index, option in enumerate(field_data.get("options", []), start=1)
            ])
        for grid_order, grid_data in enumerate(section_data.get("grids", []), start=1):
            grid = FormGrid.objects.create(
                section=section, grid_code=grid_data["grid_code"], title=str(grid_data.get("title") or grid_data["grid_code"])[:255],
                row_mode=grid_data.get("row_mode", "REPEATABLE"), min_rows=max(0, int(grid_data.get("min_rows", 0))),
                instructions=grid_data.get("instructions", ""), sort_order=grid_order,
                source_sheet=str((grid_data.get("source") or {}).get("sheet") or "")[:255],
                source_row=(grid_data.get("source") or {}).get("row"),
            )
            GridColumn.objects.bulk_create([
                GridColumn(
                    grid=grid, column_code=normalize_code(column.get("column_code") or column.get("label"), fallback=f"COLUMN_{index}"),
                    label=str(column.get("label") or f"Column {index}")[:255], field_type=column.get("field_type", "text"),
                    unit=str(column.get("unit", ""))[:50], is_required=bool(column.get("is_required", False)), sort_order=index,
                    source_sheet=str((column.get("source") or {}).get("sheet") or "")[:255],
                    source_row=(column.get("source") or {}).get("row"),
                ) for index, column in enumerate(grid_data.get("columns", []), start=1)
            ])
            row_sources = {
                str(item.get("label")): item.get("source") or {}
                for item in grid_data.get("fixed_row_sources", [])
            }
            GridRow.objects.bulk_create([
                GridRow(
                    grid=grid, row_label=str(label)[:255], sort_order=index,
                    source_sheet=str(row_sources.get(str(label), {}).get("sheet") or grid.source_sheet)[:255],
                    source_rows=list(row_sources.get(str(label), {}).get("rows") or []),
                )
                for index, label in enumerate(grid_data.get("fixed_rows", []), start=1)
            ])
    workbook_import.resulting_template = form
    workbook_import.parse_status = "CONFIRMED"
    workbook_import.detected_schema = schema
    workbook_import.save(update_fields=["resulting_template", "parse_status", "detected_schema", "updated_at"])
    return form
