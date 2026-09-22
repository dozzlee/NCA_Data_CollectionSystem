import csv
import hashlib
import io

from django.http import HttpResponse
from django.db.models import Q
from django.utils import timezone
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.services import record_audit
from apps.submissions.models import ExpectedSubmission, WORKFLOW_STATUSES
from apps.submissions.models import ReportingPeriod
from apps.forms_engine.models import FormTemplate
from apps.users.permissions import IsNCAUser
from .models import ExportLog
from .services import CANONICAL_HEADERS, canonical_row_dicts, canonical_values


ALLOWED_FILTERS = {
    "workflow_status": "workflow_status", "period": "period_id", "provider": "provider_id",
    "sector": "provider__sector", "provider_category": "provider__category",
    "form": "form_template_id", "assigned_officer": "assigned_officer_id", "due_state": "due_state",
}


def spreadsheet_safe(value):
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def validate_filters(filters):
    if not isinstance(filters, dict):
        return None, "filters must be an object."
    unknown = sorted(set(filters) - set(ALLOWED_FILTERS))
    if unknown:
        return None, f"Unsupported filters: {', '.join(unknown)}"
    query_filters = {}
    valid_workflows = {value for value, _ in WORKFLOW_STATUSES}
    for key, value in filters.items():
        if value in ("", None):
            continue
        if key in {"period", "provider", "form", "assigned_officer"}:
            if isinstance(value, bool): return None, f"{key} must be a numeric ID."
            try: value = int(value)
            except (TypeError, ValueError): return None, f"{key} must be a numeric ID."
        if key == "workflow_status" and value not in valid_workflows: return None, "Invalid workflow_status."
        if key == "workflow_status" and value != "APPROVED":
            return None, "Operational exports contain NCA-approved submissions only."
        if key == "sector" and value not in {"TELECOM", "BROADCASTING"}: return None, "Invalid sector."
        query_filters[ALLOWED_FILTERS[key]] = value
    return query_filters, None


def approved_submissions(query_filters):
    query_filters = dict(query_filters)
    query_filters["workflow_status"] = "APPROVED"
    expected = ExpectedSubmission.objects.filter(**query_filters)
    submission_ids = []
    for item in expected.prefetch_related("versions"):
        approved = item.versions.filter(reviewed_at__isnull=False).order_by("-version").first()
        if approved:
            submission_ids.append(approved.id)
    from apps.submissions.models import Submission
    return Submission.objects.filter(id__in=submission_ids)


def record_export(*, request, export_type, filters, content, row_count, mime_type):
    digest = hashlib.sha256(content).hexdigest()
    ExportLog.objects.create(
        export_type=export_type, filters=filters, generated_by=request.user,
        row_count=row_count, mime_type=mime_type, file_size=len(content), sha256=digest,
    )
    record_audit(
        user=request.user, action=f"EXPORT_{export_type}", entity_type="Export", entity_id=digest,
        after={"rows": row_count, "filters": filters, "sha256": digest, "approved_only": True},
        ip_address=request.META.get("REMOTE_ADDR"),
    )
    return digest


class CSVExportView(APIView):
    permission_classes = [IsNCAUser]

    def post(self, request):
        filters = request.data.get("filters", {})
        query_filters, error = validate_filters(filters)
        if error: return Response({"detail": error}, status=400)
        rows = canonical_values(canonical_row_dicts(approved_submissions(query_filters)))
        output = io.StringIO(newline=""); output.write("\ufeff")
        writer = csv.writer(output, lineterminator="\r\n"); writer.writerow(CANONICAL_HEADERS)
        count = 0
        for row in rows:
            writer.writerow([spreadsheet_safe(value) for value in row]); count += 1
        content = output.getvalue().encode("utf-8")
        record_export(request=request, export_type="CSV", filters=filters, content=content,
            row_count=count, mime_type="text/csv")
        filename = f"nca_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
        response = HttpResponse(content, content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        return response


class PDFExportView(APIView):
    permission_classes = [IsNCAUser]

    def post(self, request):
        filters = request.data.get("filters", {})
        query_filters, error = validate_filters(filters)
        if error:
            return Response({"detail": error}, status=400)

        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

        row_dicts = list(canonical_row_dicts(approved_submissions(query_filters)))
        output = io.BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=landscape(A3), leftMargin=10 * mm, rightMargin=10 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm,
            title="NCA approved operational data export",
        )
        styles = getSampleStyleSheet()
        cell_style = ParagraphStyle("ExportCell", parent=styles["BodyText"], fontSize=5.5, leading=6.5, alignment=TA_LEFT)
        header_style = ParagraphStyle("ExportHeader", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
        columns = [
            ("provider", "registered_company_name"), ("form", "form_code"), ("version", "form_version"),
            ("period", "period_name"), ("submission", "submission_id"), ("section", "section_name"),
            ("kind", "value_kind"), ("field/grid", "field_name"), ("grid row", "grid_row_label"),
            ("grid column", "grid_column_name"), ("value", "submitted_value"), ("status", "value_status"),
            ("explanation", "explanation"), ("disposition", "non_filled_disposition"),
            ("review", "latest_review_action"), ("compliance", "open_compliance_flags"),
        ]
        story = [
            Paragraph("NCA Approved Operational Data Export", styles["Title"]),
            Spacer(1, 3 * mm),
            Paragraph(
                f"Generated {timezone.now().isoformat()} · {len(row_dicts):,} canonical rows · Approved submissions only",
                styles["BodyText"],
            ),
            Paragraph(f"Filters: {filters or 'None'}", styles["BodyText"]),
            PageBreak(),
        ]
        table_data = [[Paragraph(label, header_style) for label, _ in columns]]
        for row in row_dicts:
            table_data.append([Paragraph(str(row.get(key, "") or ""), cell_style) for _, key in columns])
        table = LongTable(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#002d5b")),
            ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#cbd5df")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7f9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(table)
        document.build(story)
        content = output.getvalue()
        record_export(request=request, export_type="PDF", filters=filters, content=content,
            row_count=len(row_dicts), mime_type="application/pdf")
        filename = f"nca_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        response = HttpResponse(content, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["X-Content-Type-Options"] = "nosniff"
        return response


class ExportLogListView(APIView):
    permission_classes = [IsNCAUser]
    def get(self, request):
        logs = ExportLog.objects.select_related("generated_by").order_by("-generated_at")[:50]
        return Response([{"id": item.id, "export_type": item.export_type, "filters": item.filters,
            "generated_by": item.generated_by.name, "generated_at": item.generated_at, "row_count": item.row_count,
            "mime_type": item.mime_type, "file_size": item.file_size, "sha256": item.sha256} for item in logs])


def _catalogue_query(request):
    period_id = request.query_params.get("period") or None
    search = str(request.query_params.get("search") or "").strip()
    forms = FormTemplate.objects.select_related("family").prefetch_related(
        "sections__fields", "sections__grids__columns", "sections__grids__fixed_rows",
    ).order_by("form_code", "version")
    obligations = ExpectedSubmission.objects.select_related(
        "provider", "form_template", "period",
    ).prefetch_related("versions").order_by("-period__year", "-period__month", "provider__registered_name")
    if period_id:
        forms = forms.filter(expectedsubmission__period_id=period_id).distinct()
        obligations = obligations.filter(period_id=period_id)
    if search:
        forms = forms.filter(Q(form_code__icontains=search) | Q(name__icontains=search) | Q(version__icontains=search))
        obligations = obligations.filter(
            Q(form_template__form_code__icontains=search) | Q(form_template__name__icontains=search)
            | Q(provider__registered_name__icontains=search) | Q(period__name__icontains=search)
            | Q(versions__submission_reference__icontains=search)
        ).distinct()
    return forms, obligations


def _form_item(form):
    sections = list(form.sections.all())
    scalar_count = sum(section.fields.count() for section in sections)
    grid_count = sum(section.grids.count() for section in sections)
    return {
        "id": form.id, "code": form.form_code, "name": form.name, "version": form.version,
        "frequency": form.frequency, "status": form.status, "approval_status": form.approval_status,
        "sections": len(sections), "indicators": scalar_count,
        "tables": grid_count, "created_at": form.created_at,
    }


def _obligation_item(expected):
    latest = expected.versions.order_by("-version", "-id").first()
    template = expected.form_template
    return {
        "id": expected.id, "provider": expected.provider.registered_name,
        "form_code": template.form_code if template else expected.form_code_snapshot,
        "form_name": template.name if template else expected.form_name_snapshot,
        "period_id": expected.period_id, "period": expected.period.name,
        "workflow_status": expected.workflow_status, "provider_status": expected.provider_status,
        "due_at": expected.effective_due_at, "sent_at": expected.created_at,
        "submission_id": latest.id if latest else None,
        "submission_reference": latest.submission_reference if latest else "",
        "completion_pct": str(latest.completion_pct if latest else 0),
        "submitted_at": latest.submitted_at if latest else None,
    }


class ExportCatalogueView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        forms, obligations = _catalogue_query(request)
        form_items = [_form_item(form) for form in forms]
        obligation_items = [_obligation_item(item) for item in obligations]
        return Response({
            "summary": {
                "forms": len(form_items), "forms_sent": len(obligation_items),
                "submitted": sum(item["workflow_status"] in {"SUBMITTED", "UNDER_REVIEW", "RESUBMITTED", "APPROVED", "REJECTED"} for item in obligation_items),
                "approved": sum(item["workflow_status"] == "APPROVED" for item in obligation_items),
                "providers": len({item["provider"] for item in obligation_items}),
            },
            "periods": list(ReportingPeriod.objects.order_by("-year", "-month", "-quarter", "-id").values("id", "name", "frequency", "year", "month", "quarter", "status")),
            "forms": form_items,
            "submissions": obligation_items,
            "approved_data_only": True,
        })


class ExportCatalogueWorkbookView(APIView):
    permission_classes = [IsNCAUser]

    def post(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        period_id = request.data.get("period") or None
        form_ids = request.data.get("form_ids", [])
        expected_ids = request.data.get("expected_submission_ids", [])
        try:
            period_id = int(period_id) if period_id else None
            form_ids = [int(value) for value in form_ids]
            expected_ids = [int(value) for value in expected_ids]
        except (TypeError, ValueError):
            return Response({"detail": "Period, form and submission selections must use numeric IDs."}, status=400)

        query_request = request._request
        query_request.GET = query_request.GET.copy()
        if period_id:
            query_request.GET["period"] = str(period_id)
        forms, obligations = _catalogue_query(request)
        if form_ids:
            forms = forms.filter(id__in=form_ids)
        if expected_ids:
            obligations = obligations.filter(id__in=expected_ids)
        if form_ids and not expected_ids:
            obligations = obligations.filter(form_template_id__in=form_ids)
        forms = list(forms)
        obligations = list(obligations)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["NCA Data and Form Catalogue"])
        summary.append(["Generated", timezone.now().isoformat()])
        summary.append(["Selected period", ReportingPeriod.objects.filter(pk=period_id).values_list("name", flat=True).first() or "All periods"])
        summary.append(["Forms", len(forms)])
        summary.append(["Forms sent", len(obligations)])
        summary.append(["Approved submissions", sum(item.workflow_status == "APPROVED" for item in obligations)])

        forms_sheet = workbook.create_sheet("Forms")
        forms_sheet.append(["Form ID", "Code", "Name", "Version", "Frequency", "Status", "Approval", "Sections", "Indicators", "Tables"])
        for form in forms:
            item = _form_item(form)
            forms_sheet.append([item["id"], item["code"], item["name"], item["version"], item["frequency"], item["status"], item["approval_status"], item["sections"], item["indicators"], item["tables"]])

        definitions = workbook.create_sheet("Form Definitions")
        definitions.append(["Form Code", "Version", "Section", "Target Type", "Code", "Label", "Data Type", "Unit", "Required"])
        for form in forms:
            for section in form.sections.all():
                for field in section.fields.all():
                    definitions.append([form.form_code, form.version, section.title, "FIELD", field.field_code, field.label, field.field_type, field.unit, "Yes" if field.is_required else "No"])
                for grid in section.grids.all():
                    for column in grid.columns.all():
                        definitions.append([form.form_code, form.version, section.title, "TABLE", f"{grid.grid_code}.{column.column_code}", f"{grid.title} / {column.label}", column.field_type, column.unit, "Yes" if column.is_required else "No"])

        submissions_sheet = workbook.create_sheet("Forms Sent")
        submissions_sheet.append(["Task ID", "Provider", "Form", "Period", "Status", "Sent", "Due", "Submission Reference", "Submitted", "Completion %"])
        for expected in obligations:
            item = _obligation_item(expected)
            submissions_sheet.append([item["id"], item["provider"], item["form_code"], item["period"], item["workflow_status"], item["sent_at"].isoformat() if item["sent_at"] else "", item["due_at"].isoformat() if item["due_at"] else "", item["submission_reference"], item["submitted_at"].isoformat() if item["submitted_at"] else "", item["completion_pct"]])

        approved_ids = [item.id for item in obligations if item.workflow_status == "APPROVED"]
        data_sheet = workbook.create_sheet("Approved Data")
        data_sheet.append(CANONICAL_HEADERS)
        approved = approved_submissions({"id__in": approved_ids}) if approved_ids else []
        data_rows = canonical_values(canonical_row_dicts(approved)) if approved_ids else []
        row_count = 0
        for row in data_rows:
            data_sheet.append([spreadsheet_safe(value) for value in row])
            row_count += 1

        header_fill = PatternFill("solid", fgColor="002D5B")
        for sheet in workbook.worksheets:
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
            sheet.freeze_panes = "A2"
            for index, column in enumerate(sheet.columns, start=1):
                width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[get_column_letter(index)].width = width

        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        filters = {"period": period_id, "form_ids": form_ids, "expected_submission_ids": expected_ids}
        record_export(request=request, export_type="XLSX", filters=filters, content=content, row_count=row_count, mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="nca_data_form_catalogue_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        response["X-Content-Type-Options"] = "nosniff"
        return response
