"""Private XLSX-to-submission parsing and deterministic indicator matching."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.forms_engine.models import FormField, FormGrid
from apps.submissions.models import (
    ProviderEditBatch, ProviderEditItem, SubmissionValue,
)
from apps.submissions.readiness import refresh_submission_completion
from apps.submissions.workflow import complete_submission_revision
from .models import ExcelImportMappingProfile, SubmissionExcelImport, SubmissionExcelImportMatch


PARSER_VERSION = "submission-xlsx-v3"
MATCHER_VERSION = "template-coordinates-v3"
HEADER_ALIASES = {
    "code": {"code", "indicator code", "indicator number", "indicator no", "id", "reference"},
    "indicator": {"indicator", "indicator name", "name", "label", "industry data", "metric"},
    "definition": {"definition", "definitions", "description", "help text"},
    "data_type": {"data type", "datatype", "type", "format", "unit"},
    "value": {"value", "response", "data entry", "current value", "reported value", "answer"},
}
NUMERIC_TYPES = {"number", "currency", "percentage", "formula"}


def normalize(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    aliases = {
        "no": "number", "num": "number", "qty": "quantity",
        "subs": "subscriptions", "subscriber": "subscribers",
        "post paid": "postpaid", "pre paid": "prepaid",
        "cellular": "mobile", "telephony": "voice",
    }
    for source, replacement in aliases.items():
        text = re.sub(rf"\b{re.escape(source)}\b", replacement, text)
    return re.sub(r"\s+", " ", text).strip()


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _similarity(left, right):
    left, right = normalize(left), normalize(right)
    if not left or not right:
        return 0.0
    sequence = SequenceMatcher(None, left, right).ratio()
    left_tokens, right_tokens = set(left.split()), set(right.split())
    token = len(left_tokens & right_tokens) / max(len(left_tokens | right_tokens), 1)
    return max(sequence, token)


def _tokens(value):
    return set(normalize(value).split())


def _header_kind(value):
    normalized = normalize(value)
    for kind, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return kind
    return None


def _layout_fingerprint(layout):
    raw = json.dumps(layout, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _cell(rows, row_number, column_number):
    if not row_number or not column_number or row_number > len(rows):
        return None
    row = rows[row_number - 1]
    return row[column_number - 1] if column_number <= len(row) else None


def _detect_header(rows):
    best = None
    for row_number, row in enumerate(rows[:60], 1):
        found = {}
        for column, value in enumerate(row[:60], 1):
            kind = _header_kind(value)
            if kind and kind not in found:
                found[kind] = column
        score = len(found) + (2 if "indicator" in found else 0) + (1 if "value" in found else 0)
        if "indicator" in found and (best is None or score > best[0]):
            best = (score, row_number, found)
    return (best[1], best[2]) if best else (None, {})


def _detect_headers(rows):
    """Find every table header so repeated blocks on one sheet are imported."""
    headers = []
    for row_number, row in enumerate(rows, 1):
        found = {}
        for column, value in enumerate(row[:80], 1):
            kind = _header_kind(value)
            if kind and kind not in found:
                found[kind] = column
        if "indicator" in found and ("value" in found or len(found) >= 2):
            headers.append((row_number, found))
    if not headers:
        row_number, columns = _detect_header(rows)
        return [(row_number, columns)] if row_number else []
    return headers


def _period_terms(period):
    if not period:
        return set()
    terms = {normalize(period.name), str(period.year)}
    if period.month:
        month_names = ["", "january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"]
        month = month_names[period.month]
        terms.update({month, f"{month} {period.year}", f"{period.year} {period.month:02d}", f"{period.month:02d} {period.year}"})
    return {term for term in terms if term}


def _value_columns(rows, header_row, columns, *, end_row=None, period=None):
    if columns.get("value"):
        return [columns["value"]]
    metadata = set(columns.values())
    candidates = []
    max_column = max((len(row) for row in rows), default=0)
    end_row = end_row or len(rows)
    period_terms = _period_terms(period)
    for column in range(1, max_column + 1):
        if column in metadata:
            continue
        header = _cell(rows, header_row, column)
        populated = sum(
            1 for row_number in range(header_row + 1, end_row + 1)
            if _cell(rows, row_number, column) not in (None, "")
        )
        if populated:
            nearby_headers = " ".join(normalize(_cell(rows, row_number, column)) for row_number in range(max(1, header_row - 2), header_row + 1))
            period_score = max((_similarity(term, nearby_headers) for term in period_terms), default=0.0)
            candidates.append((column, header, populated, period_score))
    if not candidates:
        return []
    # Historical workbooks commonly contain many dated columns. Import the
    # rightmost populated period rather than treating every historic period as
    # a separate current response.
    period_matches = [item for item in candidates if item[3] >= 0.72]
    if period_matches:
        period_matches.sort(key=lambda item: (-item[3], -item[2], item[0]))
        return [period_matches[0][0]]
    dated = [item for item in candidates if isinstance(item[1], (date, datetime)) or re.search(r"\b(19|20)\d{2}\b", str(item[1] or ""))]
    return [dated[-1][0] if dated else max(candidates, key=lambda item: (item[2], item[0]))[0]]


def extract_candidates(path, *, period=None):
    # Read-only row streaming avoids materialising the workbook's extensive
    # styles twice. We still open formula and data-only views separately so a
    # formula is never executed and only an existing cached result is eligible.
    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    candidates, warnings = [], []
    layout = []
    try:
        for sheet in formulas.worksheets:
            if sheet.sheet_state != "visible":
                continue
            formula_rows = list(sheet.iter_rows(values_only=True))
            value_rows = list(values[sheet.title].iter_rows(values_only=True))
            layout.append({
                "sheet": sheet.title,
                "rows": [[normalize(value) for value in row[:30]] for row in formula_rows[:80]],
            })
            headers = _detect_headers(formula_rows)
            if not headers:
                warnings.append({"sheet": sheet.title, "message": "No indicator header was detected."})
                continue
            seen = set()
            for header_index, (header_row, columns) in enumerate(headers):
                end_row = headers[header_index + 1][0] - 1 if header_index + 1 < len(headers) else len(formula_rows)
                value_columns = _value_columns(formula_rows, header_row, columns, end_row=end_row, period=period)
                if not value_columns:
                    warnings.append({"sheet": sheet.title, "row": header_row, "message": "No populated value column was detected for this indicator block."})
                    continue
                for row_number in range(header_row + 1, end_row + 1):
                    name = _cell(formula_rows, row_number, columns["indicator"])
                    definition = _cell(formula_rows, row_number, columns.get("definition")) or ""
                    code = _cell(formula_rows, row_number, columns.get("code")) or ""
                    data_type = _cell(formula_rows, row_number, columns.get("data_type")) or ""
                    if not normalize(name) and not normalize(code):
                        continue
                    for value_column in value_columns:
                        value = _cell(value_rows, row_number, value_column)
                        formula_value = _cell(formula_rows, row_number, value_column)
                        if value in (None, "") and not (isinstance(formula_value, str) and formula_value.startswith("=")):
                            continue
                        locator = f"{sheet.title}!{get_column_letter(value_column)}{row_number}"
                        if locator in seen:
                            continue
                        seen.add(locator)
                        candidates.append({
                            "source_locator": locator,
                            "source_sheet": sheet.title,
                            "source_row": row_number,
                            "source_column": get_column_letter(value_column),
                            "indicator_code": str(code or "").strip(),
                            "indicator_name": str(name or code).strip(),
                            "definition": str(definition or "").strip(),
                            "source_data_type": str(data_type or "").strip(),
                            "raw_value": _json_value(value),
                            "formula_without_cache": isinstance(formula_value, str) and formula_value.startswith("=") and value is None,
                        })
    finally:
        formulas.close()
        values.close()
    return candidates, warnings, _layout_fingerprint(layout)


def _targets(template):
    result = []
    for field in FormField.objects.filter(section__form_template=template).select_related("section").prefetch_related("options"):
        if field.field_type in {"formula", "attachment"}:
            continue
        result.append({
            "target_type": "FIELD", "target_key": f"field:{field.section.section_code}:{field.field_code}".lower(),
            "field": field, "grid": None, "grid_column": None, "grid_row_id": "",
            "code": field.field_code, "name": field.label, "definition": field.help_text,
            "section": f"{field.section.section_code} {field.section.title}", "field_type": field.field_type,
            "options": [option.value for option in field.options.all()],
            "_source_coordinates": {
                (normalize(field.source_sheet), field.source_row)
            } if field.source_sheet and field.source_row else set(),
            "_code_norm": normalize(field.field_code), "_name_norm": normalize(field.label),
            "_name_tokens": _tokens(field.label), "_definition_tokens": _tokens(field.help_text),
            "_section_tokens": _tokens(f"{field.section.section_code} {field.section.title}"),
        })
    grids = FormGrid.objects.filter(section__form_template=template).select_related("section").prefetch_related("columns", "fixed_rows")
    for grid in grids:
        rows = list(grid.fixed_rows.all()) if grid.row_mode == "FIXED" else [None]
        for row in rows:
            for column_index, column in enumerate(grid.columns.all()):
                row_label = row.row_label if row else "Imported row"
                result.append({
                    "target_type": "GRID_CELL",
                    "target_key": f"grid:{grid.section.section_code}:{grid.grid_code}:{row_label}:{column.column_code}".lower(),
                    "field": None, "grid": grid, "grid_column": column, "grid_row_id": str(row.id) if row else "",
                    "code": f"{grid.grid_code} {column.column_code}",
                    "name": f"{grid.title} {row_label} {column.label}", "definition": grid.instructions,
                    "section": f"{grid.section.section_code} {grid.section.title}", "field_type": column.field_type,
                    "options": [],
                    "_source_coordinates": {
                        (normalize(row.source_sheet), int(row.source_rows[column_index]))
                    } if row and row.source_sheet and column_index < len(row.source_rows) and row.source_rows[column_index] else set(),
                    "_code_norm": normalize(f"{grid.grid_code} {column.column_code}"),
                    "_name_norm": normalize(f"{grid.title} {row.row_label} {column.label}"),
                    "_name_tokens": _tokens(f"{grid.title} {row.row_label} {column.label}"),
                    "_definition_tokens": _tokens(grid.instructions),
                    "_section_tokens": _tokens(f"{grid.section.section_code} {grid.section.title}"),
                })
    return result


def _materialize_repeatable_target(target, candidate):
    if not target or target["target_type"] != "GRID_CELL" or target["grid"].row_mode != "REPEATABLE":
        return target
    row_seed = f"{candidate['source_sheet']}:{candidate['source_row']}:{candidate['indicator_name']}"
    row_id = "excel-" + hashlib.sha256(row_seed.encode("utf-8")).hexdigest()[:24]
    materialized = {**target, "grid_row_id": row_id}
    materialized["target_key"] = (
        f"grid:{target['grid'].section.section_code}:{target['grid'].grid_code}:"
        f"{row_id}:{target['grid_column'].column_code}"
    ).lower()
    return materialized


def _convert(value, target):
    if value is None or value == "":
        return "", None
    kind = target["field_type"]
    try:
        if kind in NUMERIC_TYPES:
            text = str(value).strip().replace(",", "").replace("%", "")
            number = Decimal(text)
            return format(number, "f"), None
        if kind == "date":
            if isinstance(value, (date, datetime)):
                return value.date().isoformat() if isinstance(value, datetime) else value.isoformat(), None
            return date.fromisoformat(str(value).strip()).isoformat(), None
        if kind in {"boolean", "declaration"}:
            normalized = normalize(value)
            if normalized in {"yes", "true", "1", "checked"}:
                return "true", None
            if normalized in {"no", "false", "0", "unchecked"}:
                return "false", None
            return "", "Expected a Yes/No value."
        if kind == "select" and target["options"] and str(value) not in target["options"]:
            normalized_options = {normalize(option): option for option in target["options"]}
            if normalize(value) not in normalized_options:
                return "", "Value is not an allowed option."
            return normalized_options[normalize(value)], None
        return str(value), None
    except (InvalidOperation, ValueError, TypeError):
        return "", f"Value is not valid for {kind}."


def _rank(candidate, target):
    code_exact = bool(candidate["indicator_code"] and normalize(candidate["indicator_code"]) == normalize(target["code"]))
    name = _similarity(candidate["indicator_name"], target["name"])
    definition = _similarity(candidate["definition"], target["definition"])
    section = _similarity(candidate["source_sheet"], target["section"])
    type_score = 1.0 if normalize(candidate["source_data_type"]) in {"", normalize(target["field_type"])} else 0.0
    score = 1.0 if code_exact else (0.58 * name + 0.18 * definition + 0.16 * section + 0.08 * type_score)
    return score, {"code_exact": code_exact, "name": round(name, 4), "definition": round(definition, 4), "section": round(section, 4), "type": type_score}


def _shortlist(candidate, targets, limit=16):
    """Cheaply narrow fuzzy candidates before applying SequenceMatcher.

    Large regulatory workbooks can contain thousands of populated historical
    rows. Running three edit-distance comparisons against every form target is
    unnecessarily quadratic. Exact code/name evidence remains authoritative;
    otherwise token overlap chooses a deterministic small candidate set for the
    full weighted scorer.
    """
    source_coordinate = (normalize(candidate["source_sheet"]), candidate["source_row"])
    exact_coordinate = [
        target for target in targets
        if source_coordinate in target.get("_source_coordinates", set())
    ]
    if exact_coordinate:
        return exact_coordinate
    code_norm = normalize(candidate["indicator_code"])
    name_norm = normalize(candidate["indicator_name"])
    exact_code = [target for target in targets if code_norm and target["_code_norm"] == code_norm]
    if exact_code:
        return exact_code
    exact_name = [target for target in targets if name_norm and target["_name_norm"] == name_norm]
    if exact_name:
        return exact_name
    name_tokens = _tokens(candidate["indicator_name"])
    definition_tokens = _tokens(candidate["definition"])
    section_tokens = _tokens(candidate["source_sheet"])

    def overlap(left, right):
        return len(left & right) / max(len(left | right), 1)

    scored = []
    for target in targets:
        cheap = (
            0.62 * overlap(name_tokens, target["_name_tokens"])
            + 0.20 * overlap(definition_tokens, target["_definition_tokens"])
            + 0.18 * overlap(section_tokens, target["_section_tokens"])
        )
        scored.append((cheap, target["target_key"], target))
    scored.sort(key=lambda row: (-row[0], row[1]))
    return [row[2] for row in scored[:limit]]


def parse_and_match(excel_import):
    path = Path(settings.PRIVATE_UPLOAD_ROOT) / excel_import.backup.storage_path
    candidates, warnings, fingerprint = extract_candidates(path, period=excel_import.submission.expected.period)
    targets = _targets(excel_import.submission.expected.form_template)
    existing_fields = {value.field_id: value.value for value in excel_import.submission.values.all() if value.field_id}
    existing_grids = {(value.grid_id, value.grid_row_id, value.grid_column_id): value.value for value in excel_import.submission.values.all() if value.grid_id}
    profile = ExcelImportMappingProfile.objects.filter(
        provider=excel_import.submission.expected.provider,
        form_template=excel_import.submission.expected.form_template,
        layout_fingerprint=fingerprint,
        imports__parser_version=PARSER_VERSION,
        imports__matcher_version=MATCHER_VERSION,
    ).order_by("-version").first()
    profile_map = profile.mappings if profile else {}
    used_targets = set()
    has_template_coordinates = any(target.get("_source_coordinates") for target in targets)
    rows = []
    for candidate in candidates:
        chosen = None
        reused_key = profile_map.get(candidate["source_locator"])
        if reused_key:
            chosen = next((target for target in targets if target["target_key"] == reused_key), None)
        source_coordinate = (normalize(candidate["source_sheet"]), candidate["source_row"])
        coordinate_targets = [
            target for target in targets
            if source_coordinate in target.get("_source_coordinates", set())
        ]
        if chosen:
            ranked = [(1.0, chosen, {"match_reason": "saved_mapping_profile", "mapping_profile": profile.id})]
        elif len(coordinate_targets) == 1:
            ranked = [(1.0, coordinate_targets[0], {"match_reason": "form_template_source_coordinate"})]
        else:
            ranked = []
            for target in _shortlist(candidate, targets):
                score, evidence = _rank(candidate, target)
                evidence["match_reason"] = "exact_indicator_code" if evidence["code_exact"] else ("exact_normalized_name" if normalize(candidate["indicator_name"]) == target["_name_norm"] else "confidence_scored_match")
                ranked.append((score, target, evidence))
            ranked.sort(key=lambda item: (-item[0], item[1]["target_key"]))
        ranked = [(score, _materialize_repeatable_target(target, candidate), evidence) for score, target, evidence in ranked]
        best_score, best, evidence = ranked[0] if ranked else (0.0, None, {})
        second_score = ranked[1][0] if len(ranked) > 1 else 0.0
        status = "UNMATCHED"
        if candidate["formula_without_cache"]:
            status = "INVALID"
            evidence = {**evidence, "error": "Formula cell has no cached result."}
        elif best and (reused_key or len(coordinate_targets) == 1 or best_score >= 0.82) and best_score - second_score >= 0.05:
            status = "MATCHED"
        elif best and best_score >= 0.60:
            status = "AMBIGUOUS"
        if best and best["target_key"] in used_targets and status == "MATCHED":
            status = "DUPLICATE"
        if status == "UNMATCHED" and has_template_coordinates:
            # Populated workbook rows outside the source coordinates used to
            # build this form are headings, notes, or foreign rows—not form
            # indicators. Keep them out of the provider's unresolved list.
            status = "SKIPPED"
            evidence = {**evidence, "match_reason": "not_an_indicator_in_form_template"}
        converted, conversion_error = _convert(candidate["raw_value"], best) if best else ("", None)
        if conversion_error:
            status = "INVALID"
            evidence = {**evidence, "error": conversion_error}
        current = ""
        if best:
            current = existing_fields.get(best["field"].id, "") if best["field"] else existing_grids.get((best["grid"].id, best["grid_row_id"], best["grid_column"].id), "")
            if status == "MATCHED":
                used_targets.add(best["target_key"])
        rows.append(SubmissionExcelImportMatch(
            excel_import=excel_import, **{key: candidate[key] for key in (
                "source_locator", "source_sheet", "source_row", "source_column", "indicator_code",
                "indicator_name", "definition", "source_data_type", "raw_value",
            )},
            converted_value=converted, status=status, score=best_score, evidence={**evidence, "runner_up": round(second_score, 4)},
            target_type=best["target_type"] if best else "", target_key=best["target_key"] if best else "",
            field=best["field"] if best else None, grid=best["grid"] if best else None,
            grid_row_id=best["grid_row_id"] if best else "", grid_column=best["grid_column"] if best else None,
            current_value=current, will_overwrite=bool(current and converted != current),
        ))
    with transaction.atomic():
        excel_import.matches.all().delete()
        SubmissionExcelImportMatch.objects.bulk_create(rows)
        counts = {status: 0 for status, _ in SubmissionExcelImportMatch.STATUS_CHOICES}
        for row in rows:
            counts[row.status] += 1
        matched_target_keys = {row.target_key for row in rows if row.status == "MATCHED"}
        missing_targets = [{"target_key": target["target_key"], "label": target["name"], "required": bool(target.get("field") and target["field"].is_required)} for target in targets if target["target_key"] not in matched_target_keys]
        missing_required = FormField.objects.filter(section__form_template=excel_import.submission.expected.form_template, is_required=True).exclude(
            id__in=[row.field_id for row in rows if row.status == "MATCHED" and row.field_id]
        ).count()
        unresolved_populated = [{"source_locator": row.source_locator, "indicator": row.indicator_name, "status": row.status} for row in rows if row.raw_value not in (None, "") and row.status not in {"MATCHED", "SKIPPED"}]
        excel_import.layout_fingerprint = fingerprint
        excel_import.mapping_profile = profile
        excel_import.parser_version = PARSER_VERSION
        excel_import.matcher_version = MATCHER_VERSION
        excel_import.summary = {"detected": len(rows), "counts": counts, "missing_required": missing_required, "missing_targets": missing_targets, "unresolved_populated": unresolved_populated, "warnings": warnings}
        excel_import.status = "READY"
        excel_import.errors = []
        excel_import.save(update_fields=["parser_version", "matcher_version", "layout_fingerprint", "mapping_profile", "summary", "status", "errors", "updated_at"])
    return excel_import


def confirm_import(excel_import, user, *, confirmation_key, overwrite_ids, allow_unresolved):
    from apps.submissions.provider_workspace import provider_can_edit
    with transaction.atomic():
        excel_import = SubmissionExcelImport.objects.select_for_update().select_related(
            "submission__expected__form_template", "submission__expected__provider",
        ).get(pk=excel_import.pk)
        submission = type(excel_import.submission).objects.select_for_update().get(pk=excel_import.submission_id)
        if excel_import.status == "IMPORTED":
            if str(excel_import.confirmation_key) != str(confirmation_key):
                raise ValueError("This import was already confirmed with another confirmation key.")
            return excel_import, True
        if excel_import.status != "READY":
            raise ValueError("This import is not ready for confirmation.")
        if not provider_can_edit(user, submission):
            raise PermissionError("Your provider role cannot import values at this workflow stage.")
        if submission.revision != excel_import.source_revision:
            raise RuntimeError("The form changed after this preview was created. Reparse before importing.")
        unresolved = excel_import.matches.exclude(status__in=["MATCHED", "SKIPPED"])
        if unresolved.exists() and not allow_unresolved:
            raise ValueError("Resolve or explicitly allow the remaining unmatched indicators before importing.")
        overwrite_ids = {int(value) for value in overwrite_ids}
        matched = list(excel_import.matches.filter(status="MATCHED").select_related("field", "grid", "grid_column"))
        blocked = [row.id for row in matched if row.will_overwrite and row.id not in overwrite_ids]
        if blocked:
            raise ValueError(f"Confirm every replacement before importing. Unconfirmed match IDs: {blocked}")
        changes, profile_map = [], {}
        for row in matched:
            lookup = {"submission": submission}
            if row.field_id:
                lookup["field_id"] = row.field_id
            else:
                lookup.update({"grid_id": row.grid_id, "grid_row_id": row.grid_row_id, "grid_column_id": row.grid_column_id})
            existing = SubmissionValue.objects.filter(**lookup).first()
            before = {"value": existing.value, "value_status": existing.value_status, "explanation": existing.explanation} if existing else {}
            obj, _ = SubmissionValue.objects.update_or_create(**lookup, defaults={
                "value": row.converted_value, "value_status": "PROVIDED" if row.converted_value != "" else "MISSING",
                "explanation": "", "value_source": "EXCEL_IMPORT",
                "source_reference": f"excel-import:{excel_import.id}", "updated_by": user,
            })
            after = {"value": obj.value, "value_status": obj.value_status, "explanation": obj.explanation}
            if before != after:
                changes.append({"target_type": row.target_type, "target_id": str(row.field_id or f"{row.grid_id}:{row.grid_row_id}:{row.grid_column_id}"), "before": before, "after": after})
            row.user_confirmed = True
            row.save(update_fields=["user_confirmed"])
            profile_map[row.source_locator] = row.target_key
        base_revision = submission.revision
        resulting_revision = complete_submission_revision(submission, user) if changes else submission.revision
        if user.role == "PROVIDER_APPROVER" and changes:
            canonical = json.dumps(changes, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            batch = ProviderEditBatch.objects.create(
                submission=submission, actor=user, stage=submission.expected.workflow_status,
                section_code="EXCEL_IMPORT", base_revision=base_revision, resulting_revision=resulting_revision,
                item_count=len(changes), changes_sha256=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
            )
            ProviderEditItem.objects.bulk_create([ProviderEditItem(batch=batch, **change) for change in changes])
        profile = ExcelImportMappingProfile.objects.filter(
            provider=submission.expected.provider, form_template=submission.expected.form_template,
            layout_fingerprint=excel_import.layout_fingerprint,
        ).order_by("-version").first()
        if not profile or profile.mappings != profile_map:
            profile = ExcelImportMappingProfile.objects.create(
                provider=submission.expected.provider, form_template=submission.expected.form_template,
                layout_fingerprint=excel_import.layout_fingerprint,
                version=(profile.version + 1) if profile else 1, mappings=profile_map, created_by=user,
            )
        excel_import.status = "IMPORTED"
        excel_import.confirmed_by = user
        excel_import.confirmed_at = timezone.now()
        excel_import.confirmation_key = confirmation_key
        excel_import.resulting_revision = resulting_revision
        excel_import.mapping_profile = profile
        unresolved_rows = list(excel_import.matches.exclude(status__in=["MATCHED", "SKIPPED"]).values(
            "id", "source_locator", "indicator_name", "status",
        ))
        skipped_rows = list(excel_import.matches.filter(status="SKIPPED").values(
            "id", "source_locator", "indicator_name",
        ))
        excel_import.imported_manifest = {
            "match_ids": [row.id for row in matched],
            "changed_targets": len(changes),
            "unresolved_populated": unresolved_rows,
            "skipped_populated": skipped_rows,
            "reconciled": not unresolved_rows,
        }
        excel_import.save(update_fields=[
            "status", "confirmed_by", "confirmed_at", "confirmation_key", "resulting_revision",
            "mapping_profile", "imported_manifest", "updated_at",
        ])
        refresh_submission_completion(submission)
        return excel_import, False
