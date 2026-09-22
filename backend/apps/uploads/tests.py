import hashlib
import io
import uuid
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.test import override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APITestCase

from apps.forms_engine.models import FormField, FormSection, FormTemplate
from apps.providers.models import ProviderProfile
from apps.submissions.models import ExpectedSubmission, ReportingPeriod, Submission, SubmissionValue
from apps.users.models import Organization, User

from .excel_imports import confirm_import, parse_and_match
from .models import SubmissionExcelBackup, SubmissionExcelImport


class SubmissionExcelImportTests(APITestCase):
    def setUp(self):
        self.organization = Organization.objects.create(name="Import Provider", org_type="PROVIDER")
        self.provider = ProviderProfile.objects.create(
            organization=self.organization, provider_code="IMP", registered_name="Import Provider",
            sector="TELECOM", category="ISP", licence_type="ISP", licence_number="IMP-1",
            primary_email="provider@example.test", primary_phone="1",
        )
        self.entry = User.objects.create_user(
            "import-entry@example.test", "StrongPassword123!", name="Import Entry",
            role="PROVIDER_DATA_ENTRY", organization=self.organization,
        )
        self.admin = User.objects.create_user(
            "import-admin@nca.test", "StrongPassword123!", name="Import Admin", role="NCA_ADMIN",
        )
        self.form = FormTemplate.objects.create(
            form_code="IMPORT-TEST", name="Import Test", sector="TELECOM", provider_category="ISP",
            frequency="MONTHLY", effective_from=timezone.localdate(), status="ACTIVE",
        )
        self.section = FormSection.objects.create(form_template=self.form, section_code="service", title="Service")
        self.numeric = FormField.objects.create(
            section=self.section, field_code="subscribers", label="Active Subscribers",
            field_type="number", is_required=True,
        )
        self.text = FormField.objects.create(
            section=self.section, field_code="comment", label="Service Comment",
            field_type="text", is_required=False,
        )
        self.period = ReportingPeriod.objects.create(
            name="August 2026", frequency="MONTHLY", year=2026, month=8,
            opens_at=timezone.now() - timedelta(days=1), due_at=timezone.now() + timedelta(days=7),
            status="ACTIVE", created_by=self.admin,
        )
        self.expected = ExpectedSubmission.objects.create(
            provider=self.provider, form_template=self.form, period=self.period,
            workflow_status="DRAFT", due_state="OPEN",
        )
        self.submission = Submission.objects.create(expected=self.expected, version=1)

    def _create_import(self, root):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Subscriber Data"
        sheet.append(["Indicator Code", "Indicator", "Definition", "Data Type", "Data Entry"])
        sheet.append(["subscribers", "Active Subscribers", "Current active base", "number", 1250])
        sheet.append(["", "Unknown indicator", "Not in the form", "number", 12])
        sheet.append(["comment", "Service Comment", "Provider note", "text", "Stable service"])
        sheet.append(["", "Calculated row", "Formula without cached result", "number", "=E2+1"])
        hidden = workbook.create_sheet("Hidden History")
        hidden.sheet_state = "hidden"
        hidden.append(["Indicator", "Definition", "Data Entry"])
        hidden.append(["Active Subscribers", "Old data", 99])
        destination = Path(root) / "excel" / "import.xlsx"
        destination.parent.mkdir(parents=True)
        workbook.save(destination)
        content = destination.read_bytes()
        backup = SubmissionExcelBackup.objects.create(
            submission=self.submission, file_name="import.xlsx", file_size=len(content),
            storage_path="excel/import.xlsx", uploaded_by=self.entry, sha256=hashlib.sha256(content).hexdigest(),
            scan_status="CLEAN",
        )
        return SubmissionExcelImport.objects.create(
            submission=self.submission, backup=backup, uploaded_by=self.entry,
            source_revision=self.submission.revision, status="PARSING",
        )

    def _create_import_from_workbook(self, root, workbook, name="structured.xlsx"):
        destination = Path(root) / "excel" / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        content = destination.read_bytes()
        backup = SubmissionExcelBackup.objects.create(
            submission=self.submission, file_name=name, file_size=len(content),
            storage_path=f"excel/{name}", uploaded_by=self.entry,
            sha256=hashlib.sha256(content).hexdigest(), scan_status="CLEAN",
        )
        return SubmissionExcelImport.objects.create(
            submission=self.submission, backup=backup, uploaded_by=self.entry,
            source_revision=self.submission.revision, status="PARSING",
        )

    def test_parse_preview_and_partial_idempotent_confirmation(self):
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import(root)
            parse_and_match(item)
            item.refresh_from_db()
            self.assertEqual(item.status, "READY")
            self.assertFalse(item.matches.filter(source_sheet="Hidden History").exists())
            self.assertEqual(item.matches.get(indicator_code="subscribers").status, "MATCHED")
            self.assertEqual(item.matches.get(indicator_code="comment").status, "MATCHED")
            self.assertEqual(item.matches.get(indicator_name="Calculated row").status, "INVALID")

            key = uuid.uuid4()
            imported, replayed = confirm_import(
                item, self.entry, confirmation_key=key, overwrite_ids=[], allow_unresolved=True,
            )
            self.assertFalse(replayed)
            self.assertEqual(imported.status, "IMPORTED")
            self.assertEqual(SubmissionValue.objects.get(submission=self.submission, field=self.numeric).value, "1250")
            self.assertEqual(SubmissionValue.objects.get(submission=self.submission, field=self.text).value, "Stable service")

            replay, replayed = confirm_import(
                imported, self.entry, confirmation_key=key, overwrite_ids=[], allow_unresolved=True,
            )
            self.assertTrue(replayed)
            self.assertEqual(replay.resulting_revision, imported.resulting_revision)
            self.assertEqual(SubmissionValue.objects.filter(submission=self.submission).count(), 2)

    def test_provider_cannot_read_another_organizations_import(self):
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import(root)
            other_org = Organization.objects.create(name="Other", org_type="PROVIDER")
            other = User.objects.create_user(
                "other-import@example.test", "StrongPassword123!", name="Other",
                role="PROVIDER_DATA_ENTRY", organization=other_org,
            )
            self.client.force_authenticate(other)
            response = self.client.get(f"/api/v1/excel-imports/{item.id}/")
            self.assertEqual(response.status_code, 404)

    def test_preview_matches_are_paginated(self):
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import(root)
            parse_and_match(item)
            self.client.force_authenticate(self.entry)
            response = self.client.get(f"/api/v1/excel-imports/{item.id}/?page=2&page_size=2")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(response.data["matches"]), 2)
            self.assertEqual(response.data["matches_meta"]["count"], 4)
            self.assertEqual(response.data["matches_meta"]["page"], 2)

    def test_repeated_blocks_and_zero_values_are_not_skipped(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "August Returns"
        sheet.append(["Indicator Code", "Indicator", "August 2026"])
        sheet.append(["subscribers", "Active Subscribers", 0])
        sheet.append([])
        sheet.append(["Indicator Code", "Indicator", "August 2026"])
        sheet.append(["comment", "Service Comment", "No outage"])
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import_from_workbook(root, workbook)
            parse_and_match(item)
            item.refresh_from_db()
            self.assertEqual(item.parser_version, "submission-xlsx-v3")
            self.assertEqual(item.matches.count(), 2)
            self.assertEqual(item.matches.get(indicator_code="subscribers").raw_value, 0)
            self.assertEqual(item.matches.get(indicator_code="subscribers").status, "MATCHED")
            self.assertEqual(item.matches.get(indicator_code="comment").status, "MATCHED")
            self.assertEqual(item.summary["unresolved_populated"], [])

    def test_submission_period_column_wins_over_newer_historical_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Indicator Code", "Indicator", "July 2026", "August 2026", "September 2026"])
        sheet.append(["subscribers", "Active Subscribers", 700, 800, 900])
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import_from_workbook(root, workbook, "periods.xlsx")
            parse_and_match(item)
            match = item.matches.get(indicator_code="subscribers")
            self.assertEqual(match.raw_value, 800)
            self.assertEqual(match.source_column, "D")

    def test_original_template_coordinates_match_changed_labels_and_skip_foreign_rows(self):
        self.numeric.source_sheet = "Subscriber Data"
        self.numeric.source_row = 2
        self.numeric.save(update_fields=["source_sheet", "source_row"])
        self.text.source_sheet = "Subscriber Data"
        self.text.source_row = 4
        self.text.save(update_fields=["source_sheet", "source_row"])
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            item = self._create_import(root)
            parse_and_match(item)
            numeric = item.matches.get(source_row=2)
            foreign = item.matches.get(source_row=3)
            self.assertEqual(numeric.status, "MATCHED")
            self.assertEqual(numeric.field_id, self.numeric.id)
            self.assertEqual(numeric.evidence["match_reason"], "form_template_source_coordinate")
            self.assertEqual(foreign.status, "SKIPPED")
            self.assertFalse(item.matches.filter(status="UNMATCHED").exists())
            item.refresh_from_db()
            self.assertEqual(item.summary["counts"]["UNMATCHED"], 0)
            self.assertEqual(item.summary["unresolved_populated"], [])

    @patch("apps.uploads.views.enqueue_excel_import")
    def test_upload_returns_scanning_before_background_parse(self, enqueue):
        workbook = Workbook()
        workbook.active.append(["Indicator", "Definition", "Data Entry"])
        workbook.active.append(["Active Subscribers", "Current active base", 42])
        payload = io.BytesIO()
        workbook.save(payload)
        self.client.force_authenticate(self.entry)
        with TemporaryDirectory() as root, override_settings(PRIVATE_UPLOAD_ROOT=root):
            response = self.client.post(
                f"/api/v1/submissions/{self.submission.id}/excel-imports/",
                {"file": SimpleUploadedFile(
                    "provider-data.xlsx", payload.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )}, format="multipart",
            )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["status"], "SCANNING")
        enqueue.assert_called_once_with(response.data["id"])
