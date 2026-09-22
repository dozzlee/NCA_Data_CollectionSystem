import csv
import hashlib
import io
import os
import re
import zipfile
from datetime import timedelta
from pathlib import Path
from django.conf import settings
from django.utils import timezone
from apps.forms_engine.models import FormField, GridColumn
from apps.submissions.models import Submission, SubmissionValue
from apps.exports.models import ExportLog
from apps.exports.services import CANONICAL_HEADERS, canonical_row_dicts, canonical_values
from .models import DataRequestArtifact


def _normalized(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _source_layout_workbook(submission, user):
    """Return a populated copy of the original provider form, never a synthesized table."""
    from openpyxl import load_workbook
    from apps.submissions.models import MonthlyReportArtifact
    from apps.submissions.monthly_reports import generate_monthly_report

    existing = MonthlyReportArtifact.objects.filter(submission=submission, status="READY").first()
    if existing and Path(existing.private_path).exists():
        return existing.filename, Path(existing.private_path).read_bytes()
    if submission.expected.form_template.provider_workbook_baselines.filter(status="ACTIVE").exists():
        report = generate_monthly_report(submission.id, actor=user)
        return report.filename, Path(report.private_path).read_bytes()

    workbook_import = getattr(submission.expected.form_template, "workbook_import", None)
    if not workbook_import:
        raise ValueError(f"{submission.expected.form_template.name} has no original workbook template available.")
    source_path = Path(workbook_import.storage_path)
    if not source_path.is_absolute():
        source_path = Path(settings.PRIVATE_UPLOAD_ROOT) / source_path
    if not source_path.exists():
        raise ValueError(f"The original workbook for {submission.expected.form_template.name} is missing.")

    workbook = load_workbook(source_path, data_only=False, keep_links=True)
    values = submission.values.select_related("field__section", "grid", "grid_column", "grid__section")
    grid_rows = {}
    for value in values:
        if value.grid_id:
            grid_rows.setdefault(value.grid_id, [])
            if value.grid_row_id not in grid_rows[value.grid_id]:
                grid_rows[value.grid_id].append(value.grid_row_id)

    def matching_column(sheet, row, label):
        wanted = _normalized(label)
        for column in range(1, sheet.max_column + 1):
            if _normalized(sheet.cell(row, column).value) == wanted:
                return column
        return None

    for value in values:
        if value.value_status != "PROVIDED":
            continue
        if value.field_id and value.field.source_sheet and value.field.source_row:
            sheet = workbook[value.field.source_sheet]
            schema_section = next((section for section in (workbook_import.detected_schema or {}).get("sections", [])
                if section.get("section_code") == value.field.section.section_code), {})
            mapping = schema_section.get("column_mapping") or {}
            header_row = int(mapping.get("header_row") or max(1, value.field.source_row - 1))
            target_column = next((int(cell["column"]) for candidate in mapping.get("candidates", [])
                if int(candidate.get("row") or 0) == header_row for cell in candidate.get("columns", [])
                if "input" in str(cell.get("label", "")).lower() or "value" in str(cell.get("label", "")).lower()), None)
            if target_column is None:
                target_column = matching_column(sheet, header_row, "User Input")
            if target_column:
                sheet.cell(value.field.source_row, target_column).value = value.value
        elif value.grid_id and value.grid_column_id and value.grid.source_sheet and value.grid.source_row:
            sheet = workbook[value.grid.source_sheet]
            target_column = matching_column(sheet, value.grid.source_row, value.grid_column.label)
            if not target_column:
                continue
            fixed_row = value.grid.fixed_rows.filter(row_label=value.grid_row_id).first()
            if fixed_row and fixed_row.source_rows:
                index = list(value.grid.columns.order_by("sort_order", "id")).index(value.grid_column)
                source_rows = list(fixed_row.source_rows)
                target_row = int(source_rows[index] if index < len(source_rows) else source_rows[0])
            else:
                target_row = int(value.grid.source_row) + 1 + grid_rows[value.grid_id].index(value.grid_row_id)
            sheet.cell(target_row, target_column).value = value.value

    out = io.BytesIO(); workbook.save(out); workbook.close()
    provider = submission.expected.provider.trade_name or submission.expected.provider.registered_name
    filename = re.sub(r"[^A-Za-z0-9._-]+", "-", f"{provider}-{submission.submission_reference}").strip("-.") + ".xlsx"
    return filename, out.getvalue()


def eligible_submissions(scope):
    qs = Submission.objects.filter(
        expected__workflow_status="APPROVED",
        expected__form_template_id__in=scope["form_template_ids"],
        expected__period_id__in=scope["period_ids"],
    ).select_related("expected__provider", "expected__form_template", "expected__period")
    mode = scope.get("provider_scope")
    if mode == "SECTOR": qs = qs.filter(expected__provider__sector=scope.get("sector"))
    elif mode == "CATEGORY": qs = qs.filter(expected__provider__category=scope.get("provider_category"))
    elif mode == "SELECTED": qs = qs.filter(expected__provider_id__in=scope.get("provider_ids", []))
    latest = {}
    for submission in qs.order_by("expected_id", "-version"):
        latest.setdefault(submission.expected_id, submission)
    return list(latest.values())


def build_manifest(item):
    submissions = eligible_submissions(item.scope)
    if not submissions:
        raise ValueError("No approved submissions match the selected datasets, providers and periods. Attach a prepared extract to fulfil this request, or ask the requester to change the selection.")
    form_ids = item.scope["form_template_ids"]
    field_ids = item.scope.get("field_ids", [])
    column_ids = item.scope.get("grid_column_ids", [])
    if item.scope.get("all_fields", True):
        field_ids = list(FormField.objects.filter(section__form_template_id__in=form_ids).values_list("id", flat=True))
        column_ids = list(GridColumn.objects.filter(grid__section__form_template_id__in=form_ids).values_list("id", flat=True))
    return {"submission_ids": [s.id for s in submissions], "field_ids": field_ids, "grid_column_ids": column_ids}


def canonical_rows(item):
    manifest = item.approval_manifest or {}
    submissions = Submission.objects.filter(id__in=manifest.get("submission_ids", [])).select_related(
        "expected__provider", "expected__form_template", "expected__period"
    )
    selected_fields = set(manifest.get("field_ids", []))
    selected_columns = set(manifest.get("grid_column_ids", []))
    yield from canonical_values(canonical_row_dicts(
        submissions, selected_field_ids=selected_fields, selected_column_ids=selected_columns,
        request_reference=str(item.id),
    ))


def safe_cell(value):
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _store_scanned_upload(item, upload):
    import uuid
    from apps.uploads.scanner import scan_path
    formats = {".pdf": "application/pdf", ".csv": "text/csv", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    filename = Path(upload.name.replace("\\", "/")).name
    suffix = Path(filename).suffix.lower()
    if suffix not in formats or upload.size == 0 or upload.size > 20 * 1024 * 1024:
        raise ValueError("Attach a non-empty PDF, CSV, XLSX or DOCX file up to 20 MB.")
    root = Path(settings.PRIVATE_EXPORT_ROOT)
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"data-request-{item.id}-{uuid.uuid4().hex}{suffix}"
    try:
        digest = hashlib.sha256()
        with path.open("xb") as destination:
            for chunk in upload.chunks():
                digest.update(chunk)
                destination.write(chunk)
        if scan_path(path)["status"] != "CLEAN":
            raise ValueError("The file did not pass scanning. Please attach another file.")
        return {"path": str(path), "filename": filename, "mime_type": formats[suffix],
            "file_size": upload.size, "sha256": digest.hexdigest()}
    except Exception:
        path.unlink(missing_ok=True)
        raise


def upload_artifact(item, upload, user, enforce_requested_format=True):
    stored = _store_scanned_upload(item, upload)
    suffix = Path(stored["filename"]).suffix.lower()
    required_suffix = {"PDF": ".pdf", "CSV": ".csv", "XLSX": ".xlsx"}[item.requested_format]
    if enforce_requested_format and suffix != required_suffix:
        Path(stored["path"]).unlink(missing_ok=True)
        raise ValueError(f"This request requires a {item.requested_format} file. Attach a {required_suffix} file.")
    return DataRequestArtifact.objects.create(request=item, private_path=stored["path"], filename=stored["filename"],
        mime_type=stored["mime_type"], file_size=stored["file_size"], sha256=stored["sha256"], generated_by=user,
        expires_at=timezone.now() + timedelta(days=30))


def store_delivery_attachment(item, upload):
    return _store_scanned_upload(item, upload)


def generate_artifact(item, user):
    rows = list(canonical_rows(item))
    headers = CANONICAL_HEADERS
    root = Path(settings.PRIVATE_EXPORT_ROOT)
    root.mkdir(parents=True, exist_ok=True)
    storage_stem = f"data-request-{item.id}"
    submission_ids = (item.approval_manifest or {}).get("submission_ids", [])
    references = list(Submission.objects.filter(id__in=submission_ids).values_list("submission_reference", flat=True))
    if len(references) == 1:
        download_stem = references[0]
    else:
        download_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", item.title).strip("-._") or "data-extract"
    if item.requested_format == "CSV":
        filename, mime = download_stem + ".csv", "text/csv"
        stream = io.StringIO(); writer = csv.writer(stream); writer.writerow(headers)
        writer.writerows([[safe_cell(v) for v in row] for row in rows]); content = b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")
    elif item.requested_format == "XLSX":
        submissions = list(Submission.objects.filter(id__in=submission_ids).select_related(
            "expected__provider", "expected__form_template", "expected__period"
        ))
        workbooks = []
        for submission in submissions:
            workbooks.append(_source_layout_workbook(submission, user))
        if len(workbooks) == 1:
            filename, content = workbooks[0]
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                used = set()
                for index, (workbook_name, workbook_content) in enumerate(workbooks, start=1):
                    safe_name = Path(workbook_name).name
                    if safe_name in used:
                        safe_name = f"{index}-{safe_name}"
                    used.add(safe_name)
                    archive.writestr(safe_name, workbook_content)
            content = out.getvalue(); filename, mime = download_stem + ".zip", "application/zip"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        out = io.BytesIO(); doc = SimpleDocTemplate(out, pagesize=landscape(A3), leftMargin=18, rightMargin=18)
        styles = getSampleStyleSheet(); story = [Paragraph(item.title, styles["Title"]), Paragraph(f"Division: {item.requesting_division}", styles["BodyText"]), Spacer(1, 12)]
        width = (landscape(A3)[0] - 36) / len(headers)
        table = Table([headers] + [[safe_cell(v) for v in row] for row in rows], repeatRows=1, colWidths=[width] * len(headers))
        table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#002d5b")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.grey), ("FONTSIZE", (0,0), (-1,-1), 3.5)])); story.append(table); doc.build(story)
        content = out.getvalue(); filename, mime = download_stem + ".pdf", "application/pdf"
    attachment = (item.approval_manifest or {}).get("delivery_attachment")
    if attachment:
        attachment_path = Path(attachment["path"])
        if not attachment_path.exists():
            raise ValueError("The approved delivery attachment is missing.")
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(Path(filename).name, content)
            archive.writestr(Path(attachment["filename"]).name, attachment_path.read_bytes())
        content = out.getvalue(); filename = download_stem + "-delivery.zip"; mime = "application/zip"
    path = root / f"{storage_stem}{Path(filename).suffix}"; path.write_bytes(content)
    artifact, _ = DataRequestArtifact.objects.update_or_create(request=item, defaults={
        "private_path": str(path), "filename": filename, "mime_type": mime, "file_size": len(content), "row_count": len(rows),
        "sha256": hashlib.sha256(content).hexdigest(), "generated_by": user, "expires_at": timezone.now() + timedelta(days=30), "expired_at": None,
    })
    digest = hashlib.sha256(content).hexdigest()
    ExportLog.objects.create(export_type=item.requested_format, data_request_id=str(item.id), filters=item.scope,
        generated_by=user, file_reference=str(path), row_count=len(rows), mime_type=mime, file_size=len(content), sha256=digest)
    return artifact
