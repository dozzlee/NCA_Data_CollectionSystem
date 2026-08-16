from datetime import timedelta
from tempfile import TemporaryDirectory

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.forms_engine.models import (
    FormField,
    FormGrid,
    FormSection,
    FormTemplate,
    GridColumn,
    GridRow,
    KMZUploadRequirement,
    ValidationRule,
)
from apps.providers.models import ProviderProfile, ProviderFormAssignment
from apps.users.models import Organization, User
from apps.uploads.models import SubmissionKMZUpload
from .models import (
    CorrectionItem, ExpectedSubmission, ReportingPeriod, ReviewAction, Submission, SubmissionEvent,
    SubmissionNotification, SubmissionValue, ProviderApprovalDecision, ProviderEditBatch, PeriodFormAssignment,
)


class FormAssignmentTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user("assign-admin@nca.test", "password", name="Admin", role="NCA_ADMIN")
        self.officer = User.objects.create_user("assign-officer@nca.test", "password", name="Officer", role="NCA_OFFICER")
        self.viewer = User.objects.create_user("assign-viewer@nca.test", "password", name="Viewer", role="NCA_VIEWER")
        self.org = Organization.objects.create(name="Assigned Provider", org_type="PROVIDER")
        self.entry = User.objects.create_user("assigned-entry@example.com", "password", name="Entry", role="PROVIDER_DATA_ENTRY", organization=self.org)
        self.approver = User.objects.create_user("assigned-approver@example.com", "password", name="Approver", role="PROVIDER_APPROVER", organization=self.org)
        self.provider = ProviderProfile.objects.create(organization=self.org, registered_name="Assigned Provider", sector="TELECOM", category="ISP", licence_type="ISP", licence_number="A", primary_email="a@example.com", primary_phone="1")
        self.form = FormTemplate.objects.create(form_code="CUSTOM-ISP", name="Custom ISP", sector="TELECOM", provider_category="ISP", frequency="QUARTERLY", effective_from=timezone.localdate(), status="ACTIVE", approval_status="APPROVED", mapping_complete=True)
        self.form.family.canonical_frequency = "QUARTERLY"; self.form.family.frequency_decision_status = "APPROVED"; self.form.family.save()
        self.period = ReportingPeriod.objects.create(name="Q1", frequency="QUARTERLY", year=2026, quarter=1, opens_at=timezone.now()-timedelta(days=1), due_at=timezone.now()+timedelta(days=20), status="ACTIVE", created_by=self.admin)
        self.client.force_authenticate(self.admin)

    def test_manual_assignment_creates_one_obligation_and_notification_idempotently(self):
        payload = {"mode":"MANUAL", "provider_ids":[self.provider.id], "period_id":self.period.id}
        response = self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(PeriodFormAssignment.objects.count(), 1)
        self.assertEqual(ExpectedSubmission.objects.filter(provider=self.provider, form_template=self.form, period=self.period).count(), 1)
        self.assertTrue(SubmissionNotification.objects.filter(recipient=self.entry, title="New form assigned").exists())
        self.assertTrue(SubmissionNotification.objects.filter(recipient=self.approver, title="New form assigned").exists())
        self.assertEqual(response.data["delivery_type"], "IMMEDIATE")
        self.assertEqual(response.data["obligations_created"], 1)
        self.assertEqual(response.data["recurring_schedules_created"], 0)
        self.assertEqual(response.data["obligations"][0]["period_id"], self.period.id)
        expected = ExpectedSubmission.objects.get(provider=self.provider, form_template=self.form, period=self.period)
        self.assertEqual(expected.due_state, "OPEN")
        repeat = self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", payload, format="json")
        self.assertEqual(repeat.status_code, 200, repeat.data)
        self.assertEqual(repeat.data["duplicates"], 1)
        self.assertEqual(repeat.data["obligations_created"], 0)
        self.assertEqual(ExpectedSubmission.objects.count(), 1)

    def test_recurring_assignment_renews_on_period_activation(self):
        response = self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", {"mode":"RECURRING", "provider_ids":[self.provider.id], "effective_from":"2026-01-01"}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["delivery_type"], "SCHEDULED")
        self.assertEqual(response.data["obligations_created"], 0)
        self.assertEqual(response.data["recurring_schedules_created"], 1)
        self.assertFalse(SubmissionNotification.objects.filter(title="New form assigned").exists())
        period = ReportingPeriod.objects.create(name="Q2", frequency="QUARTERLY", year=2026, quarter=2, opens_at=timezone.now()+timedelta(days=30), due_at=timezone.now()+timedelta(days=60), status="DRAFT", created_by=self.admin)
        period.activate()
        expected = ExpectedSubmission.objects.get(period=period, provider=self.provider, form_template=self.form)
        self.assertIsNotNone(expected.recurring_assignment_id)

    def test_mismatch_requires_audited_override_reason(self):
        other_org = Organization.objects.create(name="Broadcaster", org_type="PROVIDER")
        User.objects.create_user("broadcaster-entry@example.com", "password", name="Entry", role="PROVIDER_DATA_ENTRY", organization=other_org)
        User.objects.create_user("broadcaster-approver@example.com", "password", name="Approver", role="PROVIDER_APPROVER", organization=other_org)
        other = ProviderProfile.objects.create(organization=other_org, registered_name="Broadcaster", sector="BROADCASTING", category="PAY_TV", licence_type="TV", licence_number="B", primary_email="b@example.com", primary_phone="2")
        denied = self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", {"mode":"MANUAL", "provider_ids":[other.id], "period_id":self.period.id}, format="json")
        self.assertEqual(denied.status_code, 409)
        allowed = self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", {"mode":"MANUAL", "provider_ids":[other.id], "period_id":self.period.id, "override_reason":"Temporary cross-sector data collection"}, format="json")
        self.assertEqual(allowed.status_code, 201, allowed.data)
        self.assertEqual(PeriodFormAssignment.objects.get(provider=other).mismatch_override_reason, "Temporary cross-sector data collection")

    def test_preview_lists_all_active_providers_with_account_readiness(self):
        incomplete_org = Organization.objects.create(name="Incomplete Provider", org_type="PROVIDER")
        User.objects.create_user("incomplete-entry@example.com", "password", name="Entry", role="PROVIDER_DATA_ENTRY", organization=incomplete_org)
        incomplete = ProviderProfile.objects.create(
            organization=incomplete_org, registered_name="Incomplete Provider", sector="TELECOM", category="ISP",
            licence_type="ISP", licence_number="I", primary_email="i@example.com", primary_phone="3",
        )
        response = self.client.get(
            f"/api/v1/form-templates/{self.form.id}/assignment-preview/",
            {"mode": "MANUAL", "period": self.period.id, "provider_ids": str(incomplete.id)},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertGreaterEqual(len(response.data["providers"]), 2)
        row = next(item for item in response.data["providers"] if item["provider_id"] == incomplete.id)
        self.assertTrue(row["has_data_entry"])
        self.assertFalse(row["has_approver"])
        self.assertFalse(row["ready"])
        self.assertIn("Approver", row["blocking_reason"])

    def test_manual_batch_is_atomic_when_one_provider_is_not_ready(self):
        incomplete_org = Organization.objects.create(name="Not Ready", org_type="PROVIDER")
        User.objects.create_user("not-ready-entry@example.com", "password", name="Entry", role="PROVIDER_DATA_ENTRY", organization=incomplete_org)
        incomplete = ProviderProfile.objects.create(
            organization=incomplete_org, registered_name="Not Ready", sector="TELECOM", category="ISP",
            licence_type="ISP", licence_number="N", primary_email="n@example.com", primary_phone="4",
        )
        response = self.client.post(
            f"/api/v1/form-templates/{self.form.id}/assignments/",
            {"mode": "MANUAL", "provider_ids": [self.provider.id, incomplete.id], "period_id": self.period.id},
            format="json",
        )
        self.assertEqual(response.status_code, 409, response.data)
        self.assertEqual(PeriodFormAssignment.objects.count(), 0)
        self.assertEqual(ExpectedSubmission.objects.count(), 0)

    def test_officer_can_send_but_viewer_and_provider_cannot(self):
        url = f"/api/v1/form-templates/{self.form.id}/assignments/"
        payload = {"mode": "MANUAL", "provider_ids": [self.provider.id], "period_id": self.period.id}
        self.client.force_authenticate(self.officer)
        self.assertEqual(self.client.post(url, payload, format="json").status_code, 201)
        for user in (self.viewer, self.entry):
            self.client.force_authenticate(user)
            self.assertEqual(self.client.post(url, payload, format="json").status_code, 403)

    def test_manual_send_requires_active_matching_period_and_published_template(self):
        draft_period = ReportingPeriod.objects.create(
            name="Draft Q2", frequency="QUARTERLY", year=2026, quarter=2,
            opens_at=timezone.now(), due_at=timezone.now()+timedelta(days=40), status="DRAFT", created_by=self.admin,
        )
        payload = {"mode": "MANUAL", "provider_ids": [self.provider.id], "period_id": draft_period.id}
        self.assertEqual(self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", payload, format="json").status_code, 409)
        self.form.status = "DRAFT"; self.form.approval_status = "DRAFT"; self.form.save(update_fields=["status", "approval_status"])
        payload["period_id"] = self.period.id
        self.assertEqual(self.client.post(f"/api/v1/form-templates/{self.form.id}/assignments/", payload, format="json").status_code, 409)

    def test_complete_send_provider_corrections_and_nca_read_only_review(self):
        section = FormSection.objects.create(
            form_template=self.form, section_code="service_data", title="Service Data", sort_order=1,
        )
        field = FormField.objects.create(
            section=section, field_code="subscriber_total", label="Subscriber total",
            field_type="number", is_required=True, sort_order=1,
        )
        with TemporaryDirectory() as export_root, override_settings(PRIVATE_EXPORT_ROOT=export_root):
            sent = self.client.post(
                f"/api/v1/form-templates/{self.form.id}/assignments/",
                {"mode": "MANUAL", "provider_ids": [self.provider.id], "period_id": self.period.id},
                format="json",
            )
            self.assertEqual(sent.status_code, 201, sent.data)
            expected = ExpectedSubmission.objects.get(pk=sent.data["obligations"][0]["expected_submission_id"])
            submission = expected.versions.get(version=1)

            self.client.force_authenticate(self.approver)
            monitoring = self.client.get("/api/v1/provider-workspace/submissions/?queue=awaiting_data_entry")
            self.assertEqual(monitoring.status_code, 200, monitoring.data)
            monitored = monitoring.data["results"][0]
            self.assertEqual(monitored["id"], expected.id)
            self.assertEqual(monitored["permitted_actions"], ["VIEW"])
            self.assertEqual(monitored["ownership_label"], "Shared Data Entry queue")
            self.assertEqual(monitored["data_entry_team"][0]["email"], self.entry.email)
            self.assertIsNotNone(monitored["sent_at"])
            self.assertEqual(
                self.client.put(
                    f"/api/v1/submissions/{submission.id}/sections/service_data/values/",
                    {"values": [{"field": field.id, "value": "1", "value_status": "PROVIDED"}]},
                    format="json",
                ).status_code,
                403,
            )

            self.client.force_authenticate(self.entry)
            started = self.client.post(f"/api/v1/expected-submissions/{expected.id}/start/", {}, format="json")
            self.assertEqual(started.status_code, 201, started.data)
            saved = self.client.put(
                f"/api/v1/submissions/{submission.id}/sections/service_data/values/",
                {"revision": 0, "values": [{"field": field.id, "value": "10", "value_status": "PROVIDED"}]},
                format="json",
            )
            self.assertEqual(saved.status_code, 200, saved.data)
            handed_off = self.client.post(f"/api/v1/submissions/{submission.id}/submit-for-approval/", {}, format="json")
            self.assertEqual(handed_off.status_code, 200, handed_off.data)

            self.client.force_authenticate(self.approver)
            returned = self.client.post(
                f"/api/v1/submissions/{submission.id}/provider-review/request-correction/",
                {
                    "reason": "Correct the subscriber total.",
                    "targets": [{"type": "FIELD", "id": field.id, "instruction": "Use the signed source."}],
                },
                format="json",
            )
            self.assertEqual(returned.status_code, 200, returned.data)
            self.client.force_authenticate(self.entry)
            corrected = self.client.put(
                f"/api/v1/submissions/{submission.id}/sections/service_data/values/",
                {"revision": 1, "values": [{"field": field.id, "value": "11", "value_status": "PROVIDED"}]},
                format="json",
            )
            self.assertEqual(corrected.status_code, 200, corrected.data)
            resubmitted = self.client.post(
                f"/api/v1/submissions/{submission.id}/provider-review/resubmit/", {}, format="json",
            )
            self.assertEqual(resubmitted.status_code, 200, resubmitted.data)

            self.client.force_authenticate(self.approver)
            approver_edit = self.client.put(
                f"/api/v1/submissions/{submission.id}/sections/service_data/values/",
                {"revision": 2, "values": [{"field": field.id, "value": "12", "value_status": "PROVIDED"}]},
                format="json",
            )
            self.assertEqual(approver_edit.status_code, 200, approver_edit.data)
            official = self.client.post(
                f"/api/v1/submissions/{submission.id}/provider-review/approve/",
                {"attestation": True, "change_summary": "Corrected the subscriber total."}, format="json",
            )
            self.assertEqual(official.status_code, 200, official.data)

            self.client.force_authenticate(self.officer)
            self.assertEqual(self.client.post(f"/api/v1/submissions/{submission.id}/review/start/", {}, format="json").status_code, 200)
            nca_return = self.client.post(
                f"/api/v1/submissions/{submission.id}/review/request-correction/",
                {"comment": "Recheck the subscriber total.", "targets": [{
                    "type": "FIELD", "id": field.id, "comment": "Use the reconciled register.",
                }]}, format="json",
            )
            self.assertEqual(nca_return.status_code, 200, nca_return.data)
            correction_submission = Submission.objects.get(pk=nca_return.data["correction_submission_id"])
            denied_nca_edit = self.client.put(
                f"/api/v1/submissions/{correction_submission.id}/sections/service_data/values/",
                {"values": [{"field": field.id, "value": "13", "value_status": "PROVIDED"}]}, format="json",
            )
            self.assertEqual(denied_nca_edit.status_code, 403)

            self.client.force_authenticate(self.approver)
            provider_review = self.client.get(f"/api/v1/submissions/{correction_submission.id}/provider-review-data/")
            self.assertEqual(provider_review.status_code, 200, provider_review.data)
            self.assertIn("EDIT", provider_review.data["permitted_actions"])
            correction_save = self.client.put(
                f"/api/v1/submissions/{correction_submission.id}/sections/service_data/values/",
                {"revision": 0, "values": [{"field": field.id, "value": "13", "value_status": "PROVIDED"}]},
                format="json",
            )
            self.assertEqual(correction_save.status_code, 200, correction_save.data)
            returned_to_nca = self.client.post(
                f"/api/v1/submissions/{correction_submission.id}/provider-review/approve/",
                {"attestation": True, "change_summary": "Reconciled the subscriber total."}, format="json",
            )
            self.assertEqual(returned_to_nca.status_code, 200, returned_to_nca.data)
            expected.refresh_from_db(); self.assertEqual(expected.workflow_status, "RESUBMITTED")

            self.client.force_authenticate(self.officer)
            review_data = self.client.get(f"/api/v1/submissions/{correction_submission.id}/review-data/")
            self.assertEqual(review_data.status_code, 200, review_data.data)
            self.assertEqual(review_data.data["submission"]["form_template_id"], self.form.id)
            self.assertEqual(self.client.post(f"/api/v1/submissions/{correction_submission.id}/review/start/", {}, format="json").status_code, 200)
            approved = self.client.post(
                f"/api/v1/submissions/{correction_submission.id}/review/approve/", {}, format="json",
            )
            self.assertEqual(approved.status_code, 200, approved.data)
            expected.refresh_from_db(); self.assertEqual(expected.workflow_status, "APPROVED")


class SubmissionRemediationTests(APITestCase):
    def setUp(self):
        self.org_a = Organization.objects.create(name="Provider A", org_type="PROVIDER")
        self.org_b = Organization.objects.create(name="Provider B", org_type="PROVIDER")
        self.entry_a = User.objects.create_user(
            "entry-a@example.com", "password", name="Entry A",
            role="PROVIDER_DATA_ENTRY", organization=self.org_a,
        )
        self.approver_a = User.objects.create_user(
            "approver-a@example.com", "password", name="Approver A",
            role="PROVIDER_APPROVER", organization=self.org_a,
        )
        self.entry_b = User.objects.create_user(
            "entry-b@example.com", "password", name="Entry B",
            role="PROVIDER_DATA_ENTRY", organization=self.org_b,
        )
        self.officer = User.objects.create_user(
            "officer@nca.test", "password", name="Officer", role="NCA_OFFICER",
        )
        self.viewer = User.objects.create_user(
            "viewer@nca.test", "password", name="Viewer", role="NCA_VIEWER",
        )
        self.provider_a = ProviderProfile.objects.create(
            organization=self.org_a, registered_name="Provider A", sector="TELECOM",
            category="MNO", licence_type="MNO", licence_number="A-1",
            primary_email="a@example.com", primary_phone="1",
        )
        self.provider_b = ProviderProfile.objects.create(
            organization=self.org_b, registered_name="Provider B", sector="TELECOM",
            category="MNO", licence_type="MNO", licence_number="B-1",
            primary_email="b@example.com", primary_phone="2",
        )
        self.form = FormTemplate.objects.create(
            form_code="MNO-MONTHLY", name="Monthly", sector="TELECOM",
            provider_category="MNO", frequency="MONTHLY",
            effective_from=timezone.localdate(), status="ACTIVE",
        )
        self.section = FormSection.objects.create(
            form_template=self.form, section_code="main", title="Main",
        )
        self.required_field = FormField.objects.create(
            section=self.section, field_code="required", label="Required",
            field_type="text", is_required=True,
        )
        self.period = ReportingPeriod.objects.create(
            name="Test period", frequency="MONTHLY", year=2026, month=7,
            opens_at=timezone.now() - timedelta(days=1),
            due_at=timezone.now() + timedelta(days=7),
            status="ACTIVE", created_by=self.officer,
        )
        self.expected_a = ExpectedSubmission.objects.create(
            provider=self.provider_a, form_template=self.form, period=self.period,
            workflow_status="DRAFT", due_state="OPEN",
        )
        self.expected_b = ExpectedSubmission.objects.create(
            provider=self.provider_b, form_template=self.form, period=self.period,
            workflow_status="DRAFT", due_state="OPEN",
        )
        self.submission_a = Submission.objects.create(expected=self.expected_a)
        self.submission_b = Submission.objects.create(expected=self.expected_b)

    def authenticate(self, user):
        self.client.force_authenticate(user)

    def test_review_data_uses_exact_immutable_template_and_includes_grid_cells(self):
        grid = FormGrid.objects.create(section=self.section, grid_code="traffic", title="Traffic", row_mode="REPEATABLE", min_rows=1)
        column = GridColumn.objects.create(grid=grid, column_code="minutes", label="Minutes", field_type="number", unit="minutes")
        SubmissionValue.objects.create(submission=self.submission_a, grid=grid, grid_row_id="row-1", grid_column=column, value="42", value_status="PROVIDED", updated_by=self.entry_a)
        newer = FormTemplate.objects.create(form_code="MNO-MONTHLY", name="Newer", sector="TELECOM", provider_category="MNO", frequency="MONTHLY", version="2.0", effective_from=timezone.localdate(), status="ACTIVE", mapping_basis="PRD_SECTION_11")
        FormSection.objects.create(form_template=newer, section_code="different", title="Different")

        self.authenticate(self.officer)
        response = self.client.get(f"/api/v1/submissions/{self.submission_a.id}/review-data/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["template"]["id"], self.form.id)
        self.assertEqual(response.data["template"]["sections"][0]["section_code"], "main")
        self.assertEqual(response.data["values"][0]["grid_row_id"], "row-1")
        self.assertIsNotNone(response.data["legacy_warning"])

        self.authenticate(self.viewer)
        self.assertEqual(self.client.get(f"/api/v1/submissions/{self.submission_a.id}/review-data/").status_code, 403)

    def test_provider_cannot_enumerate_another_organization(self):
        self.authenticate(self.entry_a)
        urls = [
            f"/api/v1/expected-submissions/{self.expected_b.id}/",
            f"/api/v1/submissions/{self.submission_b.id}/",
            f"/api/v1/submissions/{self.submission_b.id}/completion/",
            f"/api/v1/submissions/{self.submission_b.id}/sections/main/values/",
            f"/api/v1/submissions/{self.submission_b.id}/kmz-uploads/",
            f"/api/v1/submissions/{self.submission_b.id}/excel-backups/",
            f"/api/v1/submissions/{self.submission_b.id}/review/history/",
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 404)

        mutations = [
            (
                f"/api/v1/expected-submissions/{self.expected_b.id}/start/",
                {},
            ),
            (
                f"/api/v1/submissions/{self.submission_b.id}/submit-for-approval/",
                {},
            ),
        ]
        for url, payload in mutations:
            with self.subTest(url=url):
                self.assertEqual(
                    self.client.post(url, payload, format="json").status_code,
                    404,
                )

        save = self.client.put(
            f"/api/v1/submissions/{self.submission_b.id}/sections/main/values/",
            {"values": []},
            format="json",
        )
        self.assertEqual(save.status_code, 404)

        self.authenticate(self.approver_a)
        for action in ("official-submit", "return-to-draft"):
            with self.subTest(action=action):
                response = self.client.post(
                    f"/api/v1/submissions/{self.submission_b.id}/{action}/",
                    {},
                    format="json",
                )
                self.assertEqual(response.status_code, 404)

    def test_blank_indicators_warn_but_can_move_through_provider_handoffs(self):
        self.authenticate(self.entry_a)
        response = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/submit-for-approval/", {},
            format="json",
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PENDING_APPROVAL")
        completion = self.client.get(f"/api/v1/submissions/{self.submission_a.id}/completion/")
        self.assertTrue(completion.data["can_submit"])
        self.assertEqual(completion.data["missing_indicator_count"], 1)
        self.assertEqual(completion.data["blocking_issues"], [])

        self.authenticate(self.approver_a)
        official = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/provider-review/approve/",
            {"attestation": True}, format="json",
        )
        self.assertEqual(official.status_code, 200, official.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "SUBMITTED")
        self.authenticate(self.officer)
        started = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/start/", {}, format="json",
        )
        self.assertEqual(started.status_code, 200, started.data)
        approved = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/approve/", {"comment": "Reviewed."}, format="json",
        )
        self.assertEqual(approved.status_code, 200, approved.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "APPROVED")

    def test_approver_sees_assignment_metadata_and_shared_data_entry_team(self):
        self.authenticate(self.approver_a)
        response = self.client.get(f"/api/v1/expected-submissions/{self.expected_a.id}/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["ownership_label"], "Shared Data Entry queue")
        self.assertEqual(response.data["form_version"], "1.0")
        self.assertEqual([member["name"] for member in response.data["data_entry_team"]], ["Entry A"])
        self.assertIsNotNone(response.data["sent_at"])

    def test_required_declaration_remains_a_genuine_blocker(self):
        FormField.objects.create(
            section=self.section, field_code="declaration", label="Accuracy declaration",
            field_type="declaration", is_required=True,
        )
        self.authenticate(self.entry_a)
        completion = self.client.get(f"/api/v1/submissions/{self.submission_a.id}/completion/")
        self.assertFalse(completion.data["can_submit"])
        self.assertTrue(any(issue["code"] == "REQUIRED_DECLARATION" for issue in completion.data["blocking_issues"]))

    def test_data_entry_can_start_and_reload_saved_values(self):
        self.submission_a.delete()
        self.expected_a.workflow_status = "NOT_STARTED"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.entry_a)
        started = self.client.post(
            f"/api/v1/expected-submissions/{self.expected_a.id}/start/",
            {},
            format="json",
        )
        self.assertEqual(started.status_code, 201)
        submission_id = started.data["id"]
        saved = self.client.put(
            f"/api/v1/submissions/{submission_id}/sections/main/values/",
            {"values": [{
                "field": self.required_field.id,
                "value": "persisted value",
                "value_status": "PROVIDED",
                "explanation": "",
            }]},
            format="json",
        )
        self.assertEqual(saved.status_code, 200)
        reloaded = self.client.get(
            f"/api/v1/submissions/{submission_id}/sections/main/values/"
        )
        self.assertEqual(reloaded.status_code, 200)
        self.assertEqual(reloaded.data[0]["value"], "persisted value")

    def test_data_entry_can_edit_not_started_and_first_save_starts_draft(self):
        self.expected_a.workflow_status = "NOT_STARTED"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.entry_a)
        detail = self.client.get(f"/api/v1/expected-submissions/{self.expected_a.id}/")
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertIn("EDIT", detail.data["permitted_actions"])
        saved = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"revision": 0, "values": []}, format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "DRAFT")
        self.assertEqual(
            SubmissionEvent.objects.filter(
                submission=self.submission_a, event_type="SUBMISSION_STARTED",
            ).count(),
            1,
        )

    def test_data_entry_can_submit_untouched_not_started_form_with_optional_blanks(self):
        self.expected_a.workflow_status = "NOT_STARTED"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.entry_a)
        submitted = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/submit-for-approval/", {}, format="json",
        )
        self.assertEqual(submitted.status_code, 200, submitted.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PENDING_APPROVAL")
        events = list(
            SubmissionEvent.objects.filter(submission=self.submission_a)
            .order_by("created_at", "id")
            .values_list("event_type", flat=True)
        )
        self.assertEqual(events, ["SUBMISSION_STARTED", "SUBMITTED_FOR_APPROVAL"])

    def test_completed_submission_moves_through_provider_workflow(self):
        self.authenticate(self.entry_a)
        save = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"values": [{
                "field": self.required_field.id,
                "value": "complete",
                "value_status": "PROVIDED",
                "explanation": "",
            }]},
            format="json",
        )
        self.assertEqual(save.status_code, 200)
        self.assertTrue(save.data["can_submit"])
        submit = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/submit-for-approval/", {},
            format="json",
        )
        self.assertEqual(submit.status_code, 200)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PENDING_APPROVAL")
        queue = self.client.get("/api/v1/expected-submissions/?workflow_status=PENDING_APPROVAL")
        self.assertEqual(queue.status_code, 200)
        self.assertEqual({item["id"] for item in queue.data["results"]}, {self.expected_a.id})
        self.assertEqual(self.client.post(f"/api/v1/submissions/{self.submission_a.id}/official-submit/", {}, format="json").status_code, 403)

        self.authenticate(self.entry_b)
        other_queue = self.client.get("/api/v1/expected-submissions/?workflow_status=PENDING_APPROVAL")
        self.assertNotIn(self.expected_a.id, {item["id"] for item in other_queue.data["results"]})

        self.authenticate(self.approver_a)
        official = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/official-submit/", {"attestation": True},
            format="json",
        )
        self.assertEqual(official.status_code, 200)
        self.expected_a.refresh_from_db()
        self.submission_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "SUBMITTED")
        self.assertEqual(self.submission_a.submitted_by, self.approver_a)
        self.assertIsNotNone(self.submission_a.submitted_at)

    def test_non_filled_required_value_needs_explanation(self):
        self.authenticate(self.entry_a)
        url = f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/"
        response = self.client.put(url, {"values": [{
            "field": self.required_field.id,
            "value": "",
            "value_status": "NOT_AVAILABLE",
            "explanation": "",
        }]}, format="json")
        self.assertFalse(response.data["can_submit"])
        response = self.client.put(url, {"values": [{
            "field": self.required_field.id,
            "value": "",
            "value_status": "NOT_AVAILABLE",
            "explanation": "The authoritative source is temporarily unavailable.",
        }]}, format="json")
        self.assertTrue(response.data["can_submit"])

    def test_upload_requires_authentication_and_organization_scope(self):
        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            url = f"/api/v1/submissions/{self.submission_a.id}/excel-backups/upload/"
            unauthenticated = self.client.post(url, {}, format="multipart")
            self.assertEqual(unauthenticated.status_code, 401)
            self.authenticate(self.entry_b)
            cross_tenant = self.client.post(url, {}, format="multipart")
            self.assertEqual(cross_tenant.status_code, 404)
            self.authenticate(self.entry_a)
            upload = self.client.post(
                url,
                {"file": SimpleUploadedFile(
                    "../provider-return.xlsx", b"test workbook content",
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )},
                format="multipart",
            )
            self.assertEqual(upload.status_code, 201)
            self.assertEqual(upload.data["file_name"], "provider-return.xlsx")
            listing = self.client.get(
                f"/api/v1/submissions/{self.submission_a.id}/excel-backups/"
            )
            self.assertEqual(listing.status_code, 200)
            self.assertEqual(listing.data[0]["file_name"], "provider-return.xlsx")

    def test_provider_approver_is_read_only_until_workflow_actions(self):
        self.authenticate(self.approver_a)
        save = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"values": []},
            format="json",
        )
        upload = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/excel-backups/upload/",
            {"file": SimpleUploadedFile("backup.xlsx", b"test")},
            format="multipart",
        )
        start = self.client.post(
            f"/api/v1/expected-submissions/{self.expected_a.id}/start/",
            {},
            format="json",
        )
        self.assertEqual(save.status_code, 403)
        self.assertEqual(upload.status_code, 403)
        self.assertEqual(start.status_code, 403)

    def test_approver_can_return_then_officially_submit(self):
        self.expected_a.workflow_status = "PENDING_APPROVAL"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.approver_a)
        returned = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/return-to-draft/",
            {"reason": "The required section needs correction.", "targets": [{"type": "SECTION", "id": "main"}]},
            format="json",
        )
        self.assertEqual(returned.status_code, 200)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PROVIDER_CHANGES_REQUESTED")

        SubmissionValue.objects.create(submission=self.submission_a, field=self.required_field, value="complete", value_status="PROVIDED", updated_by=self.entry_a)
        CorrectionItem.objects.filter(source_submission=self.submission_a).update(status="ADDRESSED")
        self.expected_a.workflow_status = "PENDING_APPROVAL"
        self.expected_a.save(update_fields=["workflow_status"])
        submitted = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/official-submit/",
            {"attestation": True},
            format="json",
        )
        self.assertEqual(submitted.status_code, 200)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "SUBMITTED")

    def test_provider_correction_requires_reason_and_target_and_notifies_entry(self):
        self.expected_a.workflow_status = "PENDING_APPROVAL"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.approver_a)
        url = f"/api/v1/submissions/{self.submission_a.id}/provider-review/request-correction/"
        self.assertEqual(self.client.post(url, {}, format="json").status_code, 400)
        returned = self.client.post(url, {
            "reason": "Correct the audited value.",
            "targets": [{"type": "FIELD", "id": self.required_field.id, "instruction": "Use the approved source."}],
        }, format="json")
        self.assertEqual(returned.status_code, 200, returned.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PROVIDER_CHANGES_REQUESTED")
        self.assertTrue(CorrectionItem.objects.filter(source_submission=self.submission_a, stage="PROVIDER_APPROVAL").exists())
        self.assertTrue(SubmissionNotification.objects.filter(recipient=self.entry_a, submission=self.submission_a).exists())
        self.assertTrue(SubmissionEvent.objects.filter(submission=self.submission_a, event_type="PROVIDER_CHANGES_REQUESTED").exists())

    def test_nca_correction_clones_official_version_and_returns_through_approver(self):
        SubmissionValue.objects.create(
            submission=self.submission_a, field=self.required_field, value="original official value",
            value_status="PROVIDED", updated_by=self.entry_a,
        )
        self.expected_a.workflow_status = "UNDER_REVIEW"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.officer)
        correction = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/request-correction/",
            {"comment": "Correct the audited field.", "targets": [{
                "type": "FIELD", "id": self.required_field.id, "comment": "Use the signed source."
            }]}, format="json",
        )
        self.assertEqual(correction.status_code, 200, correction.data)
        clone = Submission.objects.get(pk=correction.data["correction_submission_id"])
        self.assertEqual(clone.supersedes, self.submission_a)
        self.assertEqual(clone.values.get(field=self.required_field).value, "original official value")

        self.authenticate(self.entry_a)
        immutable_attempt = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"revision": self.submission_a.revision, "values": [{
                "field": self.required_field.id, "value": "must not overwrite", "value_status": "PROVIDED",
            }]}, format="json",
        )
        self.assertEqual(immutable_attempt.status_code, 409)
        self.assertEqual(self.client.put(
            f"/api/v1/submissions/{clone.id}/sections/main/values/",
            {"revision": clone.revision, "values": [{
                "field": self.required_field.id, "value": "corrected value", "value_status": "PROVIDED",
            }]}, format="json",
        ).status_code, 403)
        self.authenticate(self.approver_a)
        saved = self.client.put(
            f"/api/v1/submissions/{clone.id}/sections/main/values/",
            {"revision": clone.revision, "values": [{
                "field": self.required_field.id, "value": "corrected value", "value_status": "PROVIDED",
            }]}, format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        official = self.client.post(f"/api/v1/submissions/{clone.id}/provider-review/approve/", {
            "attestation": True, "change_summary": "Corrected the NCA-targeted field from the signed source."
        }, format="json")
        self.assertEqual(official.status_code, 200, official.data)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "RESUBMITTED")
        self.submission_a.refresh_from_db()
        self.assertEqual(self.submission_a.values.get(field=self.required_field).value, "original official value")

    def test_official_approval_reruns_readiness(self):
        self.expected_a.workflow_status = "PENDING_APPROVAL"
        self.expected_a.save(update_fields=["workflow_status"])
        SubmissionValue.objects.create(
            submission=self.submission_a, field=self.required_field,
            value="not-a-number", value_status="PROVIDED", updated_by=self.entry_a,
        )
        ValidationRule.objects.create(
            form_template=self.form, field=self.required_field, rule_type="TYPE",
            parameters={"field_type": "number"}, message="Enter a valid number.",
        )
        self.authenticate(self.approver_a)
        response = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/provider-review/approve/",
            {"attestation": True}, format="json",
        )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.data["code"], "SUBMISSION_BLOCKED")
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "PENDING_APPROVAL")

    def test_section_snapshot_deletes_omitted_repeatable_cells_and_rejects_stale_revision(self):
        grid = FormGrid.objects.create(section=self.section, grid_code="rows", title="Rows", row_mode="REPEATABLE")
        column = GridColumn.objects.create(grid=grid, column_code="value", label="Value", field_type="text")
        SubmissionValue.objects.create(submission=self.submission_a, field=self.required_field, value="old", value_status="PROVIDED", updated_by=self.entry_a)
        SubmissionValue.objects.create(submission=self.submission_a, grid=grid, grid_row_id="remove-me", grid_column=column, value="old", value_status="PROVIDED", updated_by=self.entry_a)
        self.authenticate(self.entry_a)
        url = f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/"
        saved = self.client.put(url, {"revision": 0, "values": [{
            "field": self.required_field.id, "value": "new", "value_status": "PROVIDED",
        }]}, format="json")
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertFalse(SubmissionValue.objects.filter(submission=self.submission_a, grid_row_id="remove-me").exists())
        stale = self.client.put(url, {"revision": 0, "values": []}, format="json")
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.data["code"], "STALE_REVISION")

    def test_provider_workspace_counts_every_record_beyond_first_page(self):
        for index in range(55):
            period = ReportingPeriod.objects.create(
                name=f"Queue period {index}", frequency="MONTHLY", year=2027 + index, month=1,
                opens_at=timezone.now() - timedelta(days=1), due_at=timezone.now() + timedelta(days=7),
                status="ACTIVE", created_by=self.officer,
            )
            ExpectedSubmission.objects.create(
                provider=self.provider_a, form_template=self.form, period=period,
                workflow_status="DRAFT", due_state="OPEN",
            )
        self.authenticate(self.entry_a)
        summary = self.client.get("/api/v1/provider-workspace/summary/")
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.data["drafts"], 56)
        page = self.client.get("/api/v1/provider-workspace/submissions/?queue=drafts")
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.data["count"], 56)
        self.assertEqual(len(page.data["results"]), 20)
        expanded = self.client.get("/api/v1/provider-workspace/submissions/?queue=drafts&page_size=100")
        self.assertEqual(len(expanded.data["results"]), 56)

    def test_approver_can_monitor_data_entry_work_without_editing_it(self):
        self.authenticate(self.approver_a)
        summary = self.client.get("/api/v1/provider-workspace/summary/")
        self.assertEqual(summary.status_code, 200)
        self.assertEqual(summary.data["awaiting_data_entry"], 1)
        queue = self.client.get("/api/v1/provider-workspace/submissions/?queue=awaiting_data_entry")
        self.assertEqual(queue.status_code, 200)
        self.assertEqual([item["id"] for item in queue.data["results"]], [self.expected_a.id])
        self.assertEqual(queue.data["results"][0]["permitted_actions"], ["VIEW"])
        denied = self.client.post(f"/api/v1/expected-submissions/{self.expected_a.id}/start/")
        self.assertEqual(denied.status_code, 403)

    def test_approver_edit_is_audited_and_requires_change_summary(self):
        SubmissionValue.objects.create(
            submission=self.submission_a, field=self.required_field, value="entry value",
            value_status="PROVIDED", updated_by=self.entry_a,
        )
        self.expected_a.workflow_status = "PENDING_APPROVAL"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.approver_a)
        saved = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"revision": 0, "values": [{
                "field": self.required_field.id, "value": "approver value", "value_status": "PROVIDED",
            }]}, format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        batch = ProviderEditBatch.objects.get(submission=self.submission_a)
        self.assertEqual(batch.item_count, 1)
        self.assertEqual(batch.items.get().before["value"], "entry value")
        self.assertEqual(batch.items.get().after["value"], "approver value")
        missing_summary = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/provider-review/approve/",
            {"attestation": True}, format="json",
        )
        self.assertEqual(missing_summary.status_code, 400)
        self.assertEqual(missing_summary.data["code"], "CHANGE_SUMMARY_REQUIRED")
        approved = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/provider-review/approve/",
            {"attestation": True, "change_summary": "Corrected the value from the signed source."}, format="json",
        )
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertTrue(ProviderApprovalDecision.objects.filter(submission=self.submission_a, attestation=True).exists())

    def test_completion_returns_open_corrections_without_error(self):
        CorrectionItem.objects.create(
            source_submission=self.submission_a, resolution_submission=self.submission_a,
            stage="PROVIDER_APPROVAL", target_type="FIELD", target_id=str(self.required_field.id),
            instruction="Correct this field.", created_by=self.approver_a,
        )
        self.authenticate(self.entry_a)
        response = self.client.get(f"/api/v1/submissions/{self.submission_a.id}/completion/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["open_correction_item_count"], 1)
        self.assertFalse(response.data["transition_ready"])

    def test_provider_history_hides_internal_nca_notes(self):
        ReviewAction.objects.create(submission=self.submission_a, action="ADD_NOTE", comment="internal", created_by=self.officer)
        ReviewAction.objects.create(submission=self.submission_a, action="ADD_PROVIDER_COMMENT", comment="visible", is_provider_visible=True, created_by=self.officer)
        self.authenticate(self.entry_a)
        history = self.client.get(f"/api/v1/submissions/{self.submission_a.id}/review/history/")
        self.assertEqual(history.status_code, 200)
        self.assertEqual([item["comment"] for item in history.data["results"]], ["visible"])

    def test_correction_can_be_edited_reapproved_and_resubmitted(self):
        self.expected_a.workflow_status = "SUBMITTED"
        self.expected_a.save(update_fields=["workflow_status"])
        self.authenticate(self.officer)
        started = self.client.post(f"/api/v1/submissions/{self.submission_a.id}/review/start/", {}, format="json")
        self.assertEqual(started.status_code, 200)
        correction = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/request-correction/",
            {
                "comment": "Correct the required value.",
                "targets": [{
                    "type": "FIELD",
                    "id": self.required_field.id,
                    "comment": "Use the audited source.",
                }],
            },
            format="json",
        )
        self.assertEqual(correction.status_code, 200)

        self.authenticate(self.entry_a)
        correction_submission_id = correction.data["correction_submission_id"]
        self.authenticate(self.approver_a)
        saved = self.client.put(
            f"/api/v1/submissions/{correction_submission_id}/sections/main/values/",
            {"values": [{
                "field": self.required_field.id,
                "value": "corrected",
                "value_status": "PROVIDED",
                "explanation": "",
            }]},
            format="json",
        )
        self.assertTrue(saved.data["can_submit"])
        resubmitted = self.client.post(
            f"/api/v1/submissions/{correction_submission_id}/official-submit/",
            {"attestation": True, "change_summary": "Corrected the required value."},
            format="json",
        )
        self.assertEqual(resubmitted.status_code, 200)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "RESUBMITTED")

    def test_nca_review_transitions_require_a_submitted_state(self):
        self.authenticate(self.officer)
        actions = ("approve", "reject", "request-correction")
        for action in actions:
            with self.subTest(action=action):
                response = self.client.post(
                    f"/api/v1/submissions/{self.submission_a.id}/review/{action}/",
                    {"comment": "invalid transition"},
                    format="json",
                )
                self.assertEqual(response.status_code, 409)

        self.expected_a.workflow_status = "SUBMITTED"
        self.expected_a.save(update_fields=["workflow_status"])
        SubmissionValue.objects.create(submission=self.submission_a,field=self.required_field,value="complete",value_status="PROVIDED",updated_by=self.entry_a)
        started = self.client.post(f"/api/v1/submissions/{self.submission_a.id}/review/start/", {}, format="json")
        self.assertEqual(started.status_code, 200)
        approved = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/approve/",
            {"comment": "Reviewed."},
            format="json",
        )
        self.assertEqual(approved.status_code, 200)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "APPROVED")

    def test_fixed_and_repeatable_blank_grid_cells_are_completeness_warnings(self):
        fixed = FormGrid.objects.create(
            section=self.section,
            grid_code="fixed",
            title="Fixed Grid",
            row_mode="FIXED",
        )
        fixed_column = GridColumn.objects.create(
            grid=fixed,
            column_code="required",
            label="Required Fixed Value",
            field_type="text",
            is_required=True,
        )
        fixed_row = GridRow.objects.create(grid=fixed, row_label="Row 1")
        repeatable = FormGrid.objects.create(
            section=self.section,
            grid_code="repeatable",
            title="Repeatable Grid",
            row_mode="REPEATABLE",
        )
        required_repeat = GridColumn.objects.create(
            grid=repeatable,
            column_code="required",
            label="Required Repeatable Value",
            field_type="text",
            is_required=True,
        )
        optional_repeat = GridColumn.objects.create(
            grid=repeatable,
            column_code="optional",
            label="Optional Repeatable Value",
            field_type="text",
            is_required=False,
        )
        self.authenticate(self.entry_a)
        response = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"values": [
                {
                    "field": self.required_field.id,
                    "value": "complete",
                    "value_status": "PROVIDED",
                },
                {
                    "grid": fixed.id,
                    "grid_row_id": str(fixed_row.id),
                    "grid_column": fixed_column.id,
                    "value": "complete",
                    "value_status": "PROVIDED",
                },
                {
                    "grid": repeatable.id,
                    "grid_row_id": "row-1",
                    "grid_column": optional_repeat.id,
                    "value": "optional only",
                    "value_status": "PROVIDED",
                },
            ]},
            format="json",
        )
        self.assertTrue(response.data["can_submit"])
        self.assertTrue(any(
            issue["id"] == f"{repeatable.id}:row-1:{required_repeat.id}"
            for issue in response.data["completeness_warnings"]
        ))
        completed = self.client.put(
            f"/api/v1/submissions/{self.submission_a.id}/sections/main/values/",
            {"values": [
                {"field": self.required_field.id, "value": "complete", "value_status": "PROVIDED"},
                {"grid": fixed.id, "grid_row_id": str(fixed_row.id), "grid_column": fixed_column.id, "value": "complete", "value_status": "PROVIDED"},
                {"grid": repeatable.id, "grid_row_id": "row-1", "grid_column": optional_repeat.id, "value": "optional only", "value_status": "PROVIDED"},
                {"grid": repeatable.id, "grid_row_id": "row-1", "grid_column": required_repeat.id, "value": "complete", "value_status": "PROVIDED"},
            ]},
            format="json",
        )
        self.assertTrue(completed.data["can_submit"])
        self.assertLessEqual(
            completed.data["sections"][0]["provided"],
            completed.data["sections"][0]["required"],
        )

    def test_required_kmz_upload_is_exposed_and_blocks_until_uploaded(self):
        fibre_form = FormTemplate.objects.create(
            form_code="DC-DBS05",
            name="Domestic Fibre",
            sector="TELECOM",
            provider_category="MNO",
            frequency="MONTHLY",
            effective_from=timezone.localdate(),
            status="ACTIVE",
            kmz_required=True,
        )
        fibre_section = FormSection.objects.create(
            form_template=fibre_form,
            section_code="network",
            title="Network",
            kmz_upload_required=True,
        )
        requirement = KMZUploadRequirement.objects.create(
            form_template=fibre_form,
            section=fibre_section,
            category="ROUTE",
            description="Route map",
            is_required=True,
        )
        expected = ExpectedSubmission.objects.create(
            provider=self.provider_a,
            form_template=fibre_form,
            period=self.period,
            workflow_status="DRAFT",
            due_state="OPEN",
        )
        submission = Submission.objects.create(expected=expected)
        self.authenticate(self.entry_a)
        detail = self.client.get(f"/api/v1/form-templates/{fibre_form.id}/")
        self.assertEqual(detail.status_code, 200)
        exposed_ids = [
            item["id"]
            for section in detail.data["sections"]
            for item in section["kmz_requirements"]
        ]
        self.assertIn(requirement.id, exposed_ids)

        before = self.client.get(f"/api/v1/submissions/{submission.id}/completion/")
        self.assertFalse(before.data["can_submit"])
        self.assertEqual(before.data["blocking_issues"][0]["code"], "REQUIRED_KMZ_UPLOAD")

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            uploaded = self.client.post(
                f"/api/v1/submissions/{submission.id}/kmz-uploads/",
                {
                    "file": SimpleUploadedFile(
                        "../route.kmz",
                        b"PK test kmz archive",
                        content_type="application/vnd.google-earth.kmz",
                    ),
                    "requirement_id": requirement.id,
                },
                format="multipart",
            )
            self.assertEqual(uploaded.status_code, 201)
            self.assertEqual(uploaded.data["file_name"], "route.kmz")
            self.assertTrue(
                SubmissionKMZUpload.objects.filter(
                    submission=submission,
                    requirement=requirement,
                ).exists()
            )
            after = self.client.get(f"/api/v1/submissions/{submission.id}/completion/")
            self.assertTrue(after.data["can_submit"])

    def test_period_activation_matches_category_sector_and_frequency(self):
        from apps.providers.models import ProviderFormAssignment
        broadcasting = FormTemplate.objects.create(
            form_code="DC-TB02",
            name="Broadcasting",
            sector="BROADCASTING",
            provider_category="MNO",
            frequency="MONTHLY",
            effective_from=timezone.localdate(),
            status="ACTIVE",
        )
        annual = FormTemplate.objects.create(
            form_code="MNO-ANNUAL",
            name="Annual Telecom",
            sector="TELECOM",
            provider_category="MNO",
            frequency="ANNUAL",
            effective_from=timezone.localdate(),
            status="ACTIVE",
        )
        period = ReportingPeriod.objects.create(
            name="Matching period",
            frequency="MONTHLY",
            year=2026,
            month=8,
            opens_at=timezone.now() - timedelta(days=1),
            due_at=timezone.now() + timedelta(days=7),
            status="DRAFT",
            created_by=self.officer,
        )
        period.assigned_providers.add(self.provider_a)
        period.applicable_form_templates.add(self.form, broadcasting, annual)
        for form in (self.form, broadcasting, annual):
            form.mapping_complete = True
            form.approval_status = "APPROVED"
            form.save(update_fields=["mapping_complete", "approval_status"])
        ProviderFormAssignment.objects.create(
            provider=self.provider_a,
            form_family=self.form.family,
            obligation="REQUIRED",
            effective_from=period.opens_at.date(),
            source_reference="Official test register",
            confirmed_by=self.officer,
        )
        period.activate()
        created_forms = set(
            period.expected_submissions.values_list("form_template_id", flat=True)
        )
        self.assertEqual(created_forms, {self.form.id})

    def test_legacy_migration_dry_run_and_commit_preserve_history(self):
        family = self.form.family
        self.form.status = "ARCHIVED"; self.form.save(update_fields=["status"])
        corrected = FormTemplate.objects.create(
            family=family, form_code=self.form.form_code, name="Corrected", sector="TELECOM",
            provider_category="MNO", frequency="MONTHLY", version="3.0",
            effective_from=timezone.localdate(), status="ACTIVE", mapping_basis="PRD_SECTION_11",
        )
        corrected_section = FormSection.objects.create(form_template=corrected, section_code="main", title="Main")
        corrected_field = FormField.objects.create(section=corrected_section, field_code="required", label="Required", field_type="text")
        SubmissionValue.objects.create(submission=self.submission_a, field=self.required_field, value="legacy value", value_status="PROVIDED", updated_by=self.entry_a)
        call_command("migrate_legacy_forms", expected_id=[self.expected_a.id])
        self.expected_a.refresh_from_db(); self.assertIsNone(self.expected_a.replacement_id)
        call_command("migrate_legacy_forms", expected_id=[self.expected_a.id], commit=True)
        self.expected_a.refresh_from_db()
        self.assertEqual(self.expected_a.workflow_status, "ARCHIVED")
        self.assertIsNotNone(self.expected_a.replacement_id)
        migrated = self.expected_a.replacement.versions.get()
        self.assertEqual(migrated.values.get(field=corrected_field).value, "legacy value")
        self.assertEqual(self.submission_a.values.get(field=self.required_field).value, "legacy value")

    def test_data_requester_cannot_read_operations_or_review(self):
        self.authenticate(self.viewer)
        self.assertEqual(self.client.get("/api/v1/dashboard/summary/").status_code, 403)
        self.assertEqual(
            self.client.get("/api/v1/dashboard/charts/submission-trend/").status_code,
            403,
        )
        self.assertEqual(self.client.get(f"/api/v1/periods/{self.period.id}/").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/providers/").status_code, 403)
        expected = self.client.get("/api/v1/expected-submissions/")
        self.assertEqual(expected.status_code, 200)
        self.assertEqual(expected.data["count"], 0)
        self.assertEqual(self.client.get("/api/v1/compliance/flags/").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/data-catalog/").status_code, 200)
        response = self.client.post(
            f"/api/v1/submissions/{self.submission_a.id}/review/approve/",
            {"comment": "must not be accepted"}, format="json",
        )
        self.assertEqual(response.status_code, 403)

    def test_every_authenticated_role_can_export_aggregate_dashboard_xlsx(self):
        for user in [self.entry_a, self.approver_a, self.officer, self.viewer]:
            with self.subTest(role=user.role):
                self.authenticate(user)
                response = self.client.get("/api/v1/industry-dashboard/export/?format=xlsx")
                self.assertEqual(response.status_code, 200, getattr(response, "data", None))
                self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.assertNotIn(b"e70682e229f04e6dbb46da0b7b3004f6.xlsx", response.content)

    def test_restricted_dashboard_replaces_operator_rows_with_industry_totals(self):
        self.authenticate(self.entry_a)
        response = self.client.get("/api/v1/industry-dashboard/data/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["operators"], [])
        self.assertNotIn("sourceWorkbooks", response.data["metadata"])

        chart = next(
            item for item in response.data["charts"]
            if item["id"] == "voice-subs-market-share"
        )
        self.assertEqual([item["name"] for item in chart["series"]], ["Industry total"])
        self.assertEqual(chart["shareSeries"], [])
        q4_2025 = next(
            item["value"] for item in chart["series"][0]["values"]
            if item["period"] == "Q4 2025"
        )
        self.assertEqual(q4_2025, 42871955)

        operator_names = {
            "MTN", "Telecel", "AirtelTigo", "AT", "Tigo", "Airtel", "Glo", "Expresso"
        }
        for item in response.data["charts"]:
            with self.subTest(chart=item["id"]):
                names = {series["name"] for series in item["series"]}
                self.assertTrue(names)
                self.assertTrue(
                    any(
                        observation["value"] is not None
                        for series in item["series"]
                        for observation in series["values"]
                    )
                )
                self.assertTrue(names.isdisjoint(operator_names))

    def test_restricted_dashboard_export_contains_the_calculated_industry_total(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self.authenticate(self.viewer)
        response = self.client.get("/api/v1/industry-dashboard/export/?format=xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        rows = list(workbook["Aggregate Data"].iter_rows(values_only=True))
        self.assertIn(
            (
                "mobile",
                "Mobile Voice Subscriptions and Market Share per Operator",
                "subscriptions",
                "Q4 2025",
                "value",
                "Industry total",
                42871955,
            ),
            rows,
        )
        exported_series = {row[5] for row in rows[1:]}
        self.assertNotIn("MTN", exported_series)
        self.assertNotIn("Telecel", exported_series)

    def test_privileged_dashboard_retains_operator_series(self):
        self.authenticate(self.officer)
        response = self.client.get("/api/v1/industry-dashboard/data/")
        self.assertEqual(response.status_code, 200)
        chart = next(
            item for item in response.data["charts"]
            if item["id"] == "voice-subs-market-share"
        )
        self.assertIn("MTN", [item["name"] for item in chart["series"]])
        self.assertNotIn("Industry total", [item["name"] for item in chart["series"]])

    def test_viewer_mutation_matrix_is_denied(self):
        self.authenticate(self.viewer)
        requests = [
            ("post", "/api/v1/periods/", {"name": "Forbidden"}),
            ("post", "/api/v1/form-templates/", {"name": "Forbidden"}),
            ("post", "/api/v1/auth/users/", {"email": "forbidden@example.com"}),
            ("get", "/api/v1/auth/users/", {}),
            ("get", "/api/v1/audit/", {}),
            ("patch", "/api/v1/compliance/flags/999/status/", {"status": "RESOLVED"}),
            ("post", "/api/v1/compliance/generate-email/", {"expected_submission_ids": []}),
            ("patch", f"/api/v1/submissions/{self.submission_a.id}/kmz-uploads/999/review/", {
                "review_status": "ACCEPTED",
            }),
        ]
        for method, url, payload in requests:
            with self.subTest(method=method, url=url):
                response = getattr(self.client, method)(url, payload, format="json")
                self.assertEqual(response.status_code, 403)
