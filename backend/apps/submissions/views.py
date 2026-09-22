from django.conf import settings
from django.db import transaction
from django.core.serializers.json import DjangoJSONEncoder
from django.db.models import Count, F, OuterRef, Prefetch, Q, Subquery
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils.dateparse import parse_datetime
from django.db.models.functions import TruncMonth
from django.utils import timezone
from datetime import timedelta
import hashlib
import json
import os
from pathlib import Path
import uuid
from decimal import Decimal, InvalidOperation
from rest_framework import generics, status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.pagination import PageNumberPagination

from apps.audit.services import record_audit
from apps.compliance.models import TransactionalOutbox
from apps.forms_engine.models import FormField, FormGrid, GridColumn
from apps.users.permissions import (
    IsNCAUser, IsNCAEditor, IsNCAAdmin,
    IsProviderDataEntry, IsProviderApprover, IsProviderUser,
)
from .access import (
    expected_submissions_for_user,
    submissions_for_user,
    get_expected_submission_for_user,
    get_submission_for_user,
)
from .readiness import calculate_submission_readiness, refresh_submission_completion
from .models import (
    ReportingPeriod, ExpectedSubmission, Submission,
    SubmissionValue, ReviewAction, CorrectionItem, NonFilledDisposition,
    DeadlineChangeRequest, SubmissionOverride, ReminderPolicy,
    SubmissionEvent, SubmissionNotification,
    ProviderEditBatch, ProviderEditItem, ProviderApprovalDecision,
    SectionSaveReceipt,
    ProviderWorkbookBaseline, MonthlyReportArtifact,
)
from .serializers import (
    ReportingPeriodSerializer, ExpectedSubmissionSerializer,
    SubmissionSerializer, FormalSubmissionListSerializer, ProviderFormalTaskListSerializer,
    SubmissionValueSerializer, ReviewActionSerializer,
    SubmissionEventSerializer, SubmissionNotificationSerializer,
    ProviderWorkbookBaselineSerializer, MonthlyReportArtifactSerializer,
)
from .workflow import (
    audit_transition, clone_for_nca_correction, complete_submission_revision,
    lock_submission, mark_matching_corrections_addressed, mark_notification_read,
    emit_submission_event,
)
from .provider_workspace import (
    apply_workspace_queue, filter_workspace_queryset, provider_can_edit,
    permitted_actions, summary_for_user, workspace_queryset,
)
from .monthly_reports import (
    baseline_readiness, previous_month_values, replace_baseline_mappings,
    suggest_exact_baseline_mappings,
)


def write_audit(request, action, entity_type, entity_id, before=None, after=None):
    return record_audit(user=request.user, action=action, entity_type=entity_type, entity_id=entity_id,
        before=before, after=after, ip_address=request.META.get("REMOTE_ADDR"))


def validate_correction_targets(template, targets):
    valid_types = {"SECTION", "FIELD", "GRID_CELL"}
    for target in targets:
        target_type = target.get("type")
        target_id = str(target.get("id", "")).strip()
        if target_type not in valid_types or not target_id:
            return "Every correction target requires a valid type and id."
        if target_type == "SECTION" and not template.sections.filter(section_code=target_id).exists():
            return f"Unknown section correction target: {target_id}."
        if target_type == "FIELD" and not template.sections.filter(fields__id=target_id).exists():
            return f"Unknown field correction target: {target_id}."
        if target_type == "GRID_CELL":
            parts = target_id.split(":")
            if len(parts) != 3 or not parts[1] or not template.sections.filter(
                grids__id=parts[0], grids__columns__id=parts[2]
            ).exists():
                return f"Unknown grid-cell correction target: {target_id}."
    return None


def correction_target_label(template, target_type, target_id):
    """Resolve stored correction identifiers without exposing raw database IDs."""
    target_id = str(target_id or "")
    if target_type == "SECTION":
        section = template.sections.filter(section_code=target_id).first()
        return f"Section: {section.title}" if section else "Section"
    if target_type == "FIELD":
        field = FormField.objects.filter(section__form_template=template, pk=target_id).first()
        return f"Indicator: {field.label}" if field else "Indicator"
    if target_type == "GRID_CELL":
        parts = target_id.split(":")
        if len(parts) == 3:
            grid_id, row_id, column_id = parts
            grid = FormGrid.objects.filter(section__form_template=template, pk=grid_id).first()
            column = GridColumn.objects.filter(grid=grid, pk=column_id).first() if grid else None
            row = grid.fixed_rows.filter(pk=row_id).first() if grid and grid.row_mode == "FIXED" else None
            row_label = row.row_label if row else row_id
            if grid and column:
                return f"Indicator: {grid.title} / {row_label} / {column.label}"
        return "Grid indicator"
    return "Submission"


def dashboard_queryset(request):
    # Snapshot-backed obligations are immutable history after a form-catalogue
    # reset. They must not re-enter operational dashboard totals or dereference
    # a template that no longer exists.
    qs = ExpectedSubmission.objects.filter(form_template__isnull=False)
    filters = {
        "period_id": "period_id", "form_id": "form_template_id", "provider_category": "provider__category",
        "sector": "provider__sector", "assigned_officer": "assigned_officer_id", "due_state": "due_state",
        "workflow_status": "workflow_status",
    }
    for parameter, lookup in filters.items():
        if value := request.query_params.get(parameter): qs = qs.filter(**{lookup: value})
    return qs


class ProviderWorkspacePagination(PageNumberPagination):
    page_size = 20
    page_size_query_param = "page_size"
    max_page_size = 100


# ── Dashboard ────────────────────────────────────────────────────────────────

class DashboardSummaryView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        qs = dashboard_queryset(request)
        total = qs.count()
        approved = qs.filter(workflow_status="APPROVED").count()
        expected_ids = list(qs.values_list("id", flat=True))
        latest_submissions = [item.versions.order_by("-version").first() for item in qs.prefetch_related("versions")]
        latest_submissions = [item for item in latest_submissions if item]
        required_slots = sum(item.expected.form_template.sections.filter(fields__is_required=True).count() for item in latest_submissions)
        completed_required = sum(item.values.filter(field__is_required=True, value_status__in=["PROVIDED","SYSTEM_CALCULATED","NOT_APPLICABLE","NOT_AVAILABLE","NOT_REQUIRED"]).count() for item in latest_submissions)
        kmz_required = sum(item.expected.form_template.kmz_requirements.filter(is_required=True).count() for item in latest_submissions)
        kmz_clean = sum(item.kmz_uploads.filter(scan_status="CLEAN").values("requirement_id").distinct().count() for item in latest_submissions)
        from apps.compliance.models import EmailLog
        reminders = EmailLog.objects.filter(expected_submission_id__in=expected_ids, compliance_stage__in=["REMINDER","OVERDUE"])
        return Response({
            "total_expected": total,
            "not_started": qs.filter(workflow_status="NOT_STARTED").count(),
            "draft": qs.filter(workflow_status="DRAFT").count(),
            "pending_approval": qs.filter(workflow_status="PENDING_APPROVAL").count(),
            "submitted": qs.filter(workflow_status="SUBMITTED").count(),
            "under_review": qs.filter(workflow_status="UNDER_REVIEW").count(),
            "correction_requested": qs.filter(workflow_status="CORRECTION_REQUESTED").count(),
            "resubmitted": qs.filter(workflow_status="RESUBMITTED").count(),
            "approved": approved,
            "rejected": qs.filter(workflow_status="REJECTED").count(),
            "overdue": qs.filter(due_state="OVERDUE").count(),
            "due_soon": qs.filter(due_state="DUE_SOON").count(),
            "completion_pct": round((approved / total) * 100, 1) if total else 0,
            "required_value_slots": required_slots,
            "required_values_complete": completed_required,
            "required_values_missing": max(required_slots - completed_required, 0),
            "required_kmz_slots": kmz_required,
            "clean_kmz_uploads": kmz_clean,
            "required_kmz_missing": max(kmz_required - kmz_clean, 0),
            "reminders_draft": reminders.filter(status="DRAFT").count(),
            "reminders_queued": reminders.filter(status="QUEUED").count(),
            "reminders_delivered": reminders.filter(status="DELIVERED").count(),
        })


class MetricCatalogueView(APIView):
    permission_classes = [IsNCAUser]
    def get(self, request):
        return Response([
            {"code":"SUBMISSION_COMPLETION","label":"Approved submissions","unit":"percentage","source":"ExpectedSubmission.workflow_status"},
            {"code":"REQUIRED_VALUES","label":"Required value completeness","unit":"count","source":"FormField/SubmissionValue"},
            {"code":"KMZ_COMPLETENESS","label":"Clean required KMZ uploads","unit":"count","source":"KMZUploadRequirement/SubmissionKMZUpload"},
            {"code":"REMINDER_LIFECYCLE","label":"Reminder queue and delivery","unit":"count","source":"EmailLog"},
            {"code":"DUE_STATE","label":"Submission due state","unit":"count","source":"ExpectedSubmission.due_state"},
        ])


class StatusDonutView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        qs = dashboard_queryset(request)
        return Response(list(qs.values("workflow_status").annotate(count=Count("id")).order_by("workflow_status")))


class CategoryCompletionView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        qs = dashboard_queryset(request)
        cats = qs.values("provider__sector", "provider__category").annotate(
            total=Count("id"),
            approved=Count("id", filter=Q(workflow_status="APPROVED")),
        )
        return Response([{
            "sector": c["provider__sector"],
            "category": c["provider__category"],
            "completion_pct": round((c["approved"] / (c["total"] or 1)) * 100, 1),
            "total": c["total"], "approved": c["approved"],
        } for c in cats])


class SubmissionTrendView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        cutoff = timezone.now() - timedelta(days=365)
        data = (
            Submission.objects.filter(submitted_at__gte=cutoff, submitted_at__isnull=False)
            .annotate(month=TruncMonth("submitted_at"))
            .values("month").annotate(count=Count("id")).order_by("month")
        )
        return Response(list(data))


class OverdueByFormView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request):
        return Response(list(
            dashboard_queryset(request).filter(due_state="OVERDUE")
            .values("form_template__form_code", "form_template__name")
            .annotate(count=Count("id")).order_by("-count")
        ))


# ── Periods ──────────────────────────────────────────────────────────────────

class ReportingPeriodListView(generics.ListCreateAPIView):
    queryset = ReportingPeriod.objects.all()
    serializer_class = ReportingPeriodSerializer
    filterset_fields = ["frequency", "status", "year"]

    def get_permissions(self):
        if self.request.method == "POST":
            return [IsNCAEditor()]
        return [IsNCAUser()]

    def perform_create(self, serializer):
        period = serializer.save()
        write_audit(self.request, "REPORTING_PERIOD_CREATED", "ReportingPeriod", period.id, after={
            "name": period.name, "frequency": period.frequency, "year": period.year, "status": period.status,
        })


class ReportingPeriodDetailView(generics.RetrieveUpdateAPIView):
    queryset = ReportingPeriod.objects.all()
    serializer_class = ReportingPeriodSerializer

    def get_permissions(self):
        if self.request.method in ("GET", "HEAD", "OPTIONS"):
            return [IsNCAUser()]
        return [IsNCAEditor()]

    def perform_update(self, serializer):
        period = serializer.instance
        before = {"name": period.name, "status": period.status, "due_at": period.due_at.isoformat()}
        period = serializer.save()
        write_audit(self.request, "REPORTING_PERIOD_UPDATED", "ReportingPeriod", period.id, before=before, after={
            "name": period.name, "status": period.status, "due_at": period.due_at.isoformat(),
            "changed_fields": sorted(serializer.validated_data),
        })


class ActivatePeriodView(APIView):
    permission_classes = [IsNCAEditor]

    def post(self, request, pk):
        try:
            period = ReportingPeriod.objects.get(pk=pk)
        except ReportingPeriod.DoesNotExist:
            return Response({"detail": "Not found."}, status=404)
        if period.status != "DRAFT":
            return Response({"detail": "Only DRAFT periods can be activated."}, status=400)
        try:
            period.activate()
        except ValueError as exc:
            return Response({"code": "ACTIVATION_BLOCKED", "detail": str(exc)}, status=409)
        write_audit(request, "PERIOD_ACTIVATED", "ReportingPeriod", period.id)
        return Response({"detail": "Activated.", "expected_count": period.expected_submissions.count()})


class PeriodAssignmentPreviewView(APIView):
    permission_classes = [IsNCAEditor]

    def get(self, request, pk):
        from apps.providers.models import ProviderFormAssignment
        period = generics.get_object_or_404(ReportingPeriod, pk=pk)
        provider_ids = set(period.assigned_providers.values_list("id", flat=True))
        template_ids = set(period.applicable_form_templates.values_list("id", flat=True))
        recurring = ProviderFormAssignment.objects.filter(
            form_family__canonical_frequency=period.frequency,
            obligation__in=["REQUIRED", "OPTIONAL"], effective_from__lte=period.due_at.date(),
        ).filter(Q(effective_to__isnull=True) | Q(effective_to__gte=period.opens_at.date())).select_related("provider", "form_family")
        if provider_ids:
            recurring = recurring.filter(provider_id__in=provider_ids)
        rows, blockers = [], []
        for assignment in recurring:
            form = assignment.form_family.versions.filter(status="ACTIVE", approval_status="APPROVED", frequency=period.frequency).order_by("-published_at", "-id").first()
            if not form:
                blockers.append(f"No active approved {assignment.form_family.code} version is available.")
                continue
            if template_ids and form.id not in template_ids:
                continue
            rows.append({"source": "RECURRING", "provider_id": assignment.provider_id, "provider_name": assignment.provider.registered_name, "form_template": form.id, "form_code": form.form_code})
        for manual in period.manual_form_assignments.select_related("provider", "form_template"):
            rows.append({"source": "MANUAL", "provider_id": manual.provider_id, "provider_name": manual.provider.registered_name, "form_template": manual.form_template_id, "form_code": manual.form_template.form_code})
        unique = {(row["provider_id"], row["form_template"]): row for row in rows}
        return Response({"period": period.id, "pairs": list(unique.values()), "count": len(unique), "blockers": blockers, "can_activate": not blockers and bool(unique)})


class PeriodManualReminderView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        period = get_object_or_404(ReportingPeriod, pk=pk)
        selected_ids = request.data.get("expected_submission_ids", [])
        subject = str(request.data.get("subject") or "").strip()
        body = str(request.data.get("message") or "").strip()
        request_id = str(request.data.get("client_request_id") or "").strip()
        if period.status != "ACTIVE":
            return Response({"detail": "Manual reminders can only be sent for an active reporting period."}, status=409)
        if not isinstance(selected_ids, list) or not selected_ids:
            return Response({"detail": "Select at least one form obligation."}, status=400)
        if not subject or not body:
            return Response({"detail": "An email subject and message are required."}, status=400)
        try:
            selected_ids = list(dict.fromkeys(int(value) for value in selected_ids))
        except (TypeError, ValueError):
            return Response({"detail": "Every selected obligation must have a numeric ID."}, status=400)

        queryset = ExpectedSubmission.objects.select_for_update().filter(period=period, id__in=selected_ids).select_related(
            "provider", "form_template",
        ).prefetch_related("versions")
        if queryset.count() != len(selected_ids):
            return Response({"detail": "One or more selected obligations do not belong to this period."}, status=400)
        blocked = queryset.filter(workflow_status__in=["SUBMITTED", "UNDER_REVIEW", "RESUBMITTED", "APPROVED", "REJECTED", "ARCHIVED"])
        if blocked.exists():
            return Response({"detail": "Forms already sent to NCA, completed, rejected or archived cannot receive completion reminders."}, status=409)

        from apps.compliance.communications import ensure_email_handoff, handoff_data
        reminders = []
        for expected in queryset.order_by("provider__registered_name", "id"):
            submission = expected.versions.order_by("-version", "-id").first()
            if not submission:
                return Response({"detail": f"Obligation {expected.id} has no submission version."}, status=409)
            existing = None
            if request_id:
                existing = SubmissionEvent.objects.filter(
                    submission=submission, event_type="REMINDER",
                    metadata__manual_request_id=request_id,
                ).first()
            if existing:
                handoff = ensure_email_handoff(submission=submission, event=existing, actor=request.user)
            else:
                event = emit_submission_event(
                    submission=submission, actor=request.user, event_type="REMINDER",
                    message=f"NCA sent a manual reminder for {submission.submission_reference}.",
                    from_status=expected.workflow_status, to_status=expected.workflow_status,
                    audience="PROVIDER", notify=["PROVIDER_DATA_ENTRY", "PROVIDER_APPROVER"],
                    metadata={
                        "manual": True, "manual_request_id": request_id,
                        "communication_subject": subject,
                        "communication_body": body,
                        "period_id": period.id,
                    },
                )
                handoff = ensure_email_handoff(submission=submission, event=event, actor=request.user)
            reminders.append({
                "expected_submission_id": expected.id,
                "provider": expected.provider.registered_name,
                "submission_reference": submission.submission_reference,
                "email": handoff_data(handoff),
            })
        write_audit(
            request, "PERIOD_MANUAL_REMINDERS_PREPARED", "ReportingPeriod", period.id,
            after={"obligation_ids": selected_ids, "count": len(reminders), "client_request_id": request_id},
        )
        return Response({"period": period.id, "count": len(reminders), "reminders": reminders})


# ── Expected Submissions ──────────────────────────────────────────────────────

class ExpectedSubmissionListView(generics.ListAPIView):
    permission_classes = [IsAuthenticated]
    queryset = ExpectedSubmission.objects.select_related("provider", "form_template", "period", "assigned_officer")
    serializer_class = ExpectedSubmissionSerializer
    filterset_fields = [
        "workflow_status", "due_state", "form_template", "period", "provider",
        "provider__sector", "provider__category", "assigned_officer",
    ]
    search_fields = ["provider__registered_name", "form_template__form_code"]
    ordering_fields = ["period__due_at", "workflow_status", "due_state"]

    def get_queryset(self):
        queryset = expected_submissions_for_user(self.request.user).exclude(workflow_status="ARCHIVED").select_related(
            "provider", "form_template", "period", "assigned_officer"
        )
        attention = self.request.query_params.get("attention_required")
        flag_status = self.request.query_params.get("compliance_status")
        flag_type = self.request.query_params.get("compliance_flag_type")
        if attention in {"1", "true", "True"}:
            queryset = queryset.filter(compliance_flags__status__in=["OPEN", "ACKNOWLEDGED", "IN_PROGRESS"])
        if flag_status:
            queryset = queryset.filter(compliance_flags__status=flag_status)
        if flag_type:
            queryset = queryset.filter(compliance_flags__flag_type=flag_type)
        latest_event_at = SubmissionEvent.objects.filter(
            submission__expected_id=OuterRef("pk"),
        ).order_by("-created_at", "-id").values("created_at")[:1]
        return queryset.distinct().annotate(
            latest_activity_at=Subquery(latest_event_at),
        ).order_by(F("latest_activity_at").desc(nulls_last=True), "-created_at", "-id")


class ExpectedSubmissionDetailView(generics.RetrieveUpdateAPIView):
    serializer_class = ExpectedSubmissionSerializer
    http_method_names = ["get", "patch", "head", "options"]

    def get_permissions(self):
        if self.request.method == "PATCH":
            return [IsNCAEditor()]
        return [IsAuthenticated()]

    def get_queryset(self):
        return expected_submissions_for_user(self.request.user)

    def perform_update(self, serializer):
        expected = serializer.save()
        write_audit(
            self.request, "EXPECTED_SUBMISSION_UPDATED", "ExpectedSubmission", expected.id,
            after={"changed_fields": sorted(serializer.validated_data)},
        )


class ProviderWorkspaceSummaryView(APIView):
    permission_classes = [IsProviderUser]

    def get(self, request):
        return Response(summary_for_user(request.user))


class ProviderWorkspaceSubmissionListView(generics.ListAPIView):
    permission_classes = [IsProviderUser]
    serializer_class = ExpectedSubmissionSerializer
    pagination_class = ProviderWorkspacePagination

    def get_queryset(self):
        queryset = workspace_queryset(self.request.user)
        queryset = apply_workspace_queue(queryset, self.request.user, self.request.query_params.get("queue"))
        queryset = filter_workspace_queryset(queryset, self.request.query_params)
        ordering = self.request.query_params.get("ordering", "period__due_at")
        allowed = {
            "period__due_at", "-period__due_at", "workflow_status", "-workflow_status",
            "form_template__form_code", "-form_template__form_code", "created_at", "-created_at",
        }
        return queryset.order_by(ordering if ordering in allowed else "period__due_at", "id")


class ProviderFormsListView(generics.ListAPIView):
    """Active provider work; formal versions belong on Provider Submissions."""
    permission_classes = [IsProviderUser]
    serializer_class = ExpectedSubmissionSerializer
    pagination_class = ProviderWorkspacePagination

    def get_queryset(self):
        queryset = workspace_queryset(self.request.user).exclude(provider_status="CLOSED")
        queryset = filter_workspace_queryset(queryset, self.request.query_params)
        provider_status = self.request.query_params.get("provider_status")
        if provider_status:
            queryset = queryset.filter(provider_status=provider_status)
        return queryset.order_by("period__due_at", "id")


class ProviderFormSummaryView(APIView):
    permission_classes = [IsProviderUser]

    def get(self, request):
        queryset = ExpectedSubmission.objects.filter(provider__organization_id=request.user.organization_id)
        return Response({
            "in_progress": queryset.filter(provider_status="IN_PROGRESS").count(),
            "awaiting_approval": queryset.filter(provider_status="AWAITING_APPROVAL").count(),
            "corrections_required": queryset.filter(provider_status="CORRECTIONS_REQUIRED").count(),
            "total_active": queryset.exclude(provider_status="CLOSED").count(),
        })


FORMAL_REGULATORY_STATUSES = {
    "SUBMITTED", "UNDER_REVIEW", "RETURNED_FOR_CORRECTION", "APPROVED", "REJECTED",
}


def formal_submission_queryset(user):
    return submissions_for_user(user).filter(
        regulatory_status__in=FORMAL_REGULATORY_STATUSES,
    ).select_related(
        "expected__provider", "expected__form_template", "expected__period",
        "provider_approval__approver", "receipt",
    ).order_by("-submitted_at", "-created_at", "-id")


class ProviderFormalSubmissionListView(generics.ListAPIView):
    permission_classes = [IsProviderUser]
    serializer_class = ProviderFormalTaskListSerializer
    pagination_class = ProviderWorkspacePagination

    def get_queryset(self):
        formal = Submission.objects.filter(
            regulatory_status__in=FORMAL_REGULATORY_STATUSES,
        ).select_related("provider_approval__approver", "receipt").order_by("-version", "-id")
        latest_formal = Submission.objects.filter(
            expected_id=OuterRef("pk"), regulatory_status__in=FORMAL_REGULATORY_STATUSES,
        ).order_by("-version", "-id")
        queryset = expected_submissions_for_user(self.request.user).filter(
            versions__regulatory_status__in=FORMAL_REGULATORY_STATUSES,
        ).select_related("provider", "form_template", "period").prefetch_related(
            Prefetch("versions", queryset=formal, to_attr="latest_formal_versions"),
        ).annotate(
            latest_formal_status=Subquery(latest_formal.values("regulatory_status")[:1]),
            latest_formal_timestamp=Subquery(latest_formal.values("submitted_at")[:1]),
        ).distinct()
        if value := self.request.query_params.get("regulatory_status"):
            queryset = queryset.filter(latest_formal_status=value)
        if search := (self.request.query_params.get("search") or "").strip():
            queryset = queryset.filter(
                Q(versions__submission_reference__icontains=search)
                | Q(form_template__form_code__icontains=search)
                | Q(form_template__name__icontains=search)
                | Q(period__name__icontains=search)
            )
        return queryset.order_by("-latest_formal_timestamp", "-created_at", "-id")


class ProviderFormalSubmissionDetailView(generics.RetrieveAPIView):
    permission_classes = [IsProviderUser]
    serializer_class = FormalSubmissionListSerializer

    def get_queryset(self):
        return formal_submission_queryset(self.request.user)


class ExpectedSubmissionPenaltyView(APIView):
    permission_classes = [IsNCAAdmin]

    @transaction.atomic
    def patch(self, request, pk):
        expected = get_object_or_404(
            ExpectedSubmission.objects.select_for_update().select_related("provider", "period"), pk=pk,
        )
        if request.data.get("action") == "MARK_PAID":
            if expected.penalty_amount_ghs == 0:
                return Response(ExpectedSubmissionSerializer(expected, context={"request": request}).data)
            payment_reference = str(request.data.get("payment_reference") or "").strip()
            payment_note = str(request.data.get("payment_note") or "").strip()
            before = {"penalty_amount_ghs": str(expected.penalty_amount_ghs),
                      "penalty_reference": expected.penalty_reference}
            expected.penalty_amount_ghs = Decimal("0.00")
            expected.penalty_updated_by = request.user
            expected.penalty_updated_at = timezone.now()
            expected.save(update_fields=["penalty_amount_ghs", "penalty_updated_by", "penalty_updated_at"])
            write_audit(
                request, "OBLIGATION_PENALTY_PAID", "ExpectedSubmission", expected.id,
                before=before, after={"penalty_amount_ghs": "0.00",
                    "penalty_reference": expected.penalty_reference,
                    "payment_reference": payment_reference, "payment_note": payment_note},
            )
            return Response(ExpectedSubmissionSerializer(expected, context={"request": request}).data)
        try:
            amount = Decimal(str(request.data.get("penalty_amount_ghs", expected.penalty_amount_ghs)))
        except (InvalidOperation, TypeError, ValueError):
            return Response({"detail": "Penalty amount must be a valid number."}, status=400)
        if amount < 0:
            return Response({"detail": "Penalty amount cannot be negative."}, status=400)
        reference = str(request.data.get("penalty_reference", expected.penalty_reference) or "").strip()
        note = str(request.data.get("penalty_note", expected.penalty_note) or "").strip()
        if amount > 0 and not reference:
            return Response({"detail": "A penalty reference is required for a non-zero penalty."}, status=400)
        before = {
            "penalty_amount_ghs": str(expected.penalty_amount_ghs),
            "penalty_reference": expected.penalty_reference,
            "penalty_note": expected.penalty_note,
        }
        expected.penalty_amount_ghs = amount
        expected.penalty_reference = reference
        expected.penalty_note = note
        expected.penalty_updated_by = request.user
        expected.penalty_updated_at = timezone.now()
        expected.save(update_fields=[
            "penalty_amount_ghs", "penalty_reference", "penalty_note",
            "penalty_updated_by", "penalty_updated_at",
        ])
        write_audit(
            request, "OBLIGATION_PENALTY_UPDATED", "ExpectedSubmission", expected.id,
            before=before,
            after={
                "penalty_amount_ghs": str(expected.penalty_amount_ghs),
                "penalty_reference": expected.penalty_reference,
                "penalty_note": expected.penalty_note,
            },
        )
        return Response(ExpectedSubmissionSerializer(expected, context={"request": request}).data)


class NCAFormalSubmissionListView(generics.ListAPIView):
    permission_classes = [IsNCAUser]
    serializer_class = FormalSubmissionListSerializer
    pagination_class = ProviderWorkspacePagination

    def get_queryset(self):
        queryset = formal_submission_queryset(self.request.user)
        for parameter, lookup in {
            "regulatory_status": "regulatory_status", "provider": "expected__provider_id",
            "period": "expected__period_id", "form_template": "expected__form_template_id",
            "due_state": "expected__due_state", "provider__sector": "expected__provider__sector",
            "provider__category": "expected__provider__category",
        }.items():
            if value := self.request.query_params.get(parameter):
                queryset = queryset.filter(**{lookup: value})
        if self.request.query_params.get("attention_required") in {"1", "true", "True"}:
            queryset = queryset.filter(expected__compliance_flags__status__in=["OPEN", "ACKNOWLEDGED", "IN_PROGRESS"]).distinct()
        if value := self.request.query_params.get("compliance_status"):
            queryset = queryset.filter(expected__compliance_flags__status=value).distinct()
        if value := self.request.query_params.get("compliance_flag_type"):
            queryset = queryset.filter(expected__compliance_flags__flag_type=value).distinct()
        if search := (self.request.query_params.get("search") or "").strip():
            queryset = queryset.filter(
                Q(submission_reference__icontains=search)
                | Q(expected__provider__registered_name__icontains=search)
                | Q(expected__form_template__form_code__icontains=search)
                | Q(expected__form_template__name__icontains=search)
            )
        return queryset


# ── Submissions ───────────────────────────────────────────────────────────────

class StartSubmissionView(APIView):
    permission_classes = [IsProviderDataEntry]

    @transaction.atomic
    def post(self, request, pk):
        visible = get_expected_submission_for_user(request.user, pk=pk)
        expected = ExpectedSubmission.objects.select_for_update().get(pk=visible.pk)
        if expected.workflow_status != "NOT_STARTED":
            return Response({"detail": "Cannot start in current state."}, status=400)
        last = expected.versions.order_by("-version").first()
        correcting = False
        if correcting:
            if last.supersedes_id and last.submitted_at is None:
                submission = last
            else:
                submission = clone_for_nca_correction(last)
                last.correction_items.filter(stage="NCA_REVIEW", status="OPEN").update(resolution_submission=submission)
        else:
            submission = last if last and last.submitted_at is None else Submission.objects.create(expected=expected, version=1)
        prior_status = expected.workflow_status
        expected.workflow_status = "DRAFT"
        expected.provider_status = "IN_PROGRESS"
        expected.save(update_fields=["workflow_status", "provider_status"])
        audit_transition(
            request=request, submission=submission, event_type="SUBMISSION_STARTED",
            message="A correction version was opened for editing." if correcting else "The form was opened for data entry.",
            from_status=prior_status, to_status="DRAFT", audience="BOTH",
        )
        return Response(SubmissionSerializer(submission).data, status=201)


class SubmissionDetailView(generics.RetrieveAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = SubmissionSerializer

    def get_queryset(self):
        return submissions_for_user(self.request.user).select_related(
            "expected__provider", "expected__form_template", "expected__period"
        )


class SubmissionReviewDataView(APIView):
    """One exact-version payload for regulatory review; never substitutes a newer template."""
    permission_classes = [IsNCAUser]

    def get(self, request, pk):
        from apps.forms_engine.gaps import recalculate_form_gaps
        from apps.forms_engine.serializers import FormTemplateDetailSerializer, FormGapAssessmentSerializer
        from apps.compliance.serializers import ComplianceFlagSerializer

        submission = get_submission_for_user(request.user, pk=pk)
        template = submission.expected.form_template
        archived = template is None
        readiness = {
            "can_submit": False, "blocking_issues": [], "completeness_warnings": [],
            "missing_indicator_count": 0, "archived": True,
        } if archived else refresh_submission_completion(submission)
        assessments = [] if archived else recalculate_form_gaps(template, request.user)
        values = submission.values.select_related("field", "grid", "grid_column", "non_filled_disposition")
        uploads = submission.kmz_uploads.select_related("requirement", "reviewed_by")
        validation_run = submission.validation_runs.prefetch_related("results").order_by("-started_at").first()
        correction_items = CorrectionItem.objects.filter(
            Q(source_submission=submission) | Q(resolution_submission=submission)
        ).distinct()
        correction_changes = []
        if submission.supersedes_id:
            def value_key(item):
                if item.field_id:
                    return ("FIELD", str(item.field_id))
                return ("GRID_CELL", f"{item.grid_id}:{item.grid_row_id}:{item.grid_column_id}")

            previous_values = {
                value_key(item): item
                for item in submission.supersedes.values.all()
            }
            current_values = {
                value_key(item): item
                for item in submission.values.all()
            }
            resolution_corrections = [
                item for item in correction_items
                if item.resolution_submission_id == submission.id
            ]
            for target in sorted(
                set(previous_values) | set(current_values), key=str,
            ):
                before = previous_values.get(target)
                after = current_values.get(target)
                before_state = (
                    before.value if before else None,
                    before.value_status if before else None,
                    before.explanation if before else None,
                )
                after_state = (
                    after.value if after else None,
                    after.value_status if after else None,
                    after.explanation if after else None,
                )
                if before_state == after_state:
                    continue
                value_item = after or before
                section_code = (
                    value_item.field.section.section_code
                    if value_item.field_id else value_item.grid.section.section_code
                )
                item = next((candidate for candidate in resolution_corrections if (
                    candidate.target_type == "SUBMISSION"
                    or (candidate.target_type == "SECTION" and candidate.target_id == section_code)
                    or (candidate.target_type == target[0] and candidate.target_id == target[1])
                )), None)
                if item is None:
                    continue
                correction_changes.append({
                    "target_type": target[0],
                    "target_id": target[1],
                    "target_label": correction_target_label(
                        template, target[0], target[1],
                    ),
                    "instruction": item.instruction,
                    "status": item.status,
                    "before_value": before.value if before else None,
                    "before_status": before.value_status if before else None,
                    "after_value": after.value if after else None,
                    "after_status": after.value_status if after else None,
                })
        return Response({
            "submission": SubmissionSerializer(submission).data,
            "obligation": ExpectedSubmissionSerializer(
                submission.expected, context={"request": request},
            ).data,
            "template": submission.form_schema_snapshot if archived else FormTemplateDetailSerializer(template).data,
            "values": SubmissionValueSerializer(values, many=True).data,
            "requirements": FormGapAssessmentSerializer(assessments, many=True).data,
            "uploads": [{
                "id": item.id, "requirement": item.requirement_id,
                "category": item.requirement.category if item.requirement_id else item.requirement_snapshot.get("category", ""),
                "file_name": item.file_name, "file_size": item.file_size, "sha256": item.sha256,
                "scan_status": item.scan_status, "review_status": item.review_status,
                "review_note": item.review_note, "uploaded_at": item.uploaded_at,
            } for item in uploads],
            "validation": None if not validation_run else {
                "id": validation_run.id, "status": validation_run.status, "scope": validation_run.scope,
                "completed_at": validation_run.completed_at,
                "results": [{"id": result.id, "severity": result.severity, "target_type": result.target_type,
                    "target_id": result.target_id, "code": result.code, "message": result.message,
                    "details": result.details} for result in validation_run.results.all()],
            },
            "approval_blockers": readiness.get("blocking_issues", []),
            "correction_items": [{"id": item.id, "stage": item.stage, "target_type": item.target_type, "target_id": item.target_id,
                "target_label": correction_target_label(template, item.target_type, item.target_id),
                "instruction": item.instruction, "status": item.status} for item in correction_items],
            "correction_changes": correction_changes,
            # Deprecated response key retained for API compatibility. Exact-version
            # review remains in force without presenting a misleading legacy banner.
            "legacy_warning": None,
            "previous_month": previous_month_values(submission),
            "compliance_flags": ComplianceFlagSerializer(
                submission.expected.compliance_flags.all(), many=True,
            ).data,
        })


class ProviderReviewDataView(APIView):
    """Provider-tenant review payload bound to the submission's immutable template."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        from apps.forms_engine.serializers import FormTemplateDetailSerializer
        from apps.compliance.serializers import ComplianceFlagSerializer
        from apps.uploads.models import SubmissionKMZUpload, SubmissionExcelBackup

        submission = get_submission_for_user(request.user, pk=pk)
        archived = submission.expected.form_template_id is None
        readiness = {
            "can_submit": False, "blocking_issues": [], "completeness_warnings": [],
            "missing_indicator_count": 0, "archived": True,
        } if archived else refresh_submission_completion(submission)
        corrections = CorrectionItem.objects.filter(
            Q(source_submission=submission) | Q(resolution_submission=submission)
        ).select_related("created_by").distinct()
        kmz = SubmissionKMZUpload.objects.filter(submission=submission).select_related("requirement")
        backups = SubmissionExcelBackup.objects.filter(submission=submission).order_by("-uploaded_at")
        events = submission.timeline_events.filter(audience__in=["PROVIDER", "BOTH"]).select_related("actor")
        versions = submission.expected.versions.select_related("receipt", "provider_approval__approver").order_by("-version")
        edit_batches = submission.provider_edit_batches.select_related("actor").prefetch_related("items")
        return Response({
            "submission": SubmissionSerializer(submission).data,
            "template": submission.form_schema_snapshot if archived else FormTemplateDetailSerializer(submission.expected.form_template).data,
            "values": SubmissionValueSerializer(
                submission.values.select_related("field", "grid", "grid_column", "non_filled_disposition"),
                many=True,
            ).data,
            "readiness": {
                **readiness,
                "transition_ready": readiness["can_submit"] and not corrections.filter(status="OPEN").exists(),
            },
            "correction_items": [{
                "id": item.id, "stage": item.stage, "target_type": item.target_type,
                "target_id": item.target_id,
                "target_label": correction_target_label(submission.expected.form_template, item.target_type, item.target_id) if submission.expected.form_template_id else "Flagged item",
                "instruction": item.instruction,
                "status": item.status, "created_by_name": item.created_by.name,
                "created_at": item.created_at,
            } for item in corrections],
            "uploads": {
                "kmz": [{
                    "id": item.id, "requirement_id": item.requirement_id,
                    "category": item.requirement.get_category_display() if item.requirement_id else item.requirement_snapshot.get("category", ""), "file_name": item.file_name,
                    "file_size": item.file_size, "sha256": item.sha256,
                    "scan_status": item.scan_status, "review_status": item.review_status,
                    "review_note": item.review_note, "uploaded_at": item.uploaded_at,
                    "download_ready": item.scan_status == "CLEAN",
                } for item in kmz],
                "excel": [{
                    "id": item.id, "file_name": item.file_name, "file_size": item.file_size,
                    "sha256": item.sha256, "scan_status": item.scan_status,
                    "source_control_status": item.source_control_status, "uploaded_at": item.uploaded_at,
                    "download_ready": item.scan_status == "CLEAN",
                } for item in backups],
            },
            "timeline": SubmissionEventSerializer(events, many=True).data,
            "versions": [SubmissionSerializer(version).data for version in versions],
            "provider_edits": [{
                "id": batch.id, "actor": batch.actor_id, "actor_name": batch.actor.name,
                "stage": batch.stage, "section_code": batch.section_code,
                "base_revision": batch.base_revision, "resulting_revision": batch.resulting_revision,
                "item_count": batch.item_count, "changes_sha256": batch.changes_sha256,
                "created_at": batch.created_at,
                "items": [{
                    "id": item.id, "target_type": item.target_type, "target_id": item.target_id,
                    "before": item.before, "after": item.after,
                } for item in batch.items.all()],
            } for batch in edit_batches],
            "permitted_actions": permitted_actions(request.user, submission.expected, submission),
            "previous_month": previous_month_values(submission),
            "compliance_flags": ComplianceFlagSerializer(
                submission.expected.compliance_flags.all(), many=True,
            ).data,
        })


class SectionValuesView(APIView):
    def get_permissions(self):
        if self.request.method == "PUT":
            return [IsProviderUser()]
        return [IsAuthenticated()]

    def get(self, request, pk, section_code):
        submission = get_submission_for_user(request.user, pk=pk)
        values = SubmissionValue.objects.filter(
            Q(field__section__section_code=section_code)
            | Q(grid__section__section_code=section_code),
            submission=submission,
        ).select_related("field", "grid", "grid_column")
        return Response(SubmissionValueSerializer(values, many=True).data)

    @transaction.atomic
    def put(self, request, pk, section_code):
        visible = get_submission_for_user(request.user, pk=pk)
        submission = lock_submission(visible.pk)
        values_payload = request.data.get("values", [])
        if not isinstance(values_payload, list):
            return Response({"detail": "values must be a list."}, status=400)
        try:
            client_save_id = uuid.UUID(str(request.data.get("client_save_id") or uuid.uuid4()))
        except (TypeError, ValueError, AttributeError):
            return Response({"detail": "client_save_id must be a valid UUID."}, status=400)
        try:
            persisted_change_version = max(0, int(request.data.get("change_version") or 0))
        except (TypeError, ValueError):
            return Response({"detail": "change_version must be a non-negative integer."}, status=400)
        canonical_payload = json.dumps(
            {"section_code": section_code, "values": values_payload},
            ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        )
        payload_sha256 = hashlib.sha256(canonical_payload.encode("utf-8")).hexdigest()
        replay = SectionSaveReceipt.objects.filter(
            submission=submission, client_save_id=client_save_id,
        ).first()
        if replay:
            if replay.payload_sha256 != payload_sha256 or replay.section_code != section_code or replay.actor_id != request.user.id:
                return Response({
                    "code": "IDEMPOTENCY_KEY_REUSED",
                    "detail": "This client_save_id was already used for a different section save.",
                }, status=409)
            return Response({**replay.response, "replayed": True})
        if submission.expected.workflow_status == "CORRECTION_REQUESTED" and not submission.supersedes_id:
            return Response({"detail": "Official historical versions are immutable. Edit the linked correction version."}, status=409)
        if not provider_can_edit(request.user, submission):
            return Response({
                "code": "EDIT_NOT_PERMITTED",
                "detail": "Your provider role cannot edit this submission at its current workflow stage.",
            }, status=403)
        supplied_revision = request.data.get("revision")
        try:
            supplied_revision_value = int(supplied_revision) if supplied_revision is not None else None
        except (TypeError, ValueError):
            return Response({"detail": "revision must be an integer."}, status=400)
        if supplied_revision_value is not None and supplied_revision_value != submission.revision:
            return Response({
                "code": "STALE_REVISION",
                "detail": "A newer version was saved. Your unsaved values were not overwritten; reload before reapplying them.",
                "current_revision": submission.revision,
                "last_edited_by": submission.last_edited_by_id,
                "last_edited_by_name": submission.last_edited_by.name if submission.last_edited_by else None,
                "last_edited_at": submission.last_edited_at,
            }, status=409)

        template = submission.expected.form_template
        section = get_object_or_404(template.sections.all(), section_code=section_code)
        correction_items = list(
            CorrectionItem.objects.filter(
                Q(source_submission=submission, stage="PROVIDER_APPROVAL")
                | Q(resolution_submission=submission, stage="NCA_REVIEW"),
                status__in=["OPEN", "ADDRESSED"],
            )
        )
        restrict_to_corrections = submission.expected.workflow_status in {
            "PROVIDER_CHANGES_REQUESTED", "CORRECTION_REQUESTED",
        }
        existing_before = list(SubmissionValue.objects.filter(
            Q(field__section=section) | Q(grid__section=section), submission=submission,
        ))

        def value_snapshot(item):
            return {
                "value": item.value,
                "value_status": item.value_status,
                "explanation": item.explanation,
            }

        def value_key(item):
            if item.field_id:
                return ("FIELD", str(item.field_id))
            return ("GRID_CELL", f"{item.grid_id}:{item.grid_row_id}:{item.grid_column_id}")

        before = {value_key(item): value_snapshot(item) for item in existing_before}
        saved = []
        keep_scalar_ids = set()
        keep_grid_keys = set()
        target_keys = set()
        non_filled = {"NOT_APPLICABLE", "NOT_AVAILABLE", "NOT_REQUIRED"}
        valid_statuses = {choice[0] for choice in SubmissionValue._meta.get_field("value_status").choices}
        for v in values_payload:
            field_id = v.get("field")
            grid_id = v.get("grid")
            if bool(field_id) == bool(grid_id):
                return Response({"detail": "Each value must reference exactly one field or grid."}, status=400)
            if field_id and not section.fields.filter(id=field_id).exists():
                return Response({"detail": "Field does not belong to this form section."}, status=400)
            if grid_id and not section.grids.filter(id=grid_id, columns__id=v.get("grid_column")).exists():
                return Response({"detail": "Grid cell does not belong to this form section."}, status=400)
            value_status = v.get("value_status") or ("PROVIDED" if str(v.get("value", "")).strip() else "MISSING")
            if value_status not in valid_statuses or value_status == "SYSTEM_CALCULATED":
                return Response({"detail": "Invalid value status."}, status=400)
            value = "" if value_status in non_filled else str(v.get("value", ""))
            explanation = str(v.get("explanation", ""))
            if grid_id and not str(v.get("grid_row_id", "")).strip():
                return Response({"detail": "Every grid cell requires a row identifier."}, status=400)
            target_type = "FIELD" if field_id else "GRID_CELL"
            target = str(field_id or f"{grid_id}:{v.get('grid_row_id','')}:{v.get('grid_column','')}")
            allowed = True
            if restrict_to_corrections and correction_items:
                allowed = any(
                    item.target_type == "SUBMISSION"
                    or (item.target_type == "SECTION" and item.target_id == section_code)
                    or (item.target_type == "FIELD" and item.target_id == str(field_id))
                    or (item.target_type == "GRID_CELL" and item.target_id == target)
                    for item in correction_items
                )
            if field_id:
                keep_scalar_ids.add(int(field_id))
            else:
                key = (int(grid_id), str(v.get("grid_row_id", "")), int(v.get("grid_column")))
                keep_grid_keys.add(key)
            incoming_snapshot = {
                "value": value, "value_status": value_status, "explanation": explanation,
            }
            existing_key = (target_type, target)
            if not allowed:
                if before.get(existing_key, {}) != incoming_snapshot:
                    return Response({
                        "code": "CORRECTION_SCOPE_LOCKED",
                        "detail": "A locked value outside the correction request was changed.",
                        "target_type": target_type, "target_id": target,
                    }, status=403)
                continue
            target_keys.add(target)
            obj, _ = SubmissionValue.objects.update_or_create(
                submission=submission,
                field_id=field_id,
                grid_id=grid_id,
                grid_row_id=v.get("grid_row_id", ""),
                grid_column_id=v.get("grid_column"),
                defaults={
                    "value": value,
                    "value_status": value_status,
                    "explanation": explanation,
                    "value_source": "MANUAL",
                    "source_reference": "",
                    "updated_by": request.user,
                },
            )
            saved.append(obj)

        # Authoritative section snapshot: omitted editable values are removed.
        existing = SubmissionValue.objects.filter(
            Q(field__section=section) | Q(grid__section=section), submission=submission,
        )
        for current in existing:
            key = (current.grid_id, current.grid_row_id, current.grid_column_id)
            keep = current.field_id in keep_scalar_ids if current.field_id else key in keep_grid_keys
            if not keep:
                if restrict_to_corrections and correction_items:
                    target = str(current.field_id or f"{current.grid_id}:{current.grid_row_id}:{current.grid_column_id}")
                    allowed = any(
                        item.target_type == "SUBMISSION"
                        or (item.target_type == "SECTION" and item.target_id == section_code)
                        or (item.target_type == "FIELD" and item.target_id == str(current.field_id))
                        or (item.target_type == "GRID_CELL" and item.target_id == target)
                        for item in correction_items
                    )
                    if not allowed:
                        continue
                current.delete()

        if submission.expected.workflow_status == "NOT_STARTED":
            submission.expected.workflow_status = "DRAFT"
            submission.expected.provider_status = "IN_PROGRESS"
            submission.expected.save(update_fields=["workflow_status", "provider_status"])
            audit_transition(
                request=request, submission=submission, event_type="SUBMISSION_STARTED",
                message="The form was opened for data entry by saving its first section.",
                from_status="NOT_STARTED", to_status="DRAFT", audience="BOTH",
            )
        base_revision = submission.revision
        revision = complete_submission_revision(submission, request.user)

        after_values = list(SubmissionValue.objects.filter(
            Q(field__section=section) | Q(grid__section=section), submission=submission,
        ))
        after = {value_key(item): value_snapshot(item) for item in after_values}
        changes = []
        for key in sorted(set(before) | set(after)):
            previous = before.get(key, {})
            current = after.get(key, {})
            if previous != current:
                changes.append({
                    "target_type": key[0], "target_id": key[1],
                    "before": previous, "after": current,
                })
        changed_target_keys = {change["target_id"] for change in changes}
        if restrict_to_corrections and correction_items and changed_target_keys:
            mark_matching_corrections_addressed(
                submission, section_code=section_code, target_keys=changed_target_keys,
            )
        edit_batch = None
        if request.user.role == "PROVIDER_APPROVER" and changes:
            canonical = json.dumps(changes, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            edit_batch = ProviderEditBatch.objects.create(
                submission=submission, actor=request.user,
                stage=submission.expected.workflow_status, section_code=section_code,
                base_revision=base_revision, resulting_revision=revision,
                item_count=len(changes), changes_sha256=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
            )
            ProviderEditItem.objects.bulk_create([
                ProviderEditItem(batch=edit_batch, **change) for change in changes
            ])

        readiness = refresh_submission_completion(submission, section_code)
        audit_metadata = {"section": section_code, "count": len(saved), "revision": revision}
        if edit_batch:
            audit_metadata.update({
                "provider_edit_batch_id": edit_batch.id,
                "changed_target_count": edit_batch.item_count,
                "changed_target_ids": [item["target_id"] for item in changes],
                "changes_sha256": edit_batch.changes_sha256,
            })
        write_audit(request, "SECTION_VALUES_SAVED", "Submission", submission.id, after=audit_metadata)
        response_data = {
            "saved": len(saved), "revision": revision,
            "client_save_id": str(client_save_id), "base_revision": base_revision,
            "resulting_revision": revision, "persisted_change_version": persisted_change_version,
            "replayed": False,
            "last_edited_by": request.user.id, "last_edited_by_name": request.user.name,
            "last_edited_at": submission.last_edited_at.isoformat() if submission.last_edited_at else None,
            "provider_edit_batch_id": edit_batch.id if edit_batch else None,
            **readiness,
        }
        replay_response = json.loads(json.dumps(response_data, cls=DjangoJSONEncoder))
        SectionSaveReceipt.objects.create(
            submission=submission, client_save_id=client_save_id,
            section_code=section_code, actor=request.user,
            payload_sha256=payload_sha256, base_revision=base_revision,
            resulting_revision=revision, persisted_change_version=persisted_change_version,
            response=replay_response,
        )
        return Response(response_data)


class SubmissionCompletionView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        readiness = refresh_submission_completion(submission)
        open_items = CorrectionItem.objects.filter(
            Q(source_submission=submission, stage="PROVIDER_APPROVAL")
            | Q(resolution_submission=submission, stage="NCA_REVIEW"), status="OPEN"
        )
        open_data = [{
            "id": item.id, "stage": item.stage, "target_type": item.target_type,
            "target_id": item.target_id, "instruction": item.instruction,
        } for item in open_items]
        return Response({
            **readiness,
            "transition_ready": readiness["can_submit"] and not open_data,
            "open_correction_item_count": len(open_data),
            "open_correction_items": open_data,
        })


class SubmitForApprovalView(APIView):
    permission_classes = [IsProviderDataEntry]

    @transaction.atomic
    def post(self, request, pk):
        visible = get_submission_for_user(request.user, pk=pk)
        submission = lock_submission(visible.pk)
        if submission.expected.workflow_status not in ("NOT_STARTED", "DRAFT", "PROVIDER_CHANGES_REQUESTED"):
            return Response(
                {"detail": "Only an editable draft can be sent for provider approval."},
                status=400,
            )
        if submission.expected.workflow_status == "CORRECTION_REQUESTED" and not submission.supersedes_id:
            return Response({"detail": "Official historical versions cannot be resubmitted."}, status=409)
        readiness = refresh_submission_completion(submission)
        if not readiness["can_submit"]:
            return Response({
                "code": "SUBMISSION_BLOCKED",
                "detail": "Resolve blocking validation, declaration, upload or security issues before submitting. Blank indicators are allowed.",
                **readiness,
            }, status=400)
        open_items = CorrectionItem.objects.filter(
            Q(source_submission=submission, stage="PROVIDER_APPROVAL")
            | Q(resolution_submission=submission, stage="NCA_REVIEW"),
            status="OPEN",
        )
        if open_items.exists():
            return Response({
                "detail": "Address every correction item before resubmitting.",
                "open_correction_item_ids": list(open_items.values_list("id", flat=True)),
            }, status=409)
        prior_status = submission.expected.workflow_status
        if prior_status == "NOT_STARTED":
            submission.expected.workflow_status = "DRAFT"
            submission.expected.save(update_fields=["workflow_status"])
            audit_transition(
                request=request, submission=submission, event_type="SUBMISSION_STARTED",
                message="The form was opened for data entry before submission to the Provider Approver.",
                from_status="NOT_STARTED", to_status="DRAFT", audience="BOTH",
            )
            prior_status = "DRAFT"
        is_resubmission = prior_status in {"PROVIDER_CHANGES_REQUESTED", "CORRECTION_REQUESTED"} or submission.supersedes_id is not None
        submission.expected.workflow_status = "PROVIDER_RESUBMITTED" if is_resubmission else "PENDING_APPROVAL"
        submission.expected.provider_status = "AWAITING_APPROVAL"
        submission.expected.save(update_fields=["workflow_status", "provider_status"])
        handoff_event = audit_transition(
            request=request, submission=submission,
            event_type="PROVIDER_RESUBMITTED" if is_resubmission else "SUBMITTED_FOR_APPROVAL",
            message="The corrected form was resubmitted to the Provider Approver." if is_resubmission else "The form was submitted to the Provider Approver.",
            from_status=prior_status, to_status=submission.expected.workflow_status,
            audience="PROVIDER", notify=["PROVIDER_APPROVER"],
        )
        return Response({"detail": "Submitted for provider approval.", "workflow_status": submission.expected.workflow_status,
                         "event_id": handoff_event.id, "internal_notification_created": True})


class OfficialSubmitView(APIView):
    permission_classes = [IsProviderApprover]

    @transaction.atomic
    def post(self, request, pk):
        visible = get_submission_for_user(request.user, pk=pk)
        submission = lock_submission(visible.pk)
        if submission.expected.workflow_status not in ("PENDING_APPROVAL", "PROVIDER_RESUBMITTED", "CORRECTION_REQUESTED"):
            return Response({"detail": "Only a form awaiting provider approval can be officially submitted."}, status=400)
        readiness = refresh_submission_completion(submission)
        if not readiness["can_submit"]:
            return Response({
                "code": "SUBMISSION_BLOCKED",
                "detail": "Resolve blocking validation, declaration, upload or security issues before official submission. Blank indicators are allowed.",
                **readiness,
            }, status=409)
        attestation = request.data.get("attestation")
        if attestation is not True:
            return Response({
                "code": "ATTESTATION_REQUIRED",
                "detail": "Confirm the accuracy attestation before official submission.",
            }, status=400)
        open_items = CorrectionItem.objects.filter(
            Q(source_submission=submission, stage="PROVIDER_APPROVAL")
            | Q(resolution_submission=submission, stage="NCA_REVIEW"),
            status="OPEN",
        )
        if open_items.exists():
            return Response({
                "code": "OPEN_CORRECTIONS",
                "detail": "Address every correction item before official submission.",
                "open_correction_item_ids": list(open_items.values_list("id", flat=True)),
            }, status=409)
        latest_handoff = submission.timeline_events.filter(
            event_type__in=["SUBMITTED_FOR_APPROVAL", "PROVIDER_RESUBMITTED"]
        ).order_by("-created_at").first()
        edit_batches = submission.provider_edit_batches.all()
        if latest_handoff:
            edit_batches = edit_batches.filter(created_at__gte=latest_handoff.created_at)
        edit_batch_count = edit_batches.count()
        change_summary = str(request.data.get("change_summary") or "").strip()
        approval_note = str(request.data.get("approval_note") or "").strip()
        prior_status = submission.expected.workflow_status
        was_correction = submission.supersedes_id is not None
        submission.expected.workflow_status = "RESUBMITTED" if was_correction else "SUBMITTED"
        submission.expected.provider_status = "CLOSED"
        submission.expected.due_state = submission.expected.compute_due_state()
        submission.expected.save(update_fields=["workflow_status", "provider_status", "due_state"])
        submission.submitted_by = request.user
        submission.submitted_at = timezone.now()
        submission.regulatory_status = "SUBMITTED"
        submission.save(update_fields=["submitted_by", "submitted_at", "regulatory_status"])
        from .form_snapshots import snapshot_submission
        snapshot_submission(submission)
        decision = ProviderApprovalDecision.objects.create(
            submission=submission, approver=request.user, attestation=True,
            approval_note=approval_note, change_summary=change_summary,
            approver_edited=bool(edit_batch_count), edit_batch_count=edit_batch_count,
        )
        from .receipts import create_receipt
        receipt = create_receipt(submission)
        TransactionalOutbox.objects.get_or_create(topic="submission.official", aggregate_type="Submission",
            aggregate_id=str(submission.id), idempotency_key=f"official-submission:{submission.id}:{submission.version}",
            defaults={"payload": {"submission_id": submission.id, "receipt_reference": receipt.reference}})
        official_event = audit_transition(
            request=request, submission=submission, event_type="OFFICIALLY_SUBMITTED",
            message="The Provider Approver officially resubmitted the corrected form to NCA." if was_correction else "The Provider Approver officially submitted the form to NCA.",
            from_status=prior_status, to_status=submission.expected.workflow_status,
            audience="BOTH", notify=["PROVIDER_DATA_ENTRY", "NCA_REVIEWER"],
            metadata={
                "receipt_reference": receipt.reference,
                "provider_approval_decision_id": decision.id,
                "provider_edit_batch_count": edit_batch_count,
                "comments": approval_note,
                "action_required": "NCA will review the submitted form.",
            },
        )
        from apps.compliance.communications import queue_automatic_communication
        queue_automatic_communication(
            submission=submission, event=official_event, action="NCA_ACKNOWLEDGEMENT",
            actor=request.user, payload={"comments": approval_note},
        )
        baseline = ProviderWorkbookBaseline.objects.filter(
            provider_id=submission.expected.provider_id,
            form_template_id=submission.expected.form_template_id,
            status="ACTIVE", scan_status="CLEAN",
        ).first()
        if baseline:
            if was_correction and submission.supersedes_id:
                MonthlyReportArtifact.objects.filter(
                    submission_id=submission.supersedes_id, status="READY",
                ).update(status="SUPERSEDED")
            MonthlyReportArtifact.objects.update_or_create(
                submission=submission,
                defaults={
                    "baseline": baseline, "status": "PREPARING", "error_message": "",
                    "submission_revision": submission.revision,
                },
            )

            def queue_monthly_report():
                from .tasks import generate_monthly_report_task
                try:
                    generate_monthly_report_task.delay(submission.id)
                except Exception as exc:
                    MonthlyReportArtifact.objects.filter(submission=submission).update(
                        status="FAILED", error_message=f"Generation could not be queued: {exc}",
                    )

            transaction.on_commit(queue_monthly_report)
        return Response({
            "detail": "Officially submitted to NCA.",
            "workflow_status": submission.expected.workflow_status,
            "receipt_reference": receipt.reference,
            "provider_approval_decision_id": decision.id,
            "monthly_report_status": "PREPARING" if baseline else "NOT_CONFIGURED",
            "event_id": official_event.id,
            "internal_notification_created": True,
        })


class PreviousMonthValuesView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        return Response(previous_month_values(submission))


class MonthlyReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        artifact = MonthlyReportArtifact.objects.filter(submission=submission).first()
        if artifact:
            return Response(MonthlyReportArtifactSerializer(artifact).data)
        baseline = ProviderWorkbookBaseline.objects.filter(
            provider_id=submission.expected.provider_id,
            form_template_id=submission.expected.form_template_id,
            status="ACTIVE",
        ).first()
        return Response({
            "status": "NOT_SUBMITTED" if baseline else "NOT_CONFIGURED",
            "download_ready": False,
            "detail": (
                "The Provider Approver must officially submit this form before the Excel report is generated."
                if baseline else "NCA has not configured an approved provider-specific workbook baseline for this form."
            ),
        })


class MonthlyReportRetryView(APIView):
    permission_classes = [IsNCAEditor]

    def post(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        artifact = get_object_or_404(MonthlyReportArtifact, submission=submission)
        if artifact.status != "FAILED":
            return Response({"detail": "Only a failed monthly report can be retried."}, status=409)
        artifact.status = "PREPARING"
        artifact.error_message = ""
        artifact.save(update_fields=["status", "error_message", "updated_at"])
        from .tasks import generate_monthly_report_task
        try:
            generate_monthly_report_task.delay(submission.id)
        except Exception as exc:
            artifact.status = "FAILED"
            artifact.error_message = f"Generation could not be queued: {exc}"
            artifact.save(update_fields=["status", "error_message", "updated_at"])
            return Response({"detail": artifact.error_message}, status=503)
        write_audit(request, "MONTHLY_REPORT_RETRY", "MonthlyReportArtifact", artifact.id)
        return Response({"detail": "Monthly report generation restarted."}, status=202)


class MonthlyReportDownloadView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        artifact = get_object_or_404(MonthlyReportArtifact, submission=submission)
        if artifact.status != "READY" or not artifact.private_path:
            return Response({"detail": "The monthly Excel report is not ready."}, status=409)
        path = Path(artifact.private_path)
        try:
            path.resolve().relative_to(Path(settings.PRIVATE_EXPORT_ROOT).resolve())
        except ValueError:
            return Response({"detail": "The report path is invalid."}, status=410)
        if not path.exists():
            return Response({"detail": "The private report file is missing."}, status=410)
        write_audit(
            request, "MONTHLY_REPORT_DOWNLOADED", "MonthlyReportArtifact", artifact.id,
            after={"submission_id": submission.id, "sha256": artifact.sha256},
        )
        return FileResponse(
            open(path, "rb"), as_attachment=True, filename=artifact.filename,
            content_type=artifact.mime_type,
        )


class ProviderWorkbookBaselineListCreateView(APIView):
    permission_classes = [IsNCAEditor]

    def get(self, request, template_pk):
        rows = ProviderWorkbookBaseline.objects.filter(form_template_id=template_pk).select_related(
            "provider", "form_template", "contact", "created_by", "approved_by",
        )
        return Response(ProviderWorkbookBaselineSerializer(rows, many=True).data)

    @transaction.atomic
    def post(self, request, template_pk):
        from apps.forms_engine.models import FormTemplate
        from apps.providers.models import ProviderProfile, ProviderContact
        from apps.uploads.scanner import scan_path

        template = get_object_or_404(FormTemplate, pk=template_pk)
        provider = get_object_or_404(ProviderProfile, pk=request.data.get("provider"), status="ACTIVE")
        upload = request.FILES.get("file")
        if not upload or not upload.name.lower().endswith(".xlsx"):
            return Response({"detail": "A private .xlsx workbook is required."}, status=400)
        if upload.size > 20 * 1024 * 1024:
            return Response({"detail": "Workbook files may not exceed 20 MB."}, status=400)
        contact = None
        if request.data.get("contact"):
            contact = get_object_or_404(ProviderContact, pk=request.data["contact"], provider=provider, is_active=True)
        version = (
            ProviderWorkbookBaseline.objects.filter(provider=provider, form_template=template)
            .order_by("-version").values_list("version", flat=True).first() or 0
        ) + 1
        relative = Path("provider-report-baselines") / str(provider.provider_id) / f"{uuid.uuid4().hex}.xlsx"
        full_path = Path(settings.PRIVATE_UPLOAD_ROOT) / relative
        full_path.parent.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256()
        with open(full_path, "wb") as destination:
            for chunk in upload.chunks():
                destination.write(chunk)
                digest.update(chunk)
        scan = scan_path(full_path)
        row = ProviderWorkbookBaseline.objects.create(
            provider=provider, form_template=template, contact=contact, version=version,
            file_name=os.path.basename(upload.name), storage_path=str(relative), file_size=upload.size,
            sha256=digest.hexdigest(), scan_status=scan["status"], scan_engine=scan["engine"],
            scan_details=scan["details"], main_sheet=request.data.get("main_sheet") or "REVISED MNOs MONTHLY DATA",
            created_by=request.user,
        )
        if row.scan_status == "CLEAN":
            try:
                suggest_exact_baseline_mappings(row)
            except Exception as exc:
                row.mapping_summary = {"total": 0, "error": str(exc)}
                row.save(update_fields=["mapping_summary", "updated_at"])
        write_audit(request, "PROVIDER_WORKBOOK_BASELINE_UPLOADED", "ProviderWorkbookBaseline", row.id,
            after={"provider_id": provider.id, "form_template_id": template.id, "sha256": row.sha256})
        return Response(ProviderWorkbookBaselineSerializer(row).data, status=201)


class ProviderWorkbookBaselineDetailView(APIView):
    permission_classes = [IsNCAEditor]

    def get(self, request, pk):
        row = get_object_or_404(ProviderWorkbookBaseline, pk=pk)
        return Response(ProviderWorkbookBaselineSerializer(row).data)

    @transaction.atomic
    def patch(self, request, pk):
        row = get_object_or_404(ProviderWorkbookBaseline.objects.select_for_update(), pk=pk)
        if row.status != "DRAFT":
            return Response({"detail": "Only draft workbook baselines can be changed."}, status=409)
        if "mappings" in request.data:
            try:
                replace_baseline_mappings(row, request.data.get("mappings") or [])
            except (TypeError, ValueError) as exc:
                return Response({"detail": str(exc)}, status=400)
        for field in ("main_sheet", "month_header_row", "first_month_column", "contact_cells"):
            if field in request.data:
                setattr(row, field, request.data[field])
        row.save()
        write_audit(request, "PROVIDER_WORKBOOK_BASELINE_UPDATED", "ProviderWorkbookBaseline", row.id,
            after={"mapping_summary": row.mapping_summary})
        return Response(ProviderWorkbookBaselineSerializer(row).data)


class ProviderWorkbookBaselineApproveView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        row = get_object_or_404(ProviderWorkbookBaseline.objects.select_for_update(), pk=pk)
        if row.status != "DRAFT":
            return Response({"detail": "Only a draft baseline can be approved."}, status=409)
        issues = [issue for issue in baseline_readiness(row) if "not active" not in issue]
        required_field_ids = set(row.form_template.sections.filter(
            fields__is_required=True,
        ).values_list("fields__id", flat=True))
        mapped_required_ids = set(row.indicator_mappings.filter(
            field_id__in=required_field_ids,
        ).values_list("field_id", flat=True))
        missing_required = sorted(required_field_ids - mapped_required_ids)
        if missing_required:
            issues.append(f"Required form fields are not mapped: {missing_required}.")
        if issues:
            return Response({"detail": "Workbook baseline approval is blocked.", "issues": issues}, status=409)
        ProviderWorkbookBaseline.objects.filter(
            provider=row.provider, form_template=row.form_template, status="ACTIVE",
        ).exclude(pk=row.pk).update(status="ARCHIVED")
        row.status = "ACTIVE"
        row.approved_by = request.user
        row.approved_at = timezone.now()
        row.save(update_fields=["status", "approved_by", "approved_at", "updated_at"])
        write_audit(request, "PROVIDER_WORKBOOK_BASELINE_APPROVED", "ProviderWorkbookBaseline", row.id,
            after={"sha256": row.sha256, "mapping_summary": row.mapping_summary})
        return Response(ProviderWorkbookBaselineSerializer(row).data)


# ── NCA Review ────────────────────────────────────────────────────────────────

class ReviewHistoryView(generics.ListAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = ReviewActionSerializer

    def get_queryset(self):
        submission = get_submission_for_user(self.request.user, pk=self.kwargs["pk"])
        queryset = ReviewAction.objects.filter(submission=submission).select_related("created_by")
        if self.request.user.is_provider:
            queryset = queryset.filter(is_provider_visible=True)
        return queryset


class StartReviewView(APIView):
    permission_classes = [IsNCAEditor]
    @transaction.atomic
    def post(self, request, pk):
        submission = lock_submission(pk)
        if submission.regulatory_status == "DRAFT" and submission.expected.workflow_status in {"SUBMITTED", "RESUBMITTED"}:
            submission.regulatory_status = "SUBMITTED"
            submission.save(update_fields=["regulatory_status"])
        if submission.regulatory_status != "SUBMITTED":
            return Response({"detail": "Only submitted returns can enter review."}, status=409)
        prior_status = submission.expected.workflow_status
        submission.expected.workflow_status = "UNDER_REVIEW"; submission.expected.save(update_fields=["workflow_status"])
        submission.regulatory_status = "UNDER_REVIEW"
        submission.save(update_fields=["regulatory_status"])
        ReviewAction.objects.create(submission=submission, action="ADD_NOTE", comment="NCA review started.", created_by=request.user)
        event = audit_transition(request=request, submission=submission, event_type="SUBMISSION_REVIEW_STARTED",
            message="NCA regulatory review started.", from_status=prior_status,
            to_status="UNDER_REVIEW", audience="BOTH", notify=["PROVIDER_APPROVER"])
        return Response({"detail": "Review started.", "workflow_status": "UNDER_REVIEW",
                         "event_id": event.id, "internal_notification_created": True})


class ReviewApproveView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        try: submission = lock_submission(pk)
        except Submission.DoesNotExist: return Response({"detail": "Not found."}, status=404)
        if submission.regulatory_status != "UNDER_REVIEW" and submission.expected.workflow_status == "UNDER_REVIEW":
            submission.regulatory_status = "UNDER_REVIEW"
            submission.save(update_fields=["regulatory_status"])
        if submission.regulatory_status != "UNDER_REVIEW":
            return Response({"detail": "Start the regulatory review before approving this submission."}, status=409)
        readiness = refresh_submission_completion(submission)
        if not readiness["can_submit"]:
            return Response({"detail": "Blocking validation issues must be resolved before approval.", **readiness}, status=409)
        non_filled = submission.values.filter(value_status__in=("NOT_APPLICABLE", "NOT_AVAILABLE", "NOT_REQUIRED"))
        unresolved = non_filled.filter(Q(non_filled_disposition__isnull=True) | Q(non_filled_disposition__decision="REJECTED"))
        if unresolved.exists():
            return Response({"detail": "Every non-filled explanation must have an Accepted disposition before approval.", "unresolved_value_ids": list(unresolved.values_list("id", flat=True))}, status=409)
        open_corrections = submission.resolved_correction_items.filter(status="OPEN")
        if open_corrections.exists():
            return Response({"detail": "Every requested correction must be verified before approval.", "correction_item_ids": list(open_corrections.values_list("id", flat=True))}, status=409)
        rejected_or_pending_kmz = submission.kmz_uploads.filter(requirement__is_required=True).exclude(scan_status="CLEAN", review_status="ACCEPTED")
        if rejected_or_pending_kmz.exists():
            return Response({"detail": "Every required KMZ upload must be clean and accepted before approval."}, status=409)
        from apps.forms_engine.gaps import recalculate_form_gaps
        objective_gaps = [gap for gap in recalculate_form_gaps(submission.expected.form_template, request.user)
            if gap.status != "MATCHED" and gap.requirement.severity in {"BLOCKER", "HIGH"}]
        if objective_gaps:
            return Response({"detail": "The exact form version has unresolved blocker/high Section 11 requirements.",
                "form_gap_ids": [gap.id for gap in objective_gaps]}, status=409)
        prior_status = submission.expected.workflow_status
        submission.expected.workflow_status = "APPROVED"
        submission.expected.provider_status = "CLOSED"
        submission.expected.refresh_due_state()
        submission.reviewed_by = request.user
        submission.reviewed_at = timezone.now()
        submission.regulatory_status = "APPROVED"
        submission.save(update_fields=["reviewed_by", "reviewed_at", "regulatory_status"])
        submission.expected.save(update_fields=["workflow_status", "provider_status", "due_state"])
        submission.resolved_correction_items.filter(status="ADDRESSED").update(status="VERIFIED")
        ReviewAction.objects.create(
            submission=submission, action="APPROVE",
            comment=request.data.get("comment", ""), created_by=request.user,
        )
        event = audit_transition(request=request, submission=submission, event_type="SUBMISSION_APPROVED",
            message="NCA approved the official submission.", from_status=prior_status,
            to_status="APPROVED", audience="BOTH", notify=["PROVIDER_DATA_ENTRY", "PROVIDER_APPROVER"],
            metadata={"comments": str(request.data.get("comment") or "").strip(),
                      "action_required": "No further action is required."})
        return Response({"detail": "Submission approved.", "event_id": event.id, "internal_notification_created": True})


class ReviewRejectView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        try: submission = lock_submission(pk)
        except Submission.DoesNotExist: return Response({"detail": "Not found."}, status=404)
        if submission.regulatory_status != "UNDER_REVIEW" and submission.expected.workflow_status == "UNDER_REVIEW":
            submission.regulatory_status = "UNDER_REVIEW"
            submission.save(update_fields=["regulatory_status"])
        if submission.regulatory_status != "UNDER_REVIEW":
            return Response({"detail": "Start the regulatory review before rejecting this submission."}, status=409)
        if not request.data.get("comment", "").strip():
            return Response({"detail": "A rejection reason is required."}, status=400)
        comment = request.data.get("comment", "").strip()
        prior_status = submission.expected.workflow_status
        submission.expected.workflow_status = "REJECTED"
        submission.expected.provider_status = "CLOSED"
        submission.expected.save(update_fields=["workflow_status", "provider_status"])
        submission.reviewed_by = request.user
        submission.reviewed_at = timezone.now()
        submission.regulatory_status = "REJECTED"
        submission.save(update_fields=["reviewed_by", "reviewed_at", "regulatory_status"])
        ReviewAction.objects.create(
            submission=submission, action="REJECT",
            comment=comment, is_provider_visible=True, created_by=request.user,
        )
        event = audit_transition(request=request, submission=submission, event_type="SUBMISSION_REJECTED",
            message=f"NCA rejected the submission: {comment}", from_status=prior_status,
            to_status="REJECTED", audience="BOTH", notify=["PROVIDER_DATA_ENTRY", "PROVIDER_APPROVER"],
            metadata={"reason": comment, "action_required": "Review the decision and contact NCA if clarification is needed."})
        return Response({"detail": "Submission rejected.", "event_id": event.id, "internal_notification_created": True})


class ReviewRequestCorrectionView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        try: submission = lock_submission(pk)
        except Submission.DoesNotExist: return Response({"detail": "Not found."}, status=404)
        if submission.regulatory_status != "UNDER_REVIEW" and submission.expected.workflow_status == "UNDER_REVIEW":
            submission.regulatory_status = "UNDER_REVIEW"
            submission.save(update_fields=["regulatory_status"])
        if submission.regulatory_status != "UNDER_REVIEW":
            return Response({"detail": "Start the regulatory review before requesting corrections."}, status=409)
        targets = request.data.get("targets", [])  # [{type, id, reason|comment}]
        comment = str(request.data.get("comment", "")).strip()
        if not targets:
            return Response({"detail": "Select at least one section or indicator to flag."}, status=400)
        missing_reasons = [
            str(target.get("id", "")) for target in targets
            if not str(target.get("reason") or target.get("comment") or "").strip()
        ]
        if missing_reasons:
            return Response({"detail": "Every flagged section or indicator requires its own reason.", "targets": missing_reasons}, status=400)
        target_error = validate_correction_targets(submission.expected.form_template, targets)
        if target_error:
            return Response({"detail": target_error}, status=400)

        correction_submission = clone_for_nca_correction(submission)
        prior_status = submission.expected.workflow_status
        submission.expected.workflow_status = "CORRECTION_REQUESTED"
        submission.expected.provider_status = "AWAITING_APPROVAL"
        submission.expected.save(update_fields=["workflow_status", "provider_status"])
        submission.regulatory_status = "RETURNED_FOR_CORRECTION"
        submission.save(update_fields=["regulatory_status"])

        # Mark targeted fields as WAITING_CORRECTION
        for t in targets:
            if t.get("type") == "FIELD" and t.get("id"):
                SubmissionValue.objects.filter(
                    submission=submission, field_id=t["id"]
                ).update(value_status="WAITING_CORRECTION")

        ReviewAction.objects.create(
            submission=submission, action="REQUEST_CORRECTION",
            target_type="SUBMISSION", comment=comment,
            is_provider_visible=True, created_by=request.user,
        )
        for t in targets:
            target_reason = str(t.get("reason") or t.get("comment") or "").strip()
            ReviewAction.objects.create(
                submission=submission, action="REQUEST_CORRECTION",
                target_type=t.get("type", "FIELD"),
                target_id=str(t.get("id", "")),
                comment=target_reason,
                is_provider_visible=True, created_by=request.user,
            )
            CorrectionItem.objects.create(source_submission=submission, resolution_submission=correction_submission,
                stage="NCA_REVIEW", target_type=t.get("type", "FIELD"), target_id=str(t.get("id", "")),
                instruction=target_reason, created_by=request.user)
        target_labels = [{
            "label": correction_target_label(submission.expected.form_template, target.get("type"), target.get("id")),
            "reason": str(target.get("reason") or target.get("comment") or "").strip(),
        } for target in targets]
        summary = comment or "; ".join(item["reason"] for item in target_labels)
        event = audit_transition(request=request, submission=submission, event_type="CORRECTION_REQUESTED",
            message=f"NCA flagged {len(targets)} item(s) for review.", from_status=prior_status,
            to_status="CORRECTION_REQUESTED", audience="BOTH", notify=["PROVIDER_DATA_ENTRY", "PROVIDER_APPROVER"],
            metadata={"targets": targets, "target_labels": target_labels, "correction_submission_id": correction_submission.id,
                      "reason": summary, "comments": summary,
                      "action_required": "Review the flagged items and resubmit the form."})
        return Response({"detail": "Flagged items returned for review.", "targets": len(targets),
                         "correction_submission_id": correction_submission.id, "event_id": event.id,
                         "internal_notification_created": True})


class ReviewAddNoteView(APIView):
    permission_classes = [IsNCAEditor]

    @transaction.atomic
    def post(self, request, pk):
        submission = get_submission_for_user(request.user, pk=pk)
        is_provider = request.data.get("provider_visible", False)
        action = "ADD_PROVIDER_COMMENT" if is_provider else "ADD_NOTE"
        ReviewAction.objects.create(
            submission=submission, action=action,
            comment=request.data.get("comment", ""),
            is_provider_visible=is_provider, created_by=request.user,
        )
        comment = str(request.data.get("comment") or "").strip()
        if not comment:
            return Response({"detail": "A note is required."}, status=400)
        event = audit_transition(
            request=request,
            submission=submission,
            event_type="SUBMISSION_CORRESPONDENCE" if is_provider else "SUBMISSION_INTERNAL_NOTE",
            message=comment,
            from_status=submission.expected.workflow_status,
            to_status=submission.expected.workflow_status,
            audience="BOTH" if is_provider else "INTERNAL",
            notify=["PROVIDER_APPROVER"] if is_provider else [],
            metadata={"internal_note": not is_provider, "provider_visible": bool(is_provider)},
        )
        return Response({"detail": "Note added.", "event_id": event.id,
                         "internal_notification_created": bool(is_provider)})


# ── Expected Submission Management ───────────────────────────────────────────

class AssignOfficerView(APIView):
    permission_classes = [IsNCAEditor]

    def patch(self, request, pk):
        try:
            expected = ExpectedSubmission.objects.get(pk=pk)
        except ExpectedSubmission.DoesNotExist:
            return Response({"detail": "Not found."}, status=404)

        officer_id = request.data.get("assigned_officer")
        before = expected.assigned_officer_id

        if officer_id is None:
            expected.assigned_officer = None
        else:
            from apps.users.models import User
            try:
                officer = User.objects.get(pk=officer_id, role__in=["NCA_ADMIN", "NCA_OFFICER"])
            except User.DoesNotExist:
                return Response({"detail": "Officer not found or not an NCA user."}, status=400)
            expected.assigned_officer = officer

        expected.save(update_fields=["assigned_officer"])
        write_audit(request, "OFFICER_ASSIGNED", "ExpectedSubmission", pk,
                    before={"assigned_officer": before},
                    after={"assigned_officer": officer_id})
        return Response(ExpectedSubmissionSerializer(expected).data)


class OverrideDueDateView(APIView):
    permission_classes = [IsNCAEditor]

    def patch(self, request, pk):
        try:
            expected = ExpectedSubmission.objects.get(pk=pk)
        except ExpectedSubmission.DoesNotExist:
            return Response({"detail": "Not found."}, status=404)

        proposed = parse_datetime(request.data.get("due_at_override", ""))
        reason = request.data.get("reason", "").strip()
        if not proposed or proposed <= timezone.now() or not reason: return Response({"detail": "A future due_at_override and reason are required."}, status=400)
        change = DeadlineChangeRequest.objects.create(expected_submission=expected, previous_due_at=expected.effective_due_at,
            proposed_due_at=proposed, reason=reason, requested_by=request.user)
        write_audit(request, "DEADLINE_CHANGE_REQUESTED", "DeadlineChangeRequest", change.id, after={"proposed_due_at": proposed.isoformat()})
        return Response({"id": change.id, "status": change.status}, status=201)


class DeadlineDecisionView(APIView):
    permission_classes = [IsNCAEditor]
    def post(self, request, pk):
        change = get_object_or_404(DeadlineChangeRequest, pk=pk, status="PENDING")
        decision = request.data.get("decision")
        if decision not in ("APPROVED", "REJECTED"): return Response({"detail": "decision must be APPROVED or REJECTED."}, status=400)
        change.status=decision; change.decided_by=request.user; change.decided_at=timezone.now(); change.decision_note=request.data.get("note", ""); change.save()
        change.expected_submission.refresh_due_state()
        write_audit(request, f"DEADLINE_CHANGE_{decision}", "DeadlineChangeRequest", change.id)
        return Response({"id": change.id, "status": change.status})


class SubmissionOverrideView(APIView):
    permission_classes = [IsNCAEditor]
    def post(self, request, pk):
        expected = get_object_or_404(ExpectedSubmission, pk=pk)
        blocker_ids = request.data.get("blocker_ids", []); reason=request.data.get("reason", "").strip(); expires=parse_datetime(request.data.get("expires_at", ""))
        if not blocker_ids or not reason or not expires or expires <= timezone.now(): return Response({"detail": "blocker_ids, reason and a future expires_at are required."}, status=400)
        item=SubmissionOverride.objects.create(expected_submission=expected, blocker_ids=blocker_ids, reason=reason,
            evidence_reference=request.data.get("evidence_reference", ""), expires_at=expires, requested_by=request.user)
        write_audit(request, "SUBMISSION_OVERRIDE_REQUESTED", "SubmissionOverride", item.id)
        return Response({"id": item.id, "status": item.status}, status=201)


class SubmissionOverrideDecisionView(APIView):
    permission_classes = [IsNCAEditor]
    def post(self, request, pk):
        item=get_object_or_404(SubmissionOverride, pk=pk, status="PENDING"); decision=request.data.get("decision")
        if decision not in ("APPROVED", "REJECTED"): return Response({"detail":"decision must be APPROVED or REJECTED."}, status=400)
        item.status=decision; item.approved_by=request.user if decision=="APPROVED" else None; item.approved_at=timezone.now(); item.decision_note=request.data.get("note",""); item.save()
        write_audit(request, f"SUBMISSION_OVERRIDE_{decision}", "SubmissionOverride", item.id)
        return Response({"id":item.id,"status":item.status})


class NonFilledDispositionView(APIView):
    permission_classes = [IsNCAEditor]
    def post(self, request, pk, value_id):
        submission=get_object_or_404(Submission, pk=pk); value=get_object_or_404(SubmissionValue, pk=value_id, submission=submission)
        if value.value_status not in ("NOT_APPLICABLE","NOT_AVAILABLE","NOT_REQUIRED"): return Response({"detail":"Only non-filled values require disposition."}, status=400)
        decision=request.data.get("decision")
        if decision not in ("ACCEPTED","REJECTED"): return Response({"detail":"decision must be ACCEPTED or REJECTED."}, status=400)
        disposition,_=NonFilledDisposition.objects.update_or_create(value=value, defaults={"decision":decision,"note":request.data.get("note", ""),"reviewed_by":request.user})
        write_audit(request, f"NON_FILLED_{decision}", "SubmissionValue", value.id)
        return Response({"id":disposition.id,"decision":disposition.decision})


class CorrectionDiffView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, pk):
        current=get_submission_for_user(request.user, pk=pk); previous=current.expected.versions.filter(version__lt=current.version).order_by("-version").first()
        if not previous: return Response([])
        def key(v): return (v.field_id, v.grid_id, v.grid_row_id, v.grid_column_id)
        before={key(v):v for v in previous.values.all()}; after={key(v):v for v in current.values.all()}; rows=[]
        for item_key in sorted(set(before)|set(after), key=str):
            old,new=before.get(item_key),after.get(item_key)
            if (old.value if old else None)!=(new.value if new else None) or (old.value_status if old else None)!=(new.value_status if new else None):
                rows.append({"target":item_key,"before_value":old.value if old else None,"after_value":new.value if new else None,
                    "before_status":old.value_status if old else None,"after_status":new.value_status if new else None})
        return Response(rows)


class ReceiptDownloadView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request, pk):
        submission=get_submission_for_user(request.user, pk=pk)
        if not hasattr(submission,"receipt"): return Response({"detail":"Receipt not available."}, status=404)
        import os
        if not os.path.exists(submission.receipt.private_path): return Response({"detail":"Receipt file is missing; contact support."},status=410)
        write_audit(request,"SUBMISSION_RECEIPT_DOWNLOADED","SubmissionReceipt",submission.receipt.id)
        return FileResponse(open(submission.receipt.private_path,"rb"),as_attachment=True,filename=f"{submission.receipt.reference}.pdf",content_type="application/pdf")


class ReminderPolicyListCreate(generics.ListCreateAPIView):
    permission_classes = [IsNCAEditor]
    serializer_class = None
    def get(self, request, period_id):
        rows=ReminderPolicy.objects.filter(period_id=period_id).order_by("-version")
        return Response([{"id":x.id,"version":x.version,"name":x.name,"rules":x.rules,"status":x.status} for x in rows])
    def post(self, request, period_id):
        period=get_object_or_404(ReportingPeriod, pk=period_id); version=(period.reminder_policies.order_by("-version").values_list("version",flat=True).first() or 0)+1
        rules=request.data.get("rules",[])
        if not isinstance(rules,list) or not rules: return Response({"detail":"At least one reminder rule is required."},status=400)
        for rule in rules:
            if not isinstance(rule,dict) or not isinstance(rule.get("offset_days"),int) or not isinstance(rule.get("recipient_roles",[]),list) or not rule.get("template_version"):
                return Response({"detail":"Each rule requires integer offset_days, recipient_roles and template_version."},status=400)
        row=ReminderPolicy.objects.create(period=period,version=version,name=request.data.get("name",f"Policy v{version}"),rules=rules,prepared_by=request.user)
        return Response({"id":row.id,"version":row.version,"status":row.status},status=201)


class ApproveReminderPolicyView(APIView):
    permission_classes = [IsNCAEditor]
    def post(self, request, pk):
        row=get_object_or_404(ReminderPolicy,pk=pk,status="DRAFT")
        if row.prepared_by_id==request.user.id: return Response({"detail":"Maker/checker approval requires a different Admin."},status=409)
        row.period.reminder_policies.filter(status="APPROVED").update(status="ARCHIVED")
        row.status="APPROVED"; row.approved_by=request.user; row.approved_at=timezone.now(); row.save(update_fields=["status","approved_by","approved_at"])
        write_audit(request,"REMINDER_POLICY_APPROVED","ReminderPolicy",row.id,after={"version":row.version})
        return Response({"id":row.id,"version":row.version,"status":row.status})


class PeriodAssignedTemplatesView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request, pk):
        period = generics.get_object_or_404(ReportingPeriod, pk=pk)
        from apps.forms_engine.serializers import FormTemplateListSerializer
        return Response(FormTemplateListSerializer(period.applicable_form_templates.all(), many=True).data)


class PeriodAssignedProvidersView(APIView):
    permission_classes = [IsNCAUser]

    def get(self, request, pk):
        period = generics.get_object_or_404(ReportingPeriod, pk=pk)
        from apps.providers.serializers import ProviderProfileListSerializer
        return Response(ProviderProfileListSerializer(period.assigned_providers.all(), many=True).data)


class ProviderRequestCorrectionView(APIView):
    permission_classes = [IsProviderApprover]

    @transaction.atomic
    def post(self, request, pk):
        visible = get_submission_for_user(request.user, pk=pk)
        submission = lock_submission(visible.pk)
        if submission.expected.workflow_status not in ("PENDING_APPROVAL", "PROVIDER_RESUBMITTED", "CORRECTION_REQUESTED"):
            return Response({"detail": "Only a form awaiting provider approval can be returned."}, status=400)
        reason = str(request.data.get("reason") or request.data.get("comment") or "").strip()
        targets = request.data.get("targets", [])
        if not reason:
            return Response({"detail": "A clear correction reason is required."}, status=400)
        if not isinstance(targets, list) or not targets:
            return Response({"detail": "Identify at least one section, field or grid cell to correct."}, status=400)
        target_error = validate_correction_targets(submission.expected.form_template, targets)
        if target_error:
            return Response({"detail": target_error}, status=400)
        prior_status = submission.expected.workflow_status
        for target in targets:
            CorrectionItem.objects.create(
                source_submission=submission,
                resolution_submission=submission,
                stage="PROVIDER_APPROVAL",
                target_type=target["type"],
                target_id=str(target["id"]),
                instruction=str(target.get("instruction") or target.get("comment") or reason).strip(),
                created_by=request.user,
            )
        submission.expected.workflow_status = "PROVIDER_CHANGES_REQUESTED"
        submission.expected.provider_status = "CORRECTIONS_REQUIRED"
        submission.expected.save(update_fields=["workflow_status", "provider_status"])
        event = audit_transition(request=request, submission=submission, event_type="PROVIDER_CHANGES_REQUESTED",
            message=f"The Provider Approver requested corrections: {reason}", from_status=prior_status,
            to_status="PROVIDER_CHANGES_REQUESTED", audience="PROVIDER", notify=["PROVIDER_DATA_ENTRY"],
            metadata={"targets": targets, "reason": reason, "comments": reason,
                      "action_required": "Review the flagged items and update the form."})
        return Response({"detail": "Returned to Data Entry for flagged review.", "workflow_status": "PROVIDER_CHANGES_REQUESTED",
                         "event_id": event.id, "internal_notification_created": True})


class ReturnToDraftView(ProviderRequestCorrectionView):
    """Compatibility wrapper for existing clients; the reason/targets contract is mandatory."""


class SubmissionTimelineView(generics.ListAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = SubmissionEventSerializer
    pagination_class = None

    def get_queryset(self):
        submission = get_submission_for_user(self.request.user, pk=self.kwargs["pk"])
        queryset = SubmissionEvent.objects.filter(submission=submission).select_related("actor")
        if self.request.user.is_provider:
            queryset = queryset.filter(audience__in=["PROVIDER", "BOTH"])
        elif self.request.user.role in {"NCA_ADMIN", "NCA_OFFICER"}:
            queryset = queryset.exclude(audience="INTERNAL")
        else:
            return queryset.none()
        return queryset


class SubmissionNotificationListView(generics.ListAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = SubmissionNotificationSerializer

    def get_queryset(self):
        queryset = SubmissionNotification.objects.filter(recipient=self.request.user).select_related(
            "submission__expected__form_template", "submission__expected__provider", "event"
        )
        if self.request.query_params.get("unread") in {"1", "true", "True"}:
            queryset = queryset.filter(is_read=False)
        return queryset


class SubmissionNotificationSummaryView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = SubmissionNotification.objects.filter(recipient=request.user)
        grouped = {
            row["event__event_type"]: row["total"]
            for row in queryset.filter(is_read=False).values("event__event_type").annotate(total=Count("id"))
        }
        pending_approval = 0
        if request.user.role == "PROVIDER_APPROVER" and request.user.organization_id:
            pending_approval = ExpectedSubmission.objects.filter(
                provider__organization_id=request.user.organization_id,
                workflow_status__in=["PENDING_APPROVAL", "PROVIDER_RESUBMITTED", "CORRECTION_REQUESTED"],
            ).count()
        return Response({
            "unread": queryset.filter(is_read=False).count(),
            "total": queryset.count(),
            "by_event_type": grouped,
            "pending_approval": pending_approval,
        })


class SubmissionNotificationReadView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        notification = get_object_or_404(SubmissionNotification, pk=pk, recipient=request.user)
        was_read = notification.is_read
        mark_notification_read(notification)
        if not was_read:
            write_audit(request, "NOTIFICATION_READ", "SubmissionNotification", notification.id)
        return Response({"id": notification.id, "is_read": True})


class SubmissionNotificationReadAllView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        now = timezone.now()
        count = SubmissionNotification.objects.filter(recipient=request.user, is_read=False).update(is_read=True, read_at=now)
        if count:
            write_audit(
                request, "NOTIFICATIONS_MARKED_READ", "SubmissionNotification", "bulk",
                after={"count": count},
            )
        return Response({"marked_read": count})


class IndustryDashboardExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if request.user.role.startswith("PROVIDER_"):
            return Response({"detail": "Dashboard export is not available to provider accounts."}, status=403)
        if request.query_params.get("format", "xlsx").lower() != "xlsx":
            return Response({"detail": "Only aggregate XLSX export is available."}, status=400)
        from openpyxl import Workbook
        from .industry_dashboard import aggregate_rows, dataset_for_user, load_dashboard_dataset
        dataset = dataset_for_user(load_dashboard_dataset(), request.user)
        workbook = Workbook(write_only=True)
        details = workbook.create_sheet("Export Details")
        details.append(["Dataset", dataset.get("metadata", {}).get("source", "NCA Industry Dashboard")])
        details.append(["Latest period", dataset.get("metadata", {}).get("latestObservedPeriod", "")])
        details.append(["Export scope", "Aggregate chart observations only; no source workbook or provider-level rows."])
        data_sheet = workbook.create_sheet("Aggregate Data")
        headers = ["dashboard", "chart", "unit", "period", "series_kind", "series", "value"]
        data_sheet.append(headers)
        count = 0
        for row in aggregate_rows(dataset):
            data_sheet.append([row[header] for header in headers]); count += 1
        import io, hashlib
        output = io.BytesIO(); workbook.save(output); content = output.getvalue(); digest = hashlib.sha256(content).hexdigest()
        record_audit(user=request.user, action="INDUSTRY_DASHBOARD_AGGREGATE_EXPORTED",
            entity_type="IndustryDashboard", entity_id=digest,
            after={"rows": count, "sha256": digest}, ip_address=request.META.get("REMOTE_ADDR"))
        response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="nca_industry_dashboard_aggregate.xlsx"'
        response["X-Content-Type-Options"] = "nosniff"
        return response


class IndustryDashboardDataView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from .industry_dashboard import dataset_for_user, load_dashboard_dataset
        response = Response(dataset_for_user(load_dashboard_dataset(), request.user))
        response["Cache-Control"] = "private, no-store"
        response["X-Content-Type-Options"] = "nosniff"
        return response
