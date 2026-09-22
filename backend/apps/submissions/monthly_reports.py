import hashlib
import os
import re
import shutil
import uuid
import zipfile
from collections import Counter
from copy import copy
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

from django.conf import settings
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import from_excel

from apps.audit.services import record_audit
from apps.forms_engine.models import GridRow
from .models import (
    MonthlyReportArtifact, ProviderWorkbookBaseline, Submission,
    SubmissionValue, WorkbookIndicatorMapping,
)


def field_target_key(field):
    return f"field:{field.section.section_code}:{field.field_code}".lower()


def grid_target_key(grid, row_label, column):
    return f"grid:{grid.section.section_code}:{grid.grid_code}:{row_label}:{column.column_code}".lower()


def _fixed_row_label(value):
    if not value.grid_id or not value.grid_row_id or not value.grid or value.grid.row_mode != "FIXED":
        return None
    try:
        return GridRow.objects.only("row_label").get(pk=int(value.grid_row_id), grid_id=value.grid_id).row_label
    except (GridRow.DoesNotExist, TypeError, ValueError):
        return None


def value_target_key(value):
    if value.field_id and value.field:
        return field_target_key(value.field)
    row_label = _fixed_row_label(value)
    if row_label and value.grid_column:
        return grid_target_key(value.grid, row_label, value.grid_column)
    return None


def previous_month_values(submission):
    if submission.expected.form_template_id is None:
        return {"period": None, "values": {}}
    period = submission.expected.period
    if period.frequency == "MONTHLY" and period.month:
        previous_year = period.year - 1 if period.month == 1 else period.year
        previous_month = 12 if period.month == 1 else period.month - 1
        period_filter = {
            "period__frequency": "MONTHLY", "period__year": previous_year,
            "period__month": previous_month,
        }
    elif period.frequency == "ANNUAL":
        previous_year = period.year - 1
        previous_month = None
        period_filter = {"period__frequency": "ANNUAL", "period__year": previous_year}
    else:
        return {"period": None, "values": {}}
    prior_expected = (
        submission.expected.__class__.objects.filter(
            provider_id=submission.expected.provider_id,
            form_template__family_id=submission.expected.form_template.family_id,
            workflow_status="APPROVED",
            **period_filter,
        )
        .select_related("period")
        .order_by("-id")
        .first()
    )
    if not prior_expected:
        return {"period": {"year": previous_year, "month": previous_month}, "values": {}}
    prior = prior_expected.versions.filter(reviewed_at__isnull=False).order_by("-reviewed_at", "-version").first()
    if not prior:
        return {"period": {"year": previous_year, "month": previous_month}, "values": {}}
    values = {}
    queryset = prior.values.select_related(
        "field__section", "grid__section", "grid_column",
    )
    for item in queryset:
        key = value_target_key(item)
        if key:
            values[key] = item.value if item.value_status in {"PROVIDED", "SYSTEM_CALCULATED"} else None
    return {
        "period": {
            "id": prior_expected.period_id, "name": prior_expected.period.name,
            "year": previous_year, "month": previous_month,
            "submission_id": prior.id,
        },
        "values": values,
    }


def baseline_readiness(baseline):
    issues = []
    if baseline.status != "ACTIVE":
        issues.append("The provider workbook baseline is not active.")
    if baseline.scan_status != "CLEAN":
        issues.append("The provider workbook baseline has not passed malware scanning.")
    source_path = Path(settings.PRIVATE_UPLOAD_ROOT) / baseline.storage_path
    try:
        source_path.resolve().relative_to(Path(settings.PRIVATE_UPLOAD_ROOT).resolve())
    except ValueError:
        issues.append("The workbook baseline path is outside private storage.")
    if not source_path.exists():
        issues.append("The private workbook baseline file is missing.")
    mappings = baseline.indicator_mappings.all()
    if not mappings.exists():
        issues.append("No workbook indicator mappings are configured.")
    if mappings.filter(verified=False).exists():
        issues.append("Every workbook indicator mapping must be verified.")
    if mappings.filter(value_kind="CALCULATED", formula_seed_column__isnull=True).exists():
        issues.append("Every calculated row requires a verified formula seed.")
    if source_path.exists() and baseline.scan_status == "CLEAN" and mappings.exists():
        workbook = None
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
            for mapping in mappings:
                if mapping.sheet_name not in workbook.sheetnames:
                    issues.append(f'Mapped worksheet "{mapping.sheet_name}" is missing.')
                    continue
                sheet = workbook[mapping.sheet_name]
                if mapping.excel_row > sheet.max_row:
                    issues.append(
                        f"Mapped row {mapping.sheet_name}!{mapping.excel_row} is outside the worksheet."
                    )
                    continue
                if mapping.value_kind == "CALCULATED":
                    seed = sheet.cell(mapping.excel_row, mapping.formula_seed_column or 0)
                    if not isinstance(seed.value, str) or not seed.value.startswith("="):
                        issues.append(
                            f"Calculated row {mapping.sheet_name}!{mapping.excel_row} has no valid formula seed."
                        )
        except Exception as exc:
            issues.append(f"The workbook baseline could not be validated: {exc}")
        finally:
            if workbook is not None:
                workbook.close()
    return issues


def _month_key(value, workbook):
    if isinstance(value, (date, datetime)):
        return value.year, value.month
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value, workbook.epoch)
            return parsed.year, parsed.month
        except (TypeError, ValueError, OverflowError):
            return None
    return None


def replace_baseline_mappings(baseline, mapping_payload):
    """Replace a draft mapping set after validating every target against the exact form version."""
    if baseline.status != "DRAFT":
        raise ValueError("Only a draft workbook baseline can be mapped.")
    template = baseline.form_template
    source_path = Path(settings.PRIVATE_UPLOAD_ROOT) / baseline.storage_path
    workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
    seen_targets = set()
    seen_rows = set()
    prepared = []
    for item in mapping_payload:
        target_type = str(item.get("target_type") or "").upper()
        sheet_name = str(item.get("sheet_name") or "").strip()
        excel_row = int(item.get("excel_row") or 0)
        if target_type not in {"FIELD", "GRID_CELL"} or not sheet_name or excel_row < 1:
            raise ValueError("Every mapping requires a valid target type, sheet and row.")
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f'Mapped worksheet "{sheet_name}" does not exist in the baseline.')
        if excel_row > workbook[sheet_name].max_row:
            raise ValueError(f"Mapped row {sheet_name}!{excel_row} is outside the worksheet.")
        field = grid = grid_row = grid_column = None
        if target_type == "FIELD":
            field = template.sections.filter(fields__id=item.get("field")).values_list("fields", flat=True).first()
            if not field:
                raise ValueError("A mapped field does not belong to this form version.")
            from apps.forms_engine.models import FormField
            field = FormField.objects.select_related("section").get(pk=field)
            target_key = field_target_key(field)
        else:
            from apps.forms_engine.models import FormGrid, GridColumn
            grid = FormGrid.objects.select_related("section").filter(
                pk=item.get("grid"), section__form_template=template, row_mode="FIXED",
            ).first()
            grid_row = GridRow.objects.filter(pk=item.get("grid_row"), grid=grid).first() if grid else None
            grid_column = GridColumn.objects.filter(pk=item.get("grid_column"), grid=grid).first() if grid else None
            if not grid or not grid_row or not grid_column:
                raise ValueError("A mapped grid cell does not belong to this fixed grid.")
            target_key = grid_target_key(grid, grid_row.row_label, grid_column)
        if target_key in seen_targets:
            raise ValueError(f"Duplicate mapping target: {target_key}.")
        row_key = (sheet_name.casefold(), excel_row)
        if row_key in seen_rows:
            raise ValueError(f"Workbook row {sheet_name}!{excel_row} is mapped more than once.")
        seen_targets.add(target_key)
        seen_rows.add(row_key)
        prepared.append(WorkbookIndicatorMapping(
            baseline=baseline, target_key=target_key, target_type=target_type,
            field=field if target_type == "FIELD" else None,
            grid=grid if target_type == "GRID_CELL" else None,
            grid_row=grid_row if target_type == "GRID_CELL" else None,
            grid_column=grid_column if target_type == "GRID_CELL" else None,
            sheet_name=sheet_name, excel_row=excel_row,
            value_kind=str(item.get("value_kind") or "INPUT").upper(),
            source_indicator=str(item.get("source_indicator") or "")[:500],
            source_definition=str(item.get("source_definition") or ""),
            source_data_type=str(item.get("source_data_type") or "")[:100],
            formula_seed_column=item.get("formula_seed_column") or None,
            verified=bool(item.get("verified", False)), provenance=item.get("provenance") or {},
        ))
    baseline.indicator_mappings.all().delete()
    WorkbookIndicatorMapping.objects.bulk_create(prepared)
    baseline.mapping_summary = {
        "total": len(prepared),
        "verified": sum(1 for item in prepared if item.verified),
        "calculated": sum(1 for item in prepared if item.value_kind == "CALCULATED"),
    }
    baseline.save(update_fields=["mapping_summary", "updated_at"])
    return prepared


def suggest_exact_baseline_mappings(baseline):
    """Build exact-label suggestions once; persisted mappings are used for every export."""
    source_path = Path(settings.PRIVATE_UPLOAD_ROOT) / baseline.storage_path
    workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
    if baseline.main_sheet not in workbook.sheetnames:
        baseline.mapping_summary = {"total": 0, "unmatched": [], "error": "Main worksheet is missing."}
        baseline.save(update_fields=["mapping_summary", "updated_at"])
        return []
    sheet = workbook[baseline.main_sheet]

    def normalize(value):
        return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())

    rows_by_label = {}
    row_metadata = {}
    duplicates = set()
    for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        indicator = cells[2] if len(cells) > 2 else None
        row_metadata[row_number] = {
            "indicator": indicator,
            "definition": cells[3] if len(cells) > 3 else None,
            "data_type": cells[4] if len(cells) > 4 else None,
            "formula_columns": [
                column for column, value in enumerate(cells, start=1)
                if column >= baseline.first_month_column
                and isinstance(value, str) and value.startswith("=")
            ],
        }
        key = normalize(indicator)
        if not key:
            continue
        if key in rows_by_label:
            duplicates.add(key)
        else:
            rows_by_label[key] = row_number
    schema_sections = {}
    workbook_import = getattr(baseline.form_template, "workbook_import", None)
    if workbook_import:
        schema_sections = {
            section.get("section_code"): section
            for section in (workbook_import.detected_schema or {}).get("sections", [])
        }
    source_fields = {
        (section_code, item.get("field_code")): item
        for section_code, section in schema_sections.items()
        for item in section.get("fields", [])
    }
    mappings = []
    unresolved = []
    section_offsets = {}
    sections = list(baseline.form_template.sections.prefetch_related("fields", "grids__columns", "grids__fixed_rows"))
    for section in sections:
        offsets = []
        for target in section.fields.select_related("section"):
            key = normalize(target.label)
            row = rows_by_label.get(key) if key not in duplicates else None
            if not row:
                unresolved.append((section, target))
                continue
            metadata = row_metadata[row]
            formula_columns = metadata["formula_columns"]
            source = source_fields.get((section.section_code, target.field_code), {}).get("source") or {}
            source_row = source.get("row")
            if source_row:
                offsets.append(row - int(source_row))
                if not target.source_sheet:
                    target.source_sheet = str(source.get("sheet") or "")[:255]
                    target.source_row = int(source_row)
                    target.save(update_fields=["source_sheet", "source_row"])
            mappings.append({
                "target_type": "FIELD", "field": target.id,
                "sheet_name": baseline.main_sheet, "excel_row": row,
                "value_kind": "CALCULATED" if formula_columns else "INPUT",
                "formula_seed_column": formula_columns[-1] if formula_columns else None,
                "source_indicator": str(metadata["indicator"] or ""),
                "source_definition": str(metadata["definition"] or ""),
                "source_data_type": str(metadata["data_type"] or ""),
                "verified": True,
                "provenance": {"strategy": "exact-normalized-label", "baseline_sha256": baseline.sha256},
            })
        if offsets:
            section_offsets[section.section_code] = Counter(offsets).most_common(1)[0][0]

    unmatched = []
    for section, target in unresolved:
        source = source_fields.get((section.section_code, target.field_code), {}).get("source") or {}
        source_row = source.get("row")
        offset = section_offsets.get(section.section_code)
        row = int(source_row) + int(offset) if source_row and offset is not None else None
        metadata = row_metadata.get(row) if row else None
        if not metadata or normalize(metadata["indicator"]) != normalize(target.label):
            unmatched.append({"target_key": field_target_key(target), "label": target.label})
            continue
        formula_columns = metadata["formula_columns"]
        target.source_sheet = str(source.get("sheet") or "")[:255]
        target.source_row = int(source_row)
        target.save(update_fields=["source_sheet", "source_row"])
        mappings.append({
            "target_type": "FIELD", "field": target.id,
            "sheet_name": baseline.main_sheet, "excel_row": row,
            "value_kind": "CALCULATED" if formula_columns else "INPUT",
            "formula_seed_column": formula_columns[-1] if formula_columns else None,
            "source_indicator": str(metadata["indicator"] or ""),
            "source_definition": str(metadata["definition"] or ""),
            "source_data_type": str(metadata["data_type"] or ""),
            "verified": True,
            "provenance": {
                "strategy": "exact-label-with-verified-section-row-offset",
                "section_offset": offset, "source_row": source_row,
                "baseline_sha256": baseline.sha256,
            },
        })

    # Fixed matrix cells retain two independent workbook rows per logical row.
    for section in sections:
        offset = section_offsets.get(section.section_code, 6)
        schema_grids = {item.get("grid_code"): item for item in schema_sections.get(section.section_code, {}).get("grids", [])}
        for grid in section.grids.filter(row_mode="FIXED").prefetch_related("columns", "fixed_rows"):
            schema_grid = schema_grids.get(grid.grid_code) or {}
            source_rows = {item.get("label"): (item.get("source") or {}) for item in schema_grid.get("fixed_row_sources", [])}
            columns = list(grid.columns.order_by("sort_order", "id"))
            for fixed_row in grid.fixed_rows.order_by("sort_order", "id"):
                source = source_rows.get(fixed_row.row_label) or {}
                rows = list(source.get("rows") or [])
                if len(rows) != len(columns):
                    unmatched.append({"target_key": f"grid:{section.section_code}:{grid.grid_code}:{fixed_row.row_label}", "label": fixed_row.row_label})
                    continue
                fixed_row.source_sheet = str(source.get("sheet") or "")[:255]
                fixed_row.source_rows = rows
                fixed_row.save(update_fields=["source_sheet", "source_rows"])
                for index, column in enumerate(columns):
                    excel_row = int(rows[index]) + int(offset)
                    metadata = row_metadata.get(excel_row) or {}
                    if index == 0 and normalize(metadata.get("indicator")) != normalize(fixed_row.row_label):
                        unmatched.append({"target_key": grid_target_key(grid, fixed_row.row_label, column), "label": fixed_row.row_label})
                        continue
                    mappings.append({
                        "target_type": "GRID_CELL", "grid": grid.id, "grid_row": fixed_row.id,
                        "grid_column": column.id, "sheet_name": baseline.main_sheet, "excel_row": excel_row,
                        "value_kind": "INPUT", "source_indicator": fixed_row.row_label,
                        "source_definition": str(metadata.get("definition") or column.label),
                        "source_data_type": str(metadata.get("data_type") or "number"), "verified": True,
                        "provenance": {
                            "strategy": "fixed-grid-source-row-with-verified-section-offset",
                            "section_offset": offset, "source_row": rows[index], "baseline_sha256": baseline.sha256,
                        },
                    })
    replace_baseline_mappings(baseline, mappings)
    baseline.mapping_summary = {
        **baseline.mapping_summary,
        "unmatched": unmatched,
        "unmatched_count": len(unmatched),
        "duplicate_source_labels": sorted(duplicates),
    }
    baseline.save(update_fields=["mapping_summary", "updated_at"])
    return mappings


def _safe_filename(value):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "Provider").strip("._")
    return cleaned[:80] or "Provider"


def _coerce_excel_value(mapping, raw_value):
    if raw_value in (None, ""):
        return None
    data_type = (mapping.source_data_type or "").lower()
    field_type = mapping.field.field_type if mapping.field_id else (
        mapping.grid_column.field_type if mapping.grid_column_id else ""
    )
    if field_type in {"number", "currency", "percentage"} or any(
        token in data_type for token in ("number", "numeric", "integer", "decimal", "currency", "percent")
    ):
        try:
            return float(str(raw_value).replace(",", "").strip())
        except ValueError:
            return raw_value
    return raw_value


def _copy_column_presentation(sheet, source_column, target_column):
    source_letter = sheet.cell(1, source_column).column_letter
    target_letter = sheet.cell(1, target_column).column_letter
    source_dimension = sheet.column_dimensions[source_letter]
    target_dimension = sheet.column_dimensions[target_letter]
    target_dimension.width = source_dimension.width
    target_dimension.hidden = source_dimension.hidden
    target_dimension.bestFit = source_dimension.bestFit
    target_dimension.outlineLevel = source_dimension.outlineLevel
    for row in range(1, sheet.max_row + 1):
        source = sheet.cell(row, source_column)
        target = sheet.cell(row, target_column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _mapping_value(submission, mapping, values_by_field, values_by_grid):
    if mapping.field_id:
        item = values_by_field.get(mapping.field_id)
    else:
        item = values_by_grid.get((mapping.grid_id, str(mapping.grid_row_id), mapping.grid_column_id))
    if not item or item.value_status not in {"PROVIDED", "SYSTEM_CALCULATED"}:
        return None
    return _coerce_excel_value(mapping, item.value)


def _ensure_month_column(workbook, sheet, baseline, period):
    month_columns = {}
    for column in range(baseline.first_month_column, sheet.max_column + 1):
        key = _month_key(sheet.cell(baseline.month_header_row, column).value, workbook)
        if key:
            month_columns[key] = column
    if not month_columns:
        raise ValueError(f'Worksheet "{sheet.title}" has no recognizable monthly date columns.')
    target_key = (period.year, period.month)
    target_column = month_columns.get(target_key)
    inserted = False
    if target_column is None:
        last_column = max(month_columns.values())
        target_column = last_column + 1
        sheet.insert_cols(target_column, 1)
        _copy_column_presentation(sheet, last_column, target_column)
        inserted = True
    sheet.cell(baseline.month_header_row, target_column).value = date(period.year, period.month, 1)
    return target_column, inserted


def _preserve_package_extensions(source_path, output_path, main_sheet_number):
    """Restore safe package parts that openpyxl cannot round-trip.

    Formula calculation chains are deliberately not restored because inserting a
    reporting column makes the original chain stale; Excel rebuilds it on open.
    """
    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    package_rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rels_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ElementTree.register_namespace("", content_types_ns)
    ElementTree.register_namespace("", package_rels_ns)
    ElementTree.register_namespace("r", office_rels_ns)

    with zipfile.ZipFile(source_path, "r") as source_zip, zipfile.ZipFile(output_path, "r") as output_zip:
        source_entries = {info.filename: (info, source_zip.read(info.filename)) for info in source_zip.infolist()}
        output_entries = {info.filename: (info, output_zip.read(info.filename)) for info in output_zip.infolist()}

    restored_custom_xml = 0
    for name, payload in source_entries.items():
        if name.startswith("customXml/"):
            output_entries[name] = payload
            restored_custom_xml += 1

    content_types_name = "[Content_Types].xml"
    content_types = ElementTree.fromstring(output_entries[content_types_name][1])
    source_content_types = ElementTree.fromstring(source_entries[content_types_name][1])
    defaults = {(item.get("Extension"), item.get("ContentType")) for item in content_types}
    overrides = {item.get("PartName") for item in content_types}
    for item in source_content_types:
        extension = item.get("Extension")
        part_name = item.get("PartName")
        if extension == "bin" and (extension, item.get("ContentType")) not in defaults:
            content_types.append(copy(item))
        elif part_name and part_name.startswith("/customXml/") and part_name not in overrides:
            content_types.append(copy(item))

    restored_printer_settings = 0
    sheet_name = f"xl/worksheets/sheet{main_sheet_number}.xml"
    rels_name = f"xl/worksheets/_rels/sheet{main_sheet_number}.xml.rels"
    source_rels = source_entries.get(rels_name)
    if source_rels and sheet_name in output_entries:
        source_rel_root = ElementTree.fromstring(source_rels[1])
        printer_rel = next((
            item for item in source_rel_root
            if (item.get("Type") or "").endswith("/printerSettings")
        ), None)
        if printer_rel is not None:
            printer_target = (printer_rel.get("Target") or "").split("/")[-1]
            source_printer_name = f"xl/printerSettings/{printer_target}"
            if source_printer_name in source_entries:
                output_entries[source_printer_name] = source_entries[source_printer_name]
                if rels_name in output_entries:
                    output_rel_root = ElementTree.fromstring(output_entries[rels_name][1])
                    rel_info = output_entries[rels_name][0]
                else:
                    output_rel_root = ElementTree.Element(f"{{{package_rels_ns}}}Relationships")
                    rel_info = zipfile.ZipInfo(rels_name)
                relationship_id = "ncaPrinterSettings"
                if not any((item.get("Type") or "").endswith("/printerSettings") for item in output_rel_root):
                    ElementTree.SubElement(output_rel_root, f"{{{package_rels_ns}}}Relationship", {
                        "Id": relationship_id,
                        "Type": f"{office_rels_ns}/printerSettings",
                        "Target": f"../printerSettings/{printer_target}",
                    })
                output_entries[rels_name] = (
                    rel_info, ElementTree.tostring(output_rel_root, encoding="utf-8", xml_declaration=True),
                )
                sheet_root = ElementTree.fromstring(output_entries[sheet_name][1])
                page_setup = next((item for item in sheet_root if item.tag.endswith("}pageSetup")), None)
                if page_setup is not None:
                    page_setup.set(f"{{{office_rels_ns}}}id", relationship_id)
                    output_entries[sheet_name] = (
                        output_entries[sheet_name][0],
                        ElementTree.tostring(sheet_root, encoding="utf-8", xml_declaration=True),
                    )
                restored_printer_settings = 1

    output_entries[content_types_name] = (
        output_entries[content_types_name][0],
        ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True),
    )
    temporary_path = output_path.with_suffix(".package-preservation.tmp")
    with zipfile.ZipFile(temporary_path, "w") as rebuilt:
        for _, (info, payload) in output_entries.items():
            rebuilt.writestr(info, payload)
    os.replace(temporary_path, output_path)
    return {
        "restored_custom_xml_parts": restored_custom_xml,
        "restored_printer_settings_parts": restored_printer_settings,
        "intentionally_rebuilt_by_excel": ["xl/calcChain.xml"],
    }


def generate_monthly_report(submission_id, actor=None):
    submission = Submission.objects.select_related(
        "expected__provider", "expected__period", "expected__form_template__family", "submitted_by",
    ).get(pk=submission_id)
    if submission.expected.workflow_status not in {"SUBMITTED", "UNDER_REVIEW", "RESUBMITTED", "APPROVED"}:
        raise ValueError("The Provider Approver must officially submit the form before Excel generation.")
    baseline = ProviderWorkbookBaseline.objects.select_related("provider", "contact").get(
        provider_id=submission.expected.provider_id,
        form_template_id=submission.expected.form_template_id,
        status="ACTIVE",
    )
    readiness_issues = baseline_readiness(baseline)
    if readiness_issues:
        raise ValueError(" ".join(readiness_issues))
    period = submission.expected.period
    if period.frequency != "MONTHLY" or not period.month:
        raise ValueError("Monthly report generation requires a monthly reporting period.")

    artifact, _ = MonthlyReportArtifact.objects.get_or_create(
        submission=submission, defaults={"baseline": baseline, "submission_revision": submission.revision},
    )
    artifact.baseline = baseline
    artifact.status = "PREPARING"
    artifact.error_message = ""
    artifact.submission_revision = submission.revision
    artifact.save(update_fields=["baseline", "status", "error_message", "submission_revision", "updated_at"])

    source_path = Path(settings.PRIVATE_UPLOAD_ROOT) / baseline.storage_path
    output_directory = Path(settings.PRIVATE_EXPORT_ROOT) / "monthly-reports" / str(submission.id)
    output_directory.mkdir(parents=True, exist_ok=True)
    filename = f"{_safe_filename(submission.expected.provider.trade_name or submission.expected.provider.registered_name)}_Monthly_Report_{date(period.year, period.month, 1).strftime('%B_%Y')}.xlsx"
    output_path = output_directory / f"{uuid.uuid4().hex}-{filename}"
    shutil.copy2(source_path, output_path)

    workbook = load_workbook(output_path, data_only=False, keep_links=True)
    if baseline.main_sheet not in workbook.sheetnames:
        raise ValueError(f'Workbook sheet "{baseline.main_sheet}" is missing.')
    sheet = workbook[baseline.main_sheet]

    values = list(submission.values.select_related("field", "grid", "grid_column"))
    values_by_field = {item.field_id: item for item in values if item.field_id}
    values_by_grid = {
        (item.grid_id, item.grid_row_id, item.grid_column_id): item
        for item in values if item.grid_id
    }
    mappings = list(baseline.indicator_mappings.select_related("field", "grid", "grid_row", "grid_column"))
    mapped_sheet_names = list(dict.fromkeys(
        [baseline.main_sheet, *(mapping.sheet_name for mapping in mappings)]
    ))
    sheet_columns = {}
    for sheet_name in mapped_sheet_names:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f'Mapped worksheet "{sheet_name}" is missing.')
        mapped_sheet = workbook[sheet_name]
        target_column, inserted = _ensure_month_column(workbook, mapped_sheet, baseline, period)
        sheet_columns[sheet_name] = {"column": target_column, "inserted": inserted}

    input_count = formula_count = 0
    for mapping in mappings:
        mapped_sheet = workbook[mapping.sheet_name]
        target_column = sheet_columns[mapping.sheet_name]["column"]
        target_cell = mapped_sheet.cell(mapping.excel_row, target_column)
        if mapping.value_kind == "CALCULATED":
            seed_column = mapping.formula_seed_column
            seed = mapped_sheet.cell(mapping.excel_row, seed_column) if seed_column else None
            if not seed or not isinstance(seed.value, str) or not seed.value.startswith("="):
                raise ValueError(f"Calculated row {mapping.excel_row} has no valid formula seed.")
            target_cell.value = Translator(seed.value, origin=seed.coordinate).translate_formula(target_cell.coordinate)
            formula_count += 1
        else:
            target_cell.value = _mapping_value(submission, mapping, values_by_field, values_by_grid)
            input_count += 1

    contact = baseline.contact
    contact_values = {
        "provider_name": submission.expected.provider.trade_name or submission.expected.provider.registered_name,
        "contact_name": contact.name if contact else "",
        "contact_email": contact.email if contact else submission.expected.provider.primary_email,
        "contact_phone": contact.phone if contact else submission.expected.provider.primary_phone,
    }
    for key, default_cell in {
        "provider_name": "F3", "contact_name": "F4", "contact_email": "F5", "contact_phone": "F6",
    }.items():
        cell = baseline.contact_cells.get(key, default_cell)
        if cell:
            sheet[cell] = contact_values[key]

    workbook.save(output_path)
    workbook.close()
    package_preservation = _preserve_package_extensions(
        source_path, output_path, workbook.sheetnames.index(baseline.main_sheet) + 1,
    )
    sha256 = hashlib.sha256(output_path.read_bytes()).hexdigest()
    artifact.filename = filename
    artifact.private_path = str(output_path)
    artifact.file_size = output_path.stat().st_size
    artifact.sha256 = sha256
    artifact.status = "READY"
    artifact.completed_at = timezone.now()
    artifact.generation_metadata = {
        "submission_reference": submission.submission_reference,
        "target_year": period.year, "target_month": period.month,
        "target_column": sheet_columns[baseline.main_sheet]["column"],
        "month_column_inserted": sheet_columns[baseline.main_sheet]["inserted"],
        "sheet_columns": sheet_columns,
        "input_mapping_count": input_count, "formula_mapping_count": formula_count,
        "baseline_sha256": baseline.sha256, "workbook_sheets": workbook.sheetnames,
        "package_preservation": package_preservation,
    }
    artifact.save()
    record_audit(
        user=actor or submission.submitted_by, action="MONTHLY_REPORT_GENERATED",
        entity_type="MonthlyReportArtifact", entity_id=artifact.id,
        after={"submission_id": submission.id, "sha256": sha256, **artifact.generation_metadata},
    )
    return artifact
